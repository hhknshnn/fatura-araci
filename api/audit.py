# api/audit.py
# İşlem kaydı (audit log) — kim, ne zaman, ne yaptı. Sadece admin görebilir.

import io
import time
from flask import jsonify, send_file
from api.db import get_conn


def log_action(session, action, description):
    """session: flask.g.user (dict: username/displayName/role). session yoksa sessizce geçilir."""
    if not session:
        return
    conn = get_conn()
    cur  = conn.cursor()
    cur.execute('''
        INSERT INTO audit_log (username, display_name, role, action, description, created_at)
        VALUES (%s, %s, %s, %s, %s, %s)
    ''', (
        session.get('username', ''),
        session.get('displayName', session.get('username', '')),
        session.get('role', ''),
        action,
        description,
        int(time.time()),
    ))
    conn.commit()
    cur.close()
    conn.close()


def get_audit_log(limit=300):
    conn = get_conn()
    cur  = conn.cursor()
    cur.execute('''
        SELECT username, display_name, role, action, description, created_at
        FROM audit_log ORDER BY created_at DESC LIMIT %s
    ''', (limit,))
    rows = cur.fetchall()
    cur.close()
    conn.close()
    return [{
        'username':    r[0],
        'displayName': r[1],
        'role':        r[2],
        'action':      r[3],
        'description': r[4],
        'createdAt':   r[5],
    } for r in rows]


def audit_log_get():
    """GET /api/audit-log — admin only (route seviyesinde @require_auth ile korunur)"""
    return jsonify({'success': True, 'entries': get_audit_log()})


def audit_log_export():
    """GET /api/audit-log/export — işlem kayıtlarını Excel olarak indirir (admin only)"""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    entries = get_audit_log(limit=5000)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'İşlem Kayıtları'

    headers = ['Zaman', 'Kullanıcı', 'Rol', 'İşlem', 'Detay']
    header_fill = PatternFill('solid', fgColor='1F3864')
    header_font = Font(name='Arial', bold=True, color='FFFFFF', size=10)
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font      = header_font
        cell.fill      = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for row_idx, e in enumerate(entries, start=2):
        tarih = time.strftime('%Y-%m-%d %H:%M:%S', time.localtime(e['createdAt']))
        ws.cell(row=row_idx, column=1, value=tarih)
        ws.cell(row=row_idx, column=2, value=e['displayName'] or e['username'])
        ws.cell(row=row_idx, column=3, value=e['role'])
        ws.cell(row=row_idx, column=4, value=e['action'])
        ws.cell(row=row_idx, column=5, value=e['description'])

    for col_idx in range(1, len(headers) + 1):
        col_letter = ws.cell(row=1, column=col_idx).column_letter
        max_len    = len(str(ws.cell(row=1, column=col_idx).value or ''))
        for row_idx in range(2, len(entries) + 2):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[col_letter].width = min(max_len + 4, 60)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return send_file(
        buf,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f'islem_kayitlari_{time.strftime("%Y-%m-%d")}.xlsx',
    )
