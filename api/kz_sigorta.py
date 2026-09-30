# api/kz_sigorta.py
"""Kazakistan — Nakliyat Sigortası Bildirim Formu (sigorta poliçe talimatı).

INV+PL üretiminde Price List gibi ek evrak olarak iner. `templates/sigorta_kz.xlsx`
şablonunu doldurur. Her KZ faturasının sigorta bilgisi (kap, brüt, fatura TL,
ürün grupları) `kz_sigorta_bilgi` tablosunda saklanır; gruplu sevkte (ANT + İHR aynı
araç) partner faturanın bilgisi buradan okunup tek forma birleştirilir.

Partner tespiti: navlun_bekleyen_tahsis (gruplu taslak eşleşmesi) + aynı sefer_id.
Partner bu modülden önce üretildiyse (tabloda yoksa) shipments kaydından
kap/fatura TL okunur; brüt ve ürün grupları o durumda eksik kalır.
"""
import io
import os
import re
from datetime import date, datetime

import openpyxl

from api.db import get_conn

TEMPLATE_PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                             'templates', 'sigorta_kz.xlsx')


def _kap_sayisi(kap):
    m = re.search(r'\d+', str(kap or ''))
    return int(m.group(0)) if m else 0


def _tr_sayi(n):
    """2649.57 → '2649,57' (formdaki yazım)."""
    return f'{float(n):.2f}'.replace('.', ',')


def _depo_tipi(fatura_no, depo_tipi=None):
    if str(fatura_no or '').upper().startswith('ANT'):
        return 'antrepo'
    if str(fatura_no or '').upper().startswith('IHR'):
        return 'serbest'
    return depo_tipi or 'serbest'


def _tarih(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    try:
        return datetime.strptime(str(v)[:10], '%Y-%m-%d').date()
    except (TypeError, ValueError):
        return None


def _bilgi_kaydet(cur, b):
    cur.execute('''
        INSERT INTO kz_sigorta_bilgi
            (fatura_no, dosya_no, depo_tipi, kap, brut_kg, fatura_tl, urun_gruplari,
             plaka, yukleme_tarihi, guncelleme)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, now())
        ON CONFLICT (fatura_no) DO UPDATE SET
            dosya_no = EXCLUDED.dosya_no, depo_tipi = EXCLUDED.depo_tipi,
            kap = EXCLUDED.kap, brut_kg = EXCLUDED.brut_kg,
            fatura_tl = EXCLUDED.fatura_tl, urun_gruplari = EXCLUDED.urun_gruplari,
            plaka = EXCLUDED.plaka, yukleme_tarihi = EXCLUDED.yukleme_tarihi,
            guncelleme = now()
    ''', (b['fatura_no'], b['dosya_no'] or None, b['depo_tipi'], b['kap'], b['brut_kg'],
          b['fatura_tl'], '|'.join(b['urun_gruplari']), b['plaka'] or None, b['yukleme_tarihi']))


def _partner_bilgileri(cur, dosya_no, fatura_no):
    """Gruplu partner faturalarının bilgileri (kendisi hariç)."""
    if not dosya_no:
        return []
    from api.navlun import _partner_dosyalari
    dosyalar = set(_partner_dosyalari(cur, dosya_no))
    cur.execute('''
        SELECT s2.ihracat_dosya_no FROM shipments s1
        JOIN shipments s2 ON s2.sefer_id = s1.sefer_id AND s2.id <> s1.id
        WHERE s1.ihracat_dosya_no = %s AND s1.sefer_id IS NOT NULL
    ''', (dosya_no,))
    dosyalar.update(r[0] for r in cur.fetchall() if r[0])
    dosyalar.discard(dosya_no)
    if not dosyalar:
        return []

    sonuc = []
    cur.execute('''
        SELECT fatura_no, dosya_no, depo_tipi, kap, brut_kg, fatura_tl, urun_gruplari,
               plaka, yukleme_tarihi
        FROM kz_sigorta_bilgi WHERE dosya_no = ANY(%s) AND fatura_no <> %s
    ''', (list(dosyalar), fatura_no))
    for r in cur.fetchall():
        sonuc.append({
            'fatura_no': r[0], 'dosya_no': r[1], 'depo_tipi': r[2], 'kap': r[3] or 0,
            'brut_kg': float(r[4] or 0), 'fatura_tl': float(r[5] or 0),
            'urun_gruplari': [g for g in (r[6] or '').split('|') if g],
            'plaka': r[7] or '', 'yukleme_tarihi': r[8],
        })
    bulunan = {p['dosya_no'] for p in sonuc}

    # Bu modülden önce üretilmiş partner: shipments kaydından (brüt/grup yok)
    eksik = [d for d in dosyalar if d not in bulunan]
    if eksik:
        cur.execute('''
            SELECT fatura_no, ihracat_dosya_no, palet, fatura_bedeli_tl, plaka, yukleme_tarihi
            FROM shipments WHERE ihracat_dosya_no = ANY(%s) AND fatura_no <> %s
        ''', (eksik, fatura_no))
        for r in cur.fetchall():
            sonuc.append({
                'fatura_no': r[0], 'dosya_no': r[1], 'depo_tipi': _depo_tipi(r[0]),
                'kap': _kap_sayisi(r[2]), 'brut_kg': 0.0, 'fatura_tl': float(r[3] or 0),
                'urun_gruplari': [], 'plaka': r[4] or '', 'yukleme_tarihi': r[5],
            })
    return sonuc


def _form_doldur(kalemler):
    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    ws = wb.active

    # Ürün grupları — sırayı koruyarak tekilleştir
    gruplar = []
    for k in kalemler:
        for g in k['urun_gruplari']:
            if g not in gruplar:
                gruplar.append(g)
    ws['C14'] = 'Product Type : ' + ','.join(gruplar) if gruplar else None

    miktar = []
    for k in kalemler:
        etiket = 'Bonded Warehouse' if k['depo_tipi'] == 'antrepo' else 'Warehouse'
        satir = f"{etiket} {k['kap']} packages"
        if k['brut_kg'] > 0:
            satir += f" {_tr_sayi(k['brut_kg'])} BRÜT KG"
        miktar.append(satir)
    ws['C20'] = '\n'.join(miktar)

    tarihler = [t for t in (_tarih(k['yukleme_tarihi']) for k in kalemler) if t]
    ws['E20'] = datetime.combine(max(tarihler), datetime.min.time()) if tarihler else None

    plaka = next((k['plaka'] for k in kalemler if k['plaka']), '')
    ws['I24'] = plaka.replace('-', ' - ') if plaka else None

    ws['C28'] = '\n'.join(f"{k['fatura_no']} // {k['dosya_no']}" if k['dosya_no'] else k['fatura_no']
                          for k in kalemler)
    ws['I32'] = round(sum(k['fatura_tl'] for k in kalemler), 2)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_kz_sigorta(fatura_no, dosya_no, depo_tipi, kap, brut_kg, fatura_tl,
                        urun_gruplari, plaka='', yukleme_tarihi=None):
    """Faturanın bilgisini kaydeder, gruplu partnerle birleştirip formu döner.
    Dönen: (xlsx_bytes, dosya_adi_ref) — ref: dosya no'lar (ANT önce) ya da fatura no."""
    bu = {
        'fatura_no': fatura_no, 'dosya_no': dosya_no or '',
        'depo_tipi': _depo_tipi(fatura_no, depo_tipi), 'kap': _kap_sayisi(kap),
        'brut_kg': float(brut_kg or 0), 'fatura_tl': float(fatura_tl or 0),
        'urun_gruplari': list(urun_gruplari or []), 'plaka': plaka or '',
        'yukleme_tarihi': _tarih(yukleme_tarihi),
    }
    conn = get_conn()
    cur = conn.cursor()
    try:
        _bilgi_kaydet(cur, bu)
        partnerler = _partner_bilgileri(cur, bu['dosya_no'], fatura_no)
        conn.commit()
    finally:
        cur.close()
        conn.close()

    # ANT (Bonded Warehouse) önce, sonra İHR — örnek formdaki sıra
    kalemler = sorted([bu] + partnerler, key=lambda k: (k['depo_tipi'] != 'antrepo', k['fatura_no']))
    ref = ' - '.join(k['dosya_no'] for k in kalemler if k['dosya_no']) or fatura_no
    return _form_doldur(kalemler), ref
