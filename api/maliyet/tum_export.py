# api/maliyet/tum_export.py
# Tümü görünümü: L/T karşılaştırma + lojistik×ülke matrisi → Excel.

import io
from collections import defaultdict

from flask import jsonify, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def _taxes_kod(kod):
    k = str(kod or '').lower()
    return k in ('taxes', 'tax', 'customs', 'vat', 'kdv', 'gumruk', 'gümrük') or 'tax' in k


def _style():
    thin = Border(
        left=Side(style='thin', color='EDE8DF'),
        right=Side(style='thin', color='EDE8DF'),
        top=Side(style='thin', color='EDE8DF'),
        bottom=Side(style='thin', color='EDE8DF'),
    )
    return {
        'thin': thin,
        'ink': Font(name='Calibri', bold=True, color='FFFFFF', size=11),
        'title': Font(name='Calibri', bold=True, size=16, color='1A1916'),
        'sub': Font(name='Calibri', size=10, color='6B7280'),
        'lab': Font(name='Calibri', bold=True, size=10),
        'cell': Font(name='Calibri', size=10),
        'fill_ink': PatternFill('solid', fgColor='1A1916'),
        'fill_teal': PatternFill('solid', fgColor='0F766E'),
        'fill_tot': PatternFill('solid', fgColor='F0FDFA'),
        'fill_head': PatternFill('solid', fgColor='F7F4EE'),
        'right': Alignment(horizontal='right', vertical='center'),
        'left': Alignment(horizontal='left', vertical='center'),
        'center': Alignment(horizontal='center', vertical='center'),
    }


def _header_row(ws, row, values, st, fill=None):
    fill = fill or st['fill_ink']
    for i, v in enumerate(values, 1):
        c = ws.cell(row, i, v)
        c.font = st['ink']
        c.fill = fill
        c.alignment = st['center'] if i > 1 else st['left']
        c.border = st['thin']


def _autosize(ws, max_w=28):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = 10
        for cell in col:
            if cell.value is None:
                continue
            width = max(width, min(max_w, len(str(cell.value)) + 2))
        ws.column_dimensions[letter].width = width


def maliyet_tum_export():
    """POST /api/maliyet/tum/export — Tümü özet + L/T + kalem×ülke Excel."""
    body = request.get_json(silent=True) or {}
    start = str(body.get('start') or '')
    end = str(body.get('end') or '')
    ozet = body.get('ozet') if isinstance(body.get('ozet'), dict) else {}
    ulkeler = body.get('ulkeler') if isinstance(body.get('ulkeler'), list) else []
    kalemler = body.get('kalemler') if isinstance(body.get('kalemler'), list) else []
    oran = body.get('oran') if isinstance(body.get('oran'), dict) else {}

    loj_kalemler = [k for k in kalemler if not _taxes_kod(k.get('kalem_kod'))]
    ulke_loj = defaultdict(float)
    for k in loj_kalemler:
        for kod, tutar in (k.get('ulkeler') or {}).items():
            ulke_loj[kod] += float(tutar or 0)

    ulkeler = sorted(ulkeler, key=lambda u: -(ulke_loj.get(u.get('ulke'), 0) or float(u.get('gercek_eur') or 0)))
    st = _style()
    wb = Workbook()

    # ── Özet ──────────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = 'Ozet'
    ws.cell(1, 1, 'Maliyet Takip — Tümü').font = st['title']
    ws.cell(2, 1, f'Dönem: {start} — {end}').font = st['sub']
    ws.cell(4, 1, 'Gösterge').font = st['lab']
    ws.cell(4, 2, 'Değer').font = st['lab']
    loj = sum(ulke_loj.values())
    toplam = float(ozet.get('gercek_toplam_eur') or 0)
    rows = [
        ('Toplam Lojistik (EUR)', round(loj, 2)),
        ('Toplam Maliyet (EUR)', round(toplam, 2)),
        ('Lojistik / Toplam %', round(loj / toplam * 100, 2) if toplam else None),
        ('Fatura sayısı', ozet.get('fatura_sayisi') or 0),
        ('Ülke sayısı', ozet.get('ulke_sayisi') or len(ulkeler)),
    ]
    o = oran.get('ozet') or {}
    if o:
        rows.extend([
            ('Grup L/T % (cirolu aylar)', o.get('genel_oran')),
            ('Ülke ort. L/T % (basit)', o.get('ortalama_oran')),
            ('Toplam lojistik (tüm aylar)', o.get('toplam_lojistik')),
            ('Lojistik cirolu aylar', o.get('toplam_lojistik_oran')),
            ('En verimli', (o.get('en_iyi') or {}).get('label')),
            ('En yüksek L/T', (o.get('en_yuksek') or {}).get('label')),
        ])
    for i, (lab, val) in enumerate(rows, 5):
        ws.cell(i, 1, lab).font = st['cell']
        cell = ws.cell(i, 2, val)
        cell.font = st['lab']
        if isinstance(val, float):
            cell.number_format = '#,##0.00'
            cell.alignment = st['right']
    _autosize(ws)

    # ── L/T Oran ──────────────────────────────────────────────────────────
    ws2 = wb.create_sheet('LT_Oran')
    ws2.cell(1, 1, 'Logistics / Turnover').font = st['title']
    ws2.cell(2, 1, 'Lojistik ÷ ciro · Δ ort. = ülke ortalamasına fark (pp)').font = st['sub']
    _header_row(ws2, 4, ['Ülke', 'Kod', 'Lojistik EUR', 'Ciro EUR', 'L/T %', 'Δ ort. (pp)'], st)
    ort = o.get('ortalama_oran')
    oran_rows = sorted(
        oran.get('ulkeler') or [],
        key=lambda r: (r.get('oran') is None, r.get('oran') if r.get('oran') is not None else 999),
    )
    for i, r in enumerate(oran_rows, 5):
        oran_v = r.get('oran')
        delta = (oran_v - ort) if oran_v is not None and ort is not None else None
        vals = [
            r.get('label'), r.get('ulke', '').upper(),
            float(r.get('lojistik') or 0), float(r.get('ciro_eur') or 0),
            oran_v, delta,
        ]
        for c, v in enumerate(vals, 1):
            cell = ws2.cell(i, c, v)
            cell.font = st['cell']
            cell.border = st['thin']
            if c >= 3:
                cell.alignment = st['right']
                if c in (3, 4):
                    cell.number_format = '#,##0.00'
                elif c in (5, 6) and isinstance(v, (int, float)):
                    cell.number_format = '0.00'
    r = 5 + len(oran_rows)
    if o:
        ws2.cell(r + 1, 1, 'Grup toplamı').font = st['lab']
        ws2.cell(r + 1, 3, float(o.get('toplam_lojistik') or 0)).number_format = '#,##0.00'
        ws2.cell(r + 1, 4, float(o.get('toplam_ciro') or 0)).number_format = '#,##0.00'
        ws2.cell(r + 1, 5, o.get('genel_oran')).number_format = '0.00'
        for c in range(1, 7):
            ws2.cell(r + 1, c).fill = st['fill_tot']
            ws2.cell(r + 1, c).border = st['thin']
    _autosize(ws2)

    # ── L/T Aylık ─────────────────────────────────────────────────────────
    aylik = oran.get('aylik') or []
    if aylik and oran_rows:
        ws3 = wb.create_sheet('LT_Aylik')
        ws3.cell(1, 1, 'Aylık L/T %').font = st['title']
        headers = ['Ay'] + [r.get('label') or r.get('ulke') for r in oran_rows]
        _header_row(ws3, 3, headers, st)
        for i, a in enumerate(aylik, 4):
            ws3.cell(i, 1, a.get('ay')).font = st['cell']
            ws3.cell(i, 1).border = st['thin']
            bag = a.get('ulkeler') or {}
            for j, r in enumerate(oran_rows, 2):
                v = bag.get(r.get('ulke'))
                cell = ws3.cell(i, j, v)
                cell.font = st['cell']
                cell.border = st['thin']
                cell.alignment = st['right']
                if isinstance(v, (int, float)):
                    cell.number_format = '0.00'
        _autosize(ws3)

    # ── Lojistik × ülke ───────────────────────────────────────────────────
    ws4 = wb.create_sheet('Lojistik_x_Ulke')
    ws4.cell(1, 1, 'Lojistik × ülke').font = st['title']
    ws4.cell(2, 1, 'Tutarlar EUR · KDV / taxes hariç').font = st['sub']
    headers = ['Kalem'] + [(u.get('ulke') or '').upper() for u in ulkeler] + ['Toplam']
    _header_row(ws4, 4, headers, st)
    for i, k in enumerate(loj_kalemler, 5):
        ws4.cell(i, 1, k.get('kalem_ad') or k.get('kalem_kod') or '').font = st['cell']
        ws4.cell(i, 1).border = st['thin']
        for j, u in enumerate(ulkeler, 2):
            t = float((k.get('ulkeler') or {}).get(u.get('ulke')) or 0)
            cell = ws4.cell(i, j, t if t > 0.005 else None)
            cell.font = st['cell']
            cell.border = st['thin']
            cell.alignment = st['right']
            if t > 0.005:
                cell.number_format = '#,##0.00'
        tot = float(k.get('tutar_eur') or 0)
        cell = ws4.cell(i, len(ulkeler) + 2, tot)
        cell.font = st['lab']
        cell.border = st['thin']
        cell.alignment = st['right']
        cell.number_format = '#,##0.00'
        cell.fill = st['fill_tot']

    # Ülke toplam satırı
    r = 5 + len(loj_kalemler)
    ws4.cell(r, 1, 'Toplam lojistik').font = st['lab']
    ws4.cell(r, 1).fill = st['fill_teal']
    ws4.cell(r, 1).font = st['ink']
    for j, u in enumerate(ulkeler, 2):
        t = float(ulke_loj.get(u.get('ulke')) or 0)
        cell = ws4.cell(r, j, t)
        cell.font = st['ink']
        cell.fill = st['fill_teal']
        cell.alignment = st['right']
        cell.number_format = '#,##0.00'
        cell.border = st['thin']
    cell = ws4.cell(r, len(ulkeler) + 2, round(loj, 2))
    cell.font = st['ink']
    cell.fill = st['fill_teal']
    cell.alignment = st['right']
    cell.number_format = '#,##0.00'
    cell.border = st['thin']
    _autosize(ws4, max_w=22)

    # ── Ülke payı ─────────────────────────────────────────────────────────
    ws5 = wb.create_sheet('Ulke_Pay')
    ws5.cell(1, 1, 'Ülke lojistik payı').font = st['title']
    _header_row(ws5, 3, ['Ülke', 'Kod', 'Lojistik EUR', 'Pay %'], st)
    for i, u in enumerate(ulkeler, 4):
        kod = u.get('ulke')
        t = float(ulke_loj.get(kod) or 0)
        pay = (t / loj * 100) if loj else 0
        for c, v in enumerate([u.get('label'), (kod or '').upper(), t, pay], 1):
            cell = ws5.cell(i, c, v)
            cell.font = st['cell']
            cell.border = st['thin']
            if c >= 3:
                cell.alignment = st['right']
                cell.number_format = '#,##0.00' if c == 3 else '0.00'
    _autosize(ws5)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"maliyet_tumu_{start}_{end}.xlsx".replace(':', '-')
    return send_file(
        buf,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=fname,
    )
