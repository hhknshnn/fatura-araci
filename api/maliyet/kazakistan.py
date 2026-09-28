# api/maliyet/kazakistan.py
# Kazakistan: HUB depo Excel (KZT) + nakliye xls + LOGISTICS COSTS Kazakhstan.
# Depolama ayrı satırdır (Giriş+Stok+Çıkış ara toplamı); lojistikte iki kez sayılmaz.

import calendar
import datetime
import logging
import os
import re
import unicodedata

from flask import jsonify, request
from openpyxl import load_workbook

from api.db import get_conn
from api.kur import get_tcmb_kurlar
from api.maliyet.hesap import to_eur
from api.maliyet.ulke_tablo import tablo_kaydet, ulke_ciro_kaydet

_BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
_KZ_DIR = os.path.join(_BASE, 'KAZAKİSTAN MALİYET', '2026')
_LOGISTICS_XLSX = os.path.join(_BASE, 'Logistic Cost', 'LOGISTICS COSTS.xlsx')

_MONTH_TR = {
    'ocak': 1, 'january': 1, 'jan': 1,
    'subat': 2, 'şubat': 2, 'february': 2, 'feb': 2,
    'mart': 3, 'march': 3, 'mar': 3,
    'nisan': 4, 'april': 4, 'apr': 4,
    'mayis': 5, 'mayıs': 5, 'may': 5,
    'haziran': 6, 'june': 6, 'jun': 6,
    'temmuz': 7, 'july': 7, 'jul': 7,
    'agustos': 8, 'ağustos': 8, 'august': 8, 'aug': 8,
}


def _fold(s):
    s = unicodedata.normalize('NFKD', str(s or ''))
    s = ''.join(c for c in s if not unicodedata.combining(c)).lower()
    return s.translate(str.maketrans('çğıöşüı', 'cgiosui'))


def _num(value):
    if value is None or value == '':
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace('\xa0', '').replace(' ', '').replace('₸', '')
    if not s:
        return None
    if ',' in s and '.' in s:
        if s.rfind(',') > s.rfind('.'):
            s = s.replace('.', '').replace(',', '.')
        else:
            s = s.replace(',', '')
    elif ',' in s:
        if re.search(r',\d{1,2}$', s):
            s = s.replace('.', '').replace(',', '.')
        else:
            s = s.replace(',', '')
    try:
        return float(s)
    except ValueError:
        return None


def _kzt_eur(kzt, kurlar):
    e = to_eur(kzt, 'KZT', kurlar)
    return round(e, 2) if e is not None else 0.0


def _folder_ay_no(name):
    n = _fold(name)
    m = re.match(r'^(\d{1,2})\b', n)
    if m:
        v = int(m.group(1))
        if 1 <= v <= 12:
            return v
    for key, mon in _MONTH_TR.items():
        if key in n:
            return mon
    return None


def _sheet_ay_no(name):
    n = _fold(name).replace("'", '').replace(' ', '')
    for key, mon in sorted(_MONTH_TR.items(), key=lambda x: -len(x[0])):
        if key in n:
            return mon
    return None


def _is_depo_name(name):
    n = _fold(name)
    if 'nakliye' in n:
        return False
    return any(k in n for k in ('depo', 'depolama', 'islem', 'işlem'))


def _is_nakliye_name(name):
    return 'nakliye' in _fold(name)


def parse_kz_depo_sheet(ws, ay_no, yil=2026):
    """Depo sheet altındaki INBOUND/OUTBOUND/STORAGE fatura bloğu (KZT)."""
    inbound = outbound = storage_qty = storage_amt = None
    days = 0
    box_vals = []
    period_start = period_end = None
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 60, 60), values_only=True):
        cells = list(row)
        joined = ' '.join(str(c or '') for c in cells).lower()
        if str(cells[0] or '').strip().lower() in ('start', 'end') and len(cells) > 1:
            v = cells[1]
            if isinstance(v, datetime.datetime):
                d = v.date()
            elif isinstance(v, datetime.date):
                d = v
            else:
                d = None
            if d:
                if str(cells[0]).strip().lower() == 'start':
                    period_start = d
                else:
                    period_end = d
        if isinstance(cells[0], (datetime.datetime, datetime.date)):
            box = _num(cells[5]) if len(cells) > 5 else None
            if box and box > 0:
                box_vals.append(box)
                days += 1
        label = str(cells[3] or '').strip().upper() if len(cells) > 3 else ''
        if label == 'INBOUND':
            inbound = {
                'miktar': float(_num(cells[4]) or 0),
                'tutar_kzt': float(_num(cells[6]) or 0),
            }
        elif label == 'OUTBOUND':
            outbound = {
                'miktar': float(_num(cells[4]) or 0),
                'tutar_kzt': float(_num(cells[6]) or 0),
            }
        elif label == 'STORAGE':
            storage_qty = float(_num(cells[4]) or 0)  # koli-gün
            storage_amt = float(_num(cells[6]) or 0)

    if not inbound and not outbound and not storage_amt:
        return None
    inbound = inbound or {'miktar': 0.0, 'tutar_kzt': 0.0}
    outbound = outbound or {'miktar': 0.0, 'tutar_kzt': 0.0}
    last = calendar.monthrange(yil, ay_no)[1]
    gun = days or last
    avg_box = (storage_qty / gun) if storage_qty and gun else (
        (sum(box_vals) / len(box_vals)) if box_vals else 0.0)
    if not period_start:
        period_start = datetime.date(yil, ay_no, 1)
        period_end = datetime.date(yil, ay_no, last)
    return {
        'ay': f'{yil}-{ay_no:02d}',
        'donem_baslangic': period_start.isoformat(),
        'donem_bitis': period_end.isoformat(),
        'inbound': inbound,
        'outbound': outbound,
        'storage': {
            'miktar': round(avg_box, 2),
            'koli_gun': storage_qty or 0.0,
            'tutar_kzt': storage_amt or 0.0,
        },
    }


def parse_kz_depo_xlsx(path, ay_no, yil=2026):
    try:
        wb = load_workbook(path, data_only=True)
    except Exception:
        logging.exception('kz depo open %s', path)
        return None
    try:
        ws = None
        for sn in wb.sheetnames:
            if _sheet_ay_no(sn) == ay_no:
                ws = wb[sn]
                break
        if ws is None:
            return None
        return parse_kz_depo_sheet(ws, ay_no, yil)
    finally:
        wb.close()


def parse_kz_nakliye_xls(path):
    """Nakliye .xls — 'total for exhibiting' / 'Итого на выставление' (KZT, KDV dahil olabilir)."""
    try:
        import xlrd
        wb = xlrd.open_workbook(path)
    except Exception:
        logging.exception('kz nakliye open %s', path)
        return None
    sh = wb.sheet_by_index(0)
    total = None
    pallets = 0.0
    headers = [_fold(sh.cell_value(0, c)) for c in range(min(sh.ncols, 16))]
    qty_col = next((i for i, h in enumerate(headers) if 'quantity' in h or h in ('кол-во', 'kol-vo')), 6)
    amt_col = None
    for i, h in enumerate(headers):
        if 'total amount' in h or 'unit price wth vat' in h:
            amt_col = i
    if amt_col is None:
        amt_col = min(sh.ncols - 1, 10)

    for r in range(sh.nrows):
        vals = [sh.cell_value(r, c) for c in range(sh.ncols)]
        joined = ' '.join(str(v) for v in vals if v not in ('', None)).lower()
        if any(k in joined for k in (
            'итого на выставление', 'total for exhibiting', 'total for exhibiting:',
            'итого на выставление:', 'total amount for issuance',
        )):
            nums = [_num(v) for v in vals]
            nums = [n for n in nums if n and n > 1000]
            if nums:
                total = max(nums)
            continue
        try:
            qty = float(vals[qty_col] or 0) if qty_col < len(vals) else 0
            amt = float(vals[amt_col] or 0) if amt_col < len(vals) else 0
        except (TypeError, ValueError):
            continue
        if qty > 0 and amt >= 20000:
            pallets += qty
    if total is None:
        # Bazı aylarda etiket yok; sondaki büyük tutarı al
        candidates = []
        for r in range(max(0, sh.nrows - 15), sh.nrows):
            for c in range(sh.ncols):
                n = _num(sh.cell_value(r, c))
                if n and n > 1_000_000:
                    candidates.append(n)
        if candidates:
            total = max(candidates)
    if total is None:
        return None
    return {'tutar_kzt': round(total, 2), 'miktar': round(pallets, 2)}


def _logistics_2026_aylar():
    wb = load_workbook(_LOGISTICS_XLSX, data_only=True)
    ws = wb['Kazakhstan']
    cols = [('2026-01', 5, 6), ('2026-02', 7, 8), ('2026-03', 9, 10)]
    row_map = {
        'inbound': 22, 'storage': 23, 'outbound': 24,
        'depolama': 25, 'transport': 26,
    }
    ciro_row = 29
    aylar = []
    for ay, qc, cc in cols:
        a = {'ay': ay, 'ciro_eur': float(ws.cell(ciro_row, qc).value or 0)}
        for kod, rr in row_map.items():
            a[kod] = {
                'miktar': float(ws.cell(rr, qc).value or 0),
                'tutar': float(ws.cell(rr, cc).value or 0),
            }
        # LOGISTICS Total = Depolama + İç Nakliye; depolama skip_total
        aylar.append(a)
    wb.close()
    return aylar


def _ciro_sayfa1():
    out = {}
    try:
        wb = load_workbook(_LOGISTICS_XLSX, data_only=True)
        ws = wb['Sayfa1']

        def _eu(v):
            if isinstance(v, (int, float)):
                return float(v)
            s = str(v or '').replace('\xa0', '').replace('€', '').strip()
            if not s:
                return 0.0
            if ',' in s:
                s = s.replace('.', '').replace(',', '.')
            try:
                return float(re.sub(r'[^\d.-]', '', s) or 0)
            except ValueError:
                return 0.0

        for row in ws.iter_rows(min_row=4, max_row=16, values_only=True):
            ad = _fold(row[2] if len(row) > 2 else '')
            if 'kazakistan' not in ad and 'kazakhstan' not in ad:
                continue
            for ay, col in (('2026-04', 3), ('2026-05', 4), ('2026-06', 5)):
                n = _eu(row[col] if len(row) > col else None)
                if n:
                    out[ay] = n
            break
        wb.close()
    except Exception:
        logging.exception('kz sayfa1 ciro')
    return out


def _scan_month_dirs(root=None):
    root = root or _KZ_DIR
    if not os.path.isdir(root):
        return []
    out = []
    for name in sorted(os.listdir(root)):
        path = os.path.join(root, name)
        if not os.path.isdir(path):
            continue
        ay = _folder_ay_no(name)
        if not ay:
            continue
        depo = nakliye = None
        for f in os.listdir(path):
            full = os.path.join(path, f)
            low = f.lower()
            if low.endswith(('.xlsx', '.xls')):
                if _is_depo_name(f) and not depo:
                    depo = full
                elif _is_nakliye_name(f) and not nakliye:
                    nakliye = full
        out.append({'folder': name, 'ay_no': ay, 'depo': depo, 'nakliye': nakliye})
    return out


def _kalem_id_map(cur):
    cur.execute('SELECT id, kod FROM maliyet_kalemleri WHERE aktif')
    return {r[1]: r[0] for r in cur.fetchall()}


def _upsert_fatura(cur, draft, kalem_ids):
    fno = draft.get('fatura_no')
    if not fno:
        return None, 'fatura_no yok'
    kalemler = []
    for i, row in enumerate(draft.get('kalemler') or []):
        kod = row.get('kalem_kod')
        kid = kalem_ids.get(kod) or kalem_ids.get('handling')
        if not kid:
            continue
        tutar = float(row.get('tutar') or 0)
        kalemler.append({
            'kalem_id': kid, 'tarih': draft.get('donem_bitis'),
            'aciklama': row.get('aciklama') or kod, 'referans': None,
            'miktar': float(row.get('miktar') or 1),
            'birim_fiyat': float(row.get('birim_fiyat') or tutar),
            'tutar': tutar, 'sira': i,
        })
    if not kalemler:
        return None, 'kalem yok'
    tutar = float(draft.get('tutar') or sum(k['tutar'] for k in kalemler))
    cur.execute(
        'SELECT id FROM maliyet_faturalari WHERE ulke=%s AND fatura_no=%s LIMIT 1',
        ('kz', fno),
    )
    row = cur.fetchone()
    alanlar = {
        'ulke': 'kz', 'fatura_no': fno,
        'donem_baslangic': draft.get('donem_baslangic'),
        'donem_bitis': draft.get('donem_bitis'),
        'tutar': round(tutar, 2),
        'para_birimi': 'EUR',
        'fatura_tarihi': draft.get('fatura_tarihi'),
        'notlar': draft.get('notlar') or 'KZ HUB depo/nakliye',
    }
    if row:
        fid = row[0]
        cur.execute('''UPDATE maliyet_faturalari SET
            donem_baslangic=%(donem_baslangic)s, donem_bitis=%(donem_bitis)s,
            tutar=%(tutar)s, para_birimi=%(para_birimi)s, fatura_tarihi=%(fatura_tarihi)s,
            notlar=%(notlar)s WHERE id=%(id)s''', {**alanlar, 'id': fid})
    else:
        cur.execute('''INSERT INTO maliyet_faturalari
            (ulke,fatura_no,donem_baslangic,donem_bitis,tutar,para_birimi,fatura_tarihi,notlar)
            VALUES (%(ulke)s,%(fatura_no)s,%(donem_baslangic)s,%(donem_bitis)s,%(tutar)s,
                    %(para_birimi)s,%(fatura_tarihi)s,%(notlar)s) RETURNING id''', alanlar)
        fid = cur.fetchone()[0]
    cur.execute('DELETE FROM maliyet_fatura_kalemleri WHERE fatura_id=%s', (fid,))
    for k in kalemler:
        cur.execute('''INSERT INTO maliyet_fatura_kalemleri
            (fatura_id,kalem_id,tarih,aciklama,referans,miktar,birim_fiyat,tutar,sira)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)''',
            (fid, k['kalem_id'], k['tarih'], k['aciklama'], k['referans'],
             k['miktar'], k['birim_fiyat'], k['tutar'], k['sira']))
    return fid, None


def _row_from_depo_nak(ay, det, nak, kurlar, ciro):
    in_e = _kzt_eur(det['inbound']['tutar_kzt'], kurlar)
    out_e = _kzt_eur(det['outbound']['tutar_kzt'], kurlar)
    st_e = _kzt_eur(det['storage']['tutar_kzt'], kurlar)
    tr_e = _kzt_eur((nak or {}).get('tutar_kzt') or 0, kurlar)
    depo_tot = round(in_e + st_e + out_e, 2)
    return {
        'ay': ay,
        'ciro_eur': float(ciro or 0),
        'inbound': {'miktar': det['inbound']['miktar'], 'tutar': in_e},
        'storage': {'miktar': det['storage']['miktar'], 'tutar': st_e},
        'outbound': {'miktar': det['outbound']['miktar'], 'tutar': out_e},
        'depolama': {'miktar': det['storage']['miktar'], 'tutar': depo_tot},
        'transport': {
            'miktar': float((nak or {}).get('miktar') or 0),
            'tutar': tr_e,
        },
    }


def seed_kz_from_folder(root=None):
    root = root or _KZ_DIR
    months = _scan_month_dirs(root)
    log_aylar = _logistics_2026_aylar()
    ciro_map = {a['ay']: a['ciro_eur'] for a in log_aylar}
    ciro_map.update(_ciro_sayfa1())
    kurlar = get_tcmb_kurlar()
    if not kurlar.get('KZT'):
        return {'success': False, 'error': 'KZT kuru alınamadı'}

    fatura_ozet = []
    folder_tablo = {}

    conn = get_conn()
    cur = conn.cursor()
    try:
        kalem_ids = _kalem_id_map(cur)
        for item in months:
            yil = 2026
            ay_key = f'{yil}-{item["ay_no"]:02d}'
            det = parse_kz_depo_xlsx(item['depo'], item['ay_no'], yil) if item.get('depo') else None
            nak = parse_kz_nakliye_xls(item['nakliye']) if item.get('nakliye') and item['nakliye'].lower().endswith('.xls') else None
            if not det:
                continue
            row = _row_from_depo_nak(ay_key, det, nak, kurlar, ciro_map.get(ay_key, 0))
            folder_tablo[ay_key] = row
            last = det['donem_bitis']
            kalemler = [
                {'aciklama': 'Inbound', 'kalem_kod': 'box_in',
                 'miktar': row['inbound']['miktar'], 'tutar': row['inbound']['tutar']},
                {'aciklama': 'Storage (avr. box)', 'kalem_kod': 'storage',
                 'miktar': row['storage']['miktar'], 'tutar': row['storage']['tutar']},
                {'aciklama': 'Outbound', 'kalem_kod': 'box_out',
                 'miktar': row['outbound']['miktar'], 'tutar': row['outbound']['tutar']},
            ]
            if row['transport']['tutar'] > 0.005:
                kalemler.append({
                    'aciklama': 'İç nakliye', 'kalem_kod': 'transport',
                    'miktar': row['transport']['miktar'] or 1,
                    'tutar': row['transport']['tutar'],
                })
            for k in kalemler:
                k['birim_fiyat'] = (k['tutar'] / k['miktar']) if k['miktar'] else k['tutar']
            draft = {
                'fatura_no': f'KZ-HUB-{ay_key}',
                'donem_baslangic': det['donem_baslangic'],
                'donem_bitis': last,
                'fatura_tarihi': last,
                'tutar': round(sum(k['tutar'] for k in kalemler), 2),
                'kalemler': kalemler,
                'notlar': 'KZ HUB Excel (KZT→EUR)',
            }
            fid, err = _upsert_fatura(cur, draft, kalem_ids)
            fatura_ozet.append({
                'ay': ay_key, 'fatura_no': draft['fatura_no'], 'id': fid,
                'tutar': draft['tutar'], 'hata': err,
            })
        for a in log_aylar:
            if a['ay'] in folder_tablo:
                continue
            y, m = map(int, a['ay'].split('-'))
            last = calendar.monthrange(y, m)[1]
            kalemler = []
            for kod, src, ad in (
                ('box_in', 'inbound', 'Inbound'),
                ('storage', 'storage', 'Storage (avr. box)'),
                ('box_out', 'outbound', 'Outbound'),
                ('transport', 'transport', 'İç nakliye'),
            ):
                srcd = a.get(src) or {}
                tutar = float(srcd.get('tutar') or 0)
                miktar = float(srcd.get('miktar') or 0)
                if tutar <= 0 and miktar <= 0:
                    continue
                kalemler.append({
                    'aciklama': ad, 'kalem_kod': kod,
                    'miktar': miktar or 1,
                    'birim_fiyat': (tutar / miktar) if miktar else tutar,
                    'tutar': round(tutar, 2),
                })
            if not kalemler:
                continue
            draft = {
                'fatura_no': f'KZ-LOG-{a["ay"]}',
                'donem_baslangic': f'{a["ay"]}-01',
                'donem_bitis': f'{a["ay"]}-{last:02d}',
                'fatura_tarihi': f'{a["ay"]}-{last:02d}',
                'tutar': round(sum(k['tutar'] for k in kalemler), 2),
                'kalemler': kalemler,
                'notlar': 'KZ LOGISTICS COSTS Q1',
            }
            fid, err = _upsert_fatura(cur, draft, kalem_ids)
            fatura_ozet.append({
                'ay': a['ay'], 'fatura_no': draft['fatura_no'], 'id': fid,
                'tutar': draft['tutar'], 'hata': err, 'kaynak': 'logistics',
            })
        conn.commit()
    except Exception:
        conn.rollback()
        logging.exception('kz seed')
        return {'success': False, 'error': 'Kazakistan aktarımı başarısız'}
    finally:
        cur.close()
        conn.close()

    tablo = []
    seen = set()
    for a in log_aylar:
        tablo.append(dict(a))
        seen.add(a['ay'])
    for ay, row in sorted(folder_tablo.items()):
        if ay in seen:
            continue
        tablo.append(row)
        seen.add(ay)
    tablo.sort(key=lambda x: x['ay'])
    tablo_kaydet('kz', tablo)
    ulke_ciro_kaydet('kz', {a['ay']: a['ciro_eur'] for a in tablo if a.get('ciro_eur')})
    return {
        'success': True,
        'klasor': root,
        'ay_sayisi': len(tablo),
        'kzt_eur': kurlar.get('KZT'),
        'faturalar': fatura_ozet,
        'aylar': [
            {
                'ay': a['ay'], 'ciro_eur': a.get('ciro_eur') or 0,
                'depolama': a.get('depolama'),
                'transport': a.get('transport'),
                'lojistik': round(
                    float((a.get('inbound') or {}).get('tutar') or 0)
                    + float((a.get('storage') or {}).get('tutar') or 0)
                    + float((a.get('outbound') or {}).get('tutar') or 0)
                    + float((a.get('transport') or {}).get('tutar') or 0), 2),
            }
            for a in tablo
        ],
        'not': 'Oca–Mar LOGISTICS; Nis–Ağu depo Excel + nakliye xls (KZT→EUR). Depolama satırı ara toplamdır.',
    }


def maliyet_kz_seed_post():
    body = request.get_json(silent=True) or {}
    return jsonify(seed_kz_from_folder(body.get('klasor') or _KZ_DIR))


def maliyet_kz_onizle_get():
    months = _scan_month_dirs()
    kurlar = get_tcmb_kurlar()
    out = []
    for item in months:
        det = parse_kz_depo_xlsx(item['depo'], item['ay_no']) if item.get('depo') else None
        nak = None
        if item.get('nakliye') and item['nakliye'].lower().endswith('.xls'):
            nak = parse_kz_nakliye_xls(item['nakliye'])
        out.append({'folder': item['folder'], 'ay_no': item['ay_no'], 'depo': det, 'nakliye': nak})
    return jsonify({
        'success': True,
        'aylar': out,
        'logistics': _logistics_2026_aylar(),
        'ciro_q2': _ciro_sayfa1(),
        'kzt': kurlar.get('KZT'),
    })
