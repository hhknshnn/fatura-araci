# api/maliyet/makedonya.py
# Makedonya: M&M Nalog Excel + BGLG sayfa-2 (Mart/Nisan) + LOGISTICS ciro.
# BGLG PDF metin katmanı yok; 04.Nisan faturalarının 2. sayfası transkribe edildi.

import calendar
import datetime
import io
import logging
import os
import re
from collections import defaultdict
from unicodedata import normalize

from flask import jsonify, request
from openpyxl import load_workbook

from api.db import get_conn
from api.maliyet.ulke_tablo import tablo_kaydet, ulke_ciro_kaydet

_BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
_MK_DIR = os.path.join(_BASE, 'MAKEDONYA MALİYET', 'MAKEDONYA')
_LOGISTICS_XLSX = os.path.join(_BASE, 'Logistic Cost', 'LOGISTICS COSTS.xlsx')

_MONTH_TR = {
    'ocak': 1, 'subat': 2, 'şubat': 2, 'mart': 3, 'nisan': 4,
    'mayis': 5, 'mayıs': 5, 'haziran': 6, 'temmuz': 7,
    'agustos': 8, 'ağustos': 8, 'eylul': 9, 'eylül': 9,
    'ekim': 10, 'kasim': 11, 'kasım': 11, 'aralik': 12, 'aralık': 12,
}

_SHEET_MONTH = {
    'januar': 1, 'january': 1, 'jan': 1, 'ocak': 1,
    'februar': 2, 'february': 2, 'feb': 2, 'subat': 2, 'şubat': 2,
    'mart': 3, 'march': 3, 'mar': 3,
    'april': 4, 'apr': 4, 'nisan': 4,
    'maj': 5, 'may': 5, 'mayis': 5, 'mayıs': 5,
    'jun': 6, 'june': 6, 'haziran': 6,
    'jul': 7, 'july': 7, 'temmuz': 7,
    'avgust': 8, 'august': 8, 'agustos': 8, 'ağustos': 8,
    'septembar': 9, 'september': 9, 'sep': 9, 'eylul': 9, 'eylül': 9,
    'oktobar': 10, 'october': 10, 'oct': 10, 'ekim': 10,
    'novembar': 11, 'november': 11, 'nov': 11, 'kasim': 11, 'kasım': 11,
    'decembar': 12, 'december': 12, 'dec': 12, 'aralik': 12, 'aralık': 12,
}


def _fold(s):
    import unicodedata
    s = unicodedata.normalize('NFKD', str(s or ''))
    return ''.join(c for c in s if not unicodedata.combining(c)).lower()


def _folder_ay_no(name):
    n = _fold(name)
    m = re.match(r'^(\d{1,2})\b', n)
    if m:
        v = int(m.group(1))
        if 1 <= v <= 12:
            return v
    for key, mon in _MONTH_TR.items():
        if key in n:
            return mon
    return None


def _sheet_ay_yil(name):
    n = _fold(name).replace(' ', '')
    yil = None
    m = re.search(r'(20\d{2})', n)
    if m:
        yil = int(m.group(1))
    elif re.search(r'(?:^|[^0-9])26(?:$|[^0-9])', n) or n.endswith('26'):
        yil = 2026
    elif re.search(r'(?:^|[^0-9])25(?:$|[^0-9])', n):
        yil = 2025
    ay = None
    for key, mon in sorted(_SHEET_MONTH.items(), key=lambda x: -len(x[0])):
        if key in n:
            ay = mon
            break
    return yil, ay


def _map_service(service):
    s = str(service or '').lower()
    if 'customs' in s or 'hs code' in s:
        return 'taxes'
    if 'warehouse' in s or 'additional wh' in s or 'stock' in s:
        return 'storage'
    if 'delivery' in s or 'transport' in s:
        return 'transport'
    if 'workforce' in s or 'worker' in s or 'unloading' in s or 'loading' in s:
        return 'handling'
    return 'handling'


def parse_nalog_sheet(ws, yil, ay_no):
    """Tek Nalog sheet → fatura draft."""
    order_no = None
    period_start = period_end = None
    for row in ws.iter_rows(min_row=1, max_row=12, values_only=True):
        vals = list(row)
        for i, c in enumerate(vals):
            if str(c or '').strip().lower() == 'order number' and i + 1 < len(vals) and vals[i + 1]:
                order_no = str(vals[i + 1]).strip()
            if str(c or '').strip().lower() == 'od' and i + 1 < len(vals) and vals[i + 1]:
                v = vals[i + 1]
                if isinstance(v, datetime.datetime):
                    period_start = v.date()
                elif isinstance(v, datetime.date):
                    period_start = v
            if str(c or '').strip().lower() == 'do' and i + 1 < len(vals) and vals[i + 1]:
                v = vals[i + 1]
                if isinstance(v, datetime.datetime):
                    period_end = v.date()
                elif isinstance(v, datetime.date):
                    period_end = v

    kalemler = []
    header_row = None
    amount_col = 8  # I
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True), 1):
        cells = [str(c or '').strip().lower() for c in row]
        if 'service' in cells and any('amount' in c for c in cells):
            header_row = i
            for j, c in enumerate(row):
                if 'amount' in str(c or '').lower():
                    amount_col = j
                    break
            break
    if not header_row:
        return None

    for row in ws.iter_rows(min_row=header_row + 1, max_row=header_row + 30, values_only=True):
        cells = list(row)
        service = cells[1] if len(cells) > 1 else None
        if not service:
            # Ukupno satırı
            joined = ' '.join(str(c or '') for c in cells).lower()
            if 'ukupno' in joined:
                break
            continue
        svc = str(service).strip()
        if svc.lower().startswith('obra'):
            continue
        qty = cells[7] if len(cells) > 7 else None
        price = cells[6] if len(cells) > 6 else None
        amount = cells[amount_col] if len(cells) > amount_col else None
        try:
            tutar = float(amount or 0)
        except (TypeError, ValueError):
            tutar = 0.0
        try:
            miktar = float(qty or 0) or 1.0
        except (TypeError, ValueError):
            miktar = 1.0
        try:
            birim = float(price or 0) or tutar
        except (TypeError, ValueError):
            birim = tutar
        if tutar <= 0.005:
            continue
        kalemler.append({
            'aciklama': svc[:140],
            'kalem_kod': _map_service(svc),
            'miktar': miktar,
            'birim_fiyat': birim,
            'tutar': round(tutar, 2),
        })

    if not kalemler:
        return None

    if not period_start:
        period_start = datetime.date(yil, ay_no, 1)
    if not period_end:
        last = calendar.monthrange(yil, ay_no)[1]
        period_end = datetime.date(yil, ay_no, last)

    fno = f"MK-MM-{order_no}" if order_no else f"MK-NALOG-{yil}-{ay_no:02d}"
    return {
        'ulke': 'mk',
        'fatura_no': fno,
        'fatura_tarihi': period_end.isoformat(),
        'donem_baslangic': period_start.isoformat(),
        'donem_bitis': period_end.isoformat(),
        'para_birimi': 'EUR',
        'tutar': round(sum(k['tutar'] for k in kalemler), 2),
        'kalemler': kalemler,
        'kaynak': 'nalog',
        'ay': f'{yil}-{ay_no:02d}',
    }


def _scan_nalogs(root=None):
    root = root or _MK_DIR
    drafts = []
    if not os.path.isdir(root):
        return drafts
    seen = set()
    paths = []
    for dirpath, _, files in os.walk(root):
        for f in files:
            if f.lower().endswith(('.xlsx', '.xlsm')) and not f.startswith('~$'):
                if 'm&m nmk' in f.lower() and 'december' in f.lower():
                    continue  # 2025 arşiv özeti
                paths.append(os.path.join(dirpath, f))
    for path in sorted(paths):
        try:
            wb = load_workbook(path, data_only=True)
        except Exception:
            logging.exception('mk nalog open %s', path)
            continue
        folder_ay = _folder_ay_no(os.path.basename(os.path.dirname(path)))
        for sn in wb.sheetnames:
            yil, ay = _sheet_ay_yil(sn)
            if yil != 2026 or not ay:
                continue
            # klasör ayı ile çelişirse sheet ayını esas al
            draft = parse_nalog_sheet(wb[sn], yil, ay)
            if not draft:
                continue
            key = draft['fatura_no']
            if key in seen:
                continue
            seen.add(key)
            draft['_source_file'] = os.path.basename(path)
            draft['_sheet'] = sn
            drafts.append(draft)
        wb.close()
    return drafts


# Taranmış BGLG 2. sayfa (OBRAČUN SKLADIŠNINE). Dönem PDF içinden; klasör adı değil
# (Mart + Nisan faturaları 04.Nisan klasöründe).
_BGLG_PAGE2 = {
    'BGLG-06415/26': {
        'order_no': '2604008535',
        'fatura_tarihi': '2026-05-08',
        'donem_baslangic': '2026-03-01',
        'donem_bitis': '2026-03-31',
        'ay': '2026-03',
        'kalemler': [
            {'aciklama': 'Import customs clearance up to 5 HS codes',
             'kalem_kod': 'taxes', 'miktar': 2.0, 'birim_fiyat': 55.0, 'tutar': 110.0},
            {'aciklama': 'Additional HS code',
             'kalem_kod': 'taxes', 'miktar': 128.0, 'birim_fiyat': 1.5, 'tutar': 192.0},
            {'aciklama': 'Warehouse up to 10 days',
             'kalem_kod': 'storage', 'miktar': 1.0, 'birim_fiyat': 300.0, 'tutar': 300.0},
            {'aciklama': 'Delivery Palma Mall Tetovo Avia (7-16 eur pal), up to 7.0 t',
             'kalem_kod': 'transport', 'miktar': 1.0, 'birim_fiyat': 180.0, 'tutar': 180.0},
            {'aciklama': 'Additional workforce for unloading',
             'kalem_kod': 'handling', 'miktar': 2.0, 'birim_fiyat': 45.0, 'tutar': 90.0},
        ],
    },
    'BGLG-07812/26': {
        'order_no': '2605010699',
        'fatura_tarihi': '2026-06-05',
        'donem_baslangic': '2026-04-01',
        'donem_bitis': '2026-04-30',
        'ay': '2026-04',
        'kalemler': [
            {'aciklama': 'Import customs clearance up to 5 HS codes',
             'kalem_kod': 'taxes', 'miktar': 2.0, 'birim_fiyat': 55.0, 'tutar': 110.0},
            {'aciklama': 'Additional HS code',
             'kalem_kod': 'taxes', 'miktar': 58.0, 'birim_fiyat': 1.5, 'tutar': 87.0},
            {'aciklama': 'Warehouse up to 10 days',
             'kalem_kod': 'storage', 'miktar': 1.0, 'birim_fiyat': 300.0, 'tutar': 300.0},
            {'aciklama': 'Delivery Palma Mall Tetovo Avia (7-16 eur pal), up to 7.0 t',
             'kalem_kod': 'transport', 'miktar': 3.0, 'birim_fiyat': 180.0, 'tutar': 540.0},
            {'aciklama': 'Additional workforce for unloading',
             'kalem_kod': 'handling', 'miktar': 6.0, 'birim_fiyat': 45.0, 'tutar': 270.0},
        ],
    },
}


def _bglg_key(filename):
    n = _fold(filename).replace(' ', '')
    m = re.search(r'bglg[-_/]?(\d+)[-_/](\d{2})', n)
    if m:
        return f'BGLG-{m.group(1)}/{m.group(2)}'
    m = re.search(r'bglg[-_/]?(\d+)', n)
    if m:
        return f'BGLG-{m.group(1)}/26'
    return None


def _draft_from_bglg(fno, meta, path):
    kalemler = [dict(k) for k in (meta.get('kalemler') or []) if float(k.get('tutar') or 0) > 0.005]
    if not kalemler:
        return None
    return {
        'ulke': 'mk',
        'fatura_no': fno,
        'fatura_tarihi': meta.get('fatura_tarihi'),
        'donem_baslangic': meta.get('donem_baslangic'),
        'donem_bitis': meta.get('donem_bitis'),
        'para_birimi': 'EUR',
        'tutar': round(sum(k['tutar'] for k in kalemler), 2),
        'kalemler': kalemler,
        'kaynak': 'bglg',
        'ay': meta.get('ay'),
        '_source_file': os.path.basename(path),
        '_order_no': meta.get('order_no'),
    }


def _scan_bglg(root=None):
    """Klasördeki BGLG PDF'leri, transkribe edilmiş 2. sayfa ile eşle."""
    root = root or _MK_DIR
    drafts = []
    if not os.path.isdir(root):
        return drafts
    seen = set()
    for dirpath, _, files in os.walk(root):
        for f in files:
            if not f.lower().endswith('.pdf'):
                continue
            key = _bglg_key(f)
            meta = _BGLG_PAGE2.get(key)
            if not meta or key in seen:
                continue
            seen.add(key)
            draft = _draft_from_bglg(key, meta, os.path.join(dirpath, f))
            if draft:
                drafts.append(draft)
    drafts.sort(key=lambda d: d.get('ay') or '')
    return drafts


def _mm_eur(s):
    s = str(s or '').strip().replace('\xa0', '').replace(' ', '').replace('€', '')
    if not s:
        return None
    if ',' in s and '.' in s:
        s = s.replace('.', '').replace(',', '.') if s.rfind(',') > s.rfind('.') else s.replace(',', '')
    elif ',' in s:
        s = s.replace('.', '').replace(',', '.')
    try:
        return float(s)
    except ValueError:
        return None


def parse_mm_obracun_text(text, filename=None):
    """OCR/metin: M&M OBRAČUN SKLADIŠNINE kalemleri."""
    t = text or ''
    low = t.lower()
    if 'militzer' not in low and 'skladis' not in low and 'madame coco nmk' not in low:
        return None
    fno = None
    m = re.search(r'BGLG\s*[-–]?\s*(\d+)\s*/\s*(\d{2})', t, re.I)
    if m:
        fno = f'BGLG-{m.group(1)}/{m.group(2)}'
    if not fno:
        fno = _bglg_key(filename or '')
    if fno and fno in _BGLG_PAGE2:
        return _draft_from_bglg(fno, _BGLG_PAGE2[fno], filename or fno)

    d0 = d1 = None
    m = re.search(
        r'od\s+(\d{2}[./]\d{2}[./]20\d{2}).{0,80}do\s+(\d{2}[./]\d{2}[./]20\d{2})',
        t, re.I | re.S,
    )
    if m:
        def _d(s):
            s = s.replace('/', '.')
            try:
                return datetime.datetime.strptime(s, '%d.%m.%Y').date()
            except ValueError:
                return None
        d0, d1 = _d(m.group(1)), _d(m.group(2))

    services = (
        (r'import customs', 'taxes', 'Import customs clearance up to 5 HS codes'),
        (r'additional hs', 'taxes', 'Additional HS code'),
        (r'warehouse up to', 'storage', 'Warehouse up to 10 days'),
        (r'additional wh', 'storage', 'Additional WH'),
        (r'delivery palma', 'transport', 'Delivery Palma Mall Tetovo Avia'),
        (r'additional workforce', 'handling', 'Additional workforce for unloading'),
    )
    kalemler = []
    for pat, kod, ad in services:
        m = re.search(pat, t, re.I)
        if not m:
            continue
        chunk = t[m.start(): m.start() + 280]
        euros = []
        for em in re.finditer(r'([\d.]+,\d{2}|\d+[.,]\d{2})\s*€', chunk):
            v = _mm_eur(em.group(1))
            if v is not None:
                euros.append(v)
        if not euros:
            continue
        tutar = euros[-1]
        if tutar <= 0.005:
            continue
        qty = 1.0
        qm = re.search(r'(\d+)[.,](\d{3})\b', chunk)
        if qm:
            try:
                qty = float(f'{int(qm.group(1))}.{qm.group(2)}')
            except ValueError:
                qty = 1.0
        elif len(euros) >= 2 and euros[0] > 0:
            qty = round(tutar / euros[0], 4) or 1.0
        kalemler.append({
            'aciklama': ad, 'kalem_kod': kod,
            'miktar': qty, 'birim_fiyat': (tutar / qty) if qty else tutar, 'tutar': round(tutar, 2),
        })
    if not kalemler:
        return None
    if not fno:
        fno = f"MK-MM-{(d1 or d0 or datetime.date.today()).isoformat()}"
    if not d0 and d1:
        d0 = d1.replace(day=1)
    if not d1 and d0:
        last = calendar.monthrange(d0.year, d0.month)[1]
        d1 = d0.replace(day=last)
    if not d0:
        return None
    return {
        'ulke': 'mk',
        'fatura_no': fno,
        'fatura_tarihi': d1.isoformat(),
        'donem_baslangic': d0.isoformat(),
        'donem_bitis': d1.isoformat(),
        'para_birimi': 'EUR',
        'tutar': round(sum(k['tutar'] for k in kalemler), 2),
        'kalemler': kalemler,
        'kaynak': 'bglg-ocr',
        'ay': f'{d0.year}-{d0.month:02d}',
        '_source_file': os.path.basename(filename or '') or None,
    }


def parse_bglg_upload(pdf_bytes, filename=None, text=None):
    """Yüklenen BGLG PDF: bilinen transkript veya OCR."""
    key = _bglg_key(filename or '')
    if key and key in _BGLG_PAGE2:
        return _draft_from_bglg(key, _BGLG_PAGE2[key], filename or key)
    return parse_mm_obracun_text(text or '', filename)


def parse_nalog_upload(xlsx_bytes):
    """Yüklenen M&M Nalog Excel → ay taslakları."""
    drafts = []
    try:
        wb = load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    except Exception:
        logging.exception('mk nalog upload')
        return drafts
    try:
        for sn in wb.sheetnames:
            yil, ay = _sheet_ay_yil(sn)
            if not yil or not ay:
                continue
            draft = parse_nalog_sheet(wb[sn], yil, ay)
            if draft:
                drafts.append(draft)
    finally:
        wb.close()
    return drafts


def _logistics_2026_aylar():
    wb = load_workbook(_LOGISTICS_XLSX, data_only=True)
    ws = wb['Macedonia']
    cols = [
        ('2026-01', 5, 6), ('2026-02', 7, 8), ('2026-03', 9, 10),
        ('2026-04', 14, 15), ('2026-05', 16, 17), ('2026-06', 18, 19),
    ]
    row_map = {'inbound': 21, 'storage': 22, 'outbound': 23, 'transport': 24}
    ciro_row = 27
    aylar = []
    for ay, qc, cc in cols:
        a = {'ay': ay, 'ciro_eur': float(ws.cell(ciro_row, qc).value or 0)}
        for kod, rr in row_map.items():
            a[kod] = {
                'miktar': float(ws.cell(rr, qc).value or 0),
                'tutar': float(ws.cell(rr, cc).value or 0),
            }
        aylar.append(a)
    wb.close()
    return aylar


def _draft_from_logistics_month(a):
    ay = a['ay']
    y, m = map(int, ay.split('-'))
    last = calendar.monthrange(y, m)[1]
    map_kod = {
        'inbound': ('storage', 'Warehouse / Inbound (truck)'),  # MK schemada inbound ≈ truck WH
        'storage': ('storage', 'Stock (Truck)'),
        'outbound': ('pallet_out', 'Outbound'),
        'transport': ('transport', 'Transportation (delivery)'),
    }
    # LOGISTICS: Inbound satırı warehouse truck bedeli — tablo inbound kolonuna yazılır
    kalemler = []
    for src, (kod, ad) in [
        ('inbound', ('storage', 'Warehouse up to 10 days')),
        ('storage', ('storage', 'Stock (Truck)')),
        ('outbound', ('pallet_out', 'Outbound')),
        ('transport', ('transport', 'Transportation (delivery)')),
    ]:
        tutar = float((a.get(src) or {}).get('tutar') or 0)
        miktar = float((a.get(src) or {}).get('miktar') or 0)
        if tutar <= 0 and miktar <= 0:
            continue
        # Defter: inbound maliyeti storage kalemi (truck WH); tablo ayrı tutulur
        defter_kod = 'storage' if src == 'inbound' else kod
        kalemler.append({
            'aciklama': ad, 'kalem_kod': defter_kod,
            'miktar': miktar or 1, 'birim_fiyat': (tutar / miktar) if miktar else tutar,
            'tutar': round(tutar, 2),
            '_tablo': src,
        })
    if not kalemler:
        return None
    return {
        'ulke': 'mk',
        'fatura_no': f'MK-LOG-{ay}',
        'fatura_tarihi': f'{ay}-{last:02d}',
        'donem_baslangic': f'{ay}-01',
        'donem_bitis': f'{ay}-{last:02d}',
        'para_birimi': 'EUR',
        'tutar': round(sum(k['tutar'] for k in kalemler), 2),
        'kalemler': kalemler,
        'kaynak': 'logistics-seed',
        'ay': ay,
    }


def _tablo_from_nalog(ay, draft, ciro=0.0):
    bag = defaultdict(lambda: {'miktar': 0.0, 'tutar': 0.0})
    for k in draft.get('kalemler') or []:
        kod = k.get('kalem_kod')
        tutar = float(k.get('tutar') or 0)
        miktar = float(k.get('miktar') or 0)
        if kod in ('taxes',):
            continue  # L/T lojistik dışı
        if kod == 'storage':
            # Warehouse truck → inbound kolonu (LOGISTICS ile uyum)
            bag['inbound']['miktar'] += miktar
            bag['inbound']['tutar'] += tutar
        elif kod == 'pallet_out':
            bag['outbound']['miktar'] += miktar
            bag['outbound']['tutar'] += tutar
        elif kod in ('transport', 'handling'):
            if kod == 'transport':
                bag['transport']['miktar'] += miktar
            bag['transport']['tutar'] += tutar
    return {
        'ay': ay,
        'ciro_eur': float(ciro or 0),
        'inbound': {'miktar': round(bag['inbound']['miktar'], 4), 'tutar': round(bag['inbound']['tutar'], 2)},
        'storage': {'miktar': round(bag['storage']['miktar'], 4), 'tutar': round(bag['storage']['tutar'], 2)},
        'outbound': {'miktar': round(bag['outbound']['miktar'], 4), 'tutar': round(bag['outbound']['tutar'], 2)},
        'transport': {'miktar': round(bag['transport']['miktar'], 4), 'tutar': round(bag['transport']['tutar'], 2)},
    }


def _kalem_id_map(cur):
    cur.execute('SELECT id, kod FROM maliyet_kalemleri WHERE aktif')
    return {r[1]: r[0] for r in cur.fetchall()}


def _upsert_fatura(cur, draft, kalem_ids):
    fno = draft.get('fatura_no')
    if not fno:
        return None, 'fatura_no yok'
    kalemler = []
    for i, row in enumerate(draft.get('kalemler') or []):
        kod = row.get('kalem_kod')
        kid = kalem_ids.get(kod) or kalem_ids.get('handling')
        if not kid:
            continue
        tutar = float(row.get('tutar') or 0)
        kalemler.append({
            'kalem_id': kid, 'tarih': draft.get('donem_bitis'),
            'aciklama': row.get('aciklama') or kod, 'referans': None,
            'miktar': float(row.get('miktar') or 0) or 1,
            'birim_fiyat': float(row.get('birim_fiyat') or tutar),
            'tutar': tutar, 'sira': i,
        })
    if not kalemler:
        return None, 'kalem yok'
    tutar = float(draft.get('tutar') or sum(k['tutar'] for k in kalemler))
    cur.execute(
        'SELECT id FROM maliyet_faturalari WHERE ulke=%s AND fatura_no=%s LIMIT 1',
        ('mk', fno),
    )
    row = cur.fetchone()
    alanlar = {
        'ulke': 'mk', 'fatura_no': fno,
        'donem_baslangic': draft.get('donem_baslangic'),
        'donem_bitis': draft.get('donem_bitis'),
        'tutar': round(tutar, 2), 'para_birimi': 'EUR',
        'fatura_tarihi': draft.get('fatura_tarihi'),
        'notlar': f"MK {draft.get('kaynak') or 'import'}",
    }
    if row:
        fid = row[0]
        cur.execute('''UPDATE maliyet_faturalari SET
            donem_baslangic=%(donem_baslangic)s, donem_bitis=%(donem_bitis)s,
            tutar=%(tutar)s, para_birimi=%(para_birimi)s, fatura_tarihi=%(fatura_tarihi)s,
            notlar=%(notlar)s WHERE id=%(id)s''', {**alanlar, 'id': fid})
    else:
        cur.execute('''INSERT INTO maliyet_faturalari
            (ulke,fatura_no,donem_baslangic,donem_bitis,tutar,para_birimi,fatura_tarihi,notlar)
            VALUES (%(ulke)s,%(fatura_no)s,%(donem_baslangic)s,%(donem_bitis)s,%(tutar)s,
                    %(para_birimi)s,%(fatura_tarihi)s,%(notlar)s) RETURNING id''', alanlar)
        fid = cur.fetchone()[0]
    cur.execute('DELETE FROM maliyet_fatura_kalemleri WHERE fatura_id=%s', (fid,))
    for k in kalemler:
        cur.execute('''INSERT INTO maliyet_fatura_kalemleri
            (fatura_id,kalem_id,tarih,aciklama,referans,miktar,birim_fiyat,tutar,sira)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)''',
            (fid, k['kalem_id'], k['tarih'], k['aciklama'], k['referans'],
             k['miktar'], k['birim_fiyat'], k['tutar'], k['sira']))
    return fid, None


def seed_mk_from_folder(root=None):
    root = root or _MK_DIR
    fatura_ozet = []
    nalog_drafts = _scan_nalogs(root)
    bglg_drafts = _scan_bglg(root)
    log_aylar = _logistics_2026_aylar()
    ciro_map = {a['ay']: a['ciro_eur'] for a in log_aylar}
    nalog_by_ay = {d['ay']: d for d in nalog_drafts}
    bglg_by_ay = {d['ay']: d for d in bglg_drafts}
    # Nalog > BGLG sayfa-2 > LOGISTICS sentetik
    primary_by_ay = dict(bglg_by_ay)
    primary_by_ay.update(nalog_by_ay)

    conn = get_conn()
    cur = conn.cursor()
    try:
        kalem_ids = _kalem_id_map(cur)
        # 1) Nalog faturaları
        for draft in nalog_drafts:
            fid, err = _upsert_fatura(cur, draft, kalem_ids)
            fatura_ozet.append({
                'ay': draft['ay'], 'fatura_no': draft['fatura_no'], 'id': fid,
                'tutar': draft.get('tutar'), 'hata': err, 'kaynak': 'nalog',
                'sheet': draft.get('_sheet'),
            })
        # 2) BGLG (Nalog olmayan aylar — Mart/Nisan)
        for draft in bglg_drafts:
            if draft['ay'] in nalog_by_ay:
                continue
            fid, err = _upsert_fatura(cur, draft, kalem_ids)
            fatura_ozet.append({
                'ay': draft['ay'], 'fatura_no': draft['fatura_no'], 'id': fid,
                'tutar': draft.get('tutar'), 'hata': err, 'kaynak': 'bglg',
                'file': draft.get('_source_file'),
            })
        # 3) LOGISTICS — Nalog/BGLG yoksa sentetik
        for a in log_aylar:
            if a['ay'] in primary_by_ay:
                continue
            loj = sum(float((a.get(k) or {}).get('tutar') or 0) for k in ('inbound', 'storage', 'outbound', 'transport'))
            if loj <= 0.005:
                continue
            draft = _draft_from_logistics_month(a)
            if not draft:
                continue
            fid, err = _upsert_fatura(cur, draft, kalem_ids)
            fatura_ozet.append({
                'ay': a['ay'], 'fatura_no': draft['fatura_no'], 'id': fid,
                'tutar': draft.get('tutar'), 'hata': err, 'kaynak': 'logistics',
            })
        # Eski sentetik MK-LOG, artık gerçek fatura olan aylardan sil
        cur.execute(
            "SELECT id, fatura_no FROM maliyet_faturalari WHERE ulke='mk' AND fatura_no LIKE 'MK-LOG-%'"
        )
        for fid, fno in list(cur.fetchall() or []):
            ay = str(fno or '').replace('MK-LOG-', '')
            if ay in primary_by_ay:
                cur.execute('DELETE FROM maliyet_faturalari WHERE id=%s', (fid,))
        conn.commit()
    except Exception:
        conn.rollback()
        logging.exception('mk seed')
        return {'success': False, 'error': 'Makedonya aktarımı başarısız'}
    finally:
        cur.close()
        conn.close()

    # Tablo: Nalog/BGLG varsa ondan (taxes hariç); yoksa logistics
    tablo = []
    for a in log_aylar:
        ay = a['ay']
        src = primary_by_ay.get(ay)
        if src:
            tablo.append(_tablo_from_nalog(ay, src, ciro_map.get(ay, 0)))
        else:
            row = dict(a)
            row['ciro_eur'] = ciro_map.get(ay, 0)
            tablo.append(row)

    tablo_kaydet('mk', tablo)
    ulke_ciro_kaydet('mk', {a['ay']: a['ciro_eur'] for a in tablo if a.get('ciro_eur')})

    return {
        'success': True,
        'klasor': root,
        'ay_sayisi': len(tablo),
        'faturalar': fatura_ozet,
        'aylar': [
            {
                'ay': a['ay'], 'ciro_eur': a.get('ciro_eur') or 0,
                'lojistik': round(
                    float((a.get('inbound') or {}).get('tutar') or 0)
                    + float((a.get('storage') or {}).get('tutar') or 0)
                    + float((a.get('outbound') or {}).get('tutar') or 0)
                    + float((a.get('transport') or {}).get('tutar') or 0), 2),
            }
            for a in tablo
        ],
        'not': 'Nalog Excel birincil; yoksa BGLG sayfa-2 (Mart/Nisan). Gümrük taxes defterde, L/T dışı.',
    }


def maliyet_mk_seed_post():
    body = request.get_json(silent=True) or {}
    return jsonify(seed_mk_from_folder(body.get('klasor') or _MK_DIR))


def maliyet_mk_onizle_get():
    nalogs = _scan_nalogs()
    bglgs = _scan_bglg()
    return jsonify({
        'success': True,
        'nalog_aylar': [
            {'ay': d['ay'], 'fatura_no': d['fatura_no'], 'tutar': d['tutar'],
             'kalemler': d['kalemler'], 'sheet': d.get('_sheet')}
            for d in nalogs
        ],
        'bglg_aylar': [
            {'ay': d['ay'], 'fatura_no': d['fatura_no'], 'tutar': d['tutar'],
             'kalemler': d['kalemler'], 'file': d.get('_source_file')}
            for d in bglgs
        ],
        'logistics_aylar': [
            {'ay': a['ay'], 'ciro_eur': a['ciro_eur'],
             'lojistik': round(sum(float(a[k]['tutar']) for k in ('inbound', 'storage', 'outbound', 'transport')), 2)}
            for a in _logistics_2026_aylar()
        ],
    })
