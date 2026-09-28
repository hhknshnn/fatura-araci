# api/maliyet/kosova.py
# Kosova: Dardania Logistics fature_shitje PDF + LOGISTICS COSTS Kosovo ciro.

import calendar
import datetime
import io
import logging
import os
import re
from collections import defaultdict

import pdfplumber
from flask import jsonify, request
from openpyxl import load_workbook

from api.db import get_conn
from api.maliyet.ulke_tablo import tablo_kaydet, ulke_ciro_kaydet

_BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
_XK_DIR = os.path.join(_BASE, 'KOSOVA MALİYET', 'KOSOVA')
_LOGISTICS_XLSX = os.path.join(_BASE, 'Logistic Cost', 'LOGISTICS COSTS.xlsx')

_MONTH_TR = {
    'ocak': 1, 'subat': 2, 'şubat': 2, 'mart': 3, 'nisan': 4,
    'mayis': 5, 'mayıs': 5, 'haziran': 6, 'temmuz': 7,
    'agustos': 8, 'ağustos': 8, 'eylul': 9, 'eylül': 9,
    'ekim': 10, 'kasim': 11, 'kasım': 11, 'aralik': 12, 'aralık': 12,
    'kosova': None,
}


def _fold(s):
    import unicodedata
    s = unicodedata.normalize('NFKD', str(s or ''))
    return ''.join(c for c in s if not unicodedata.combining(c)).lower()


def _parse_eur(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace('\xa0', '').replace(' ', '').replace('€', '')
    if not s:
        return None
    if ',' in s and '.' in s:
        if s.rfind(',') > s.rfind('.'):
            s = s.replace('.', '').replace(',', '.')
        else:
            s = s.replace(',', '')
    elif ',' in s:
        if re.search(r',\d{2}$', s):
            s = s.replace('.', '').replace(',', '.')
        else:
            s = s.replace(',', '')
    try:
        return float(s)
    except ValueError:
        return None


def _parse_date(value):
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    s = str(value or '').strip()
    for fmt in ('%d/%m/%Y', '%d.%m.%Y', '%d-%m-%Y', '%Y-%m-%d'):
        try:
            return datetime.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _pdf_text(path_or_bytes):
    if isinstance(path_or_bytes, (bytes, bytearray)):
        opener = lambda: pdfplumber.open(io.BytesIO(path_or_bytes))
    else:
        opener = lambda: pdfplumber.open(path_or_bytes)
    with opener() as pdf:
        return '\n'.join((p.extract_text(x_tolerance=2, y_tolerance=3) or '') for p in pdf.pages)


def _folder_ay_no(name):
    n = _fold(name)
    m = re.match(r'^(\d{1,2})\b', n)
    if m:
        v = int(m.group(1))
        if 1 <= v <= 12:
            return v
    for key, mon in _MONTH_TR.items():
        if mon and key in n:
            return mon
    return None


def _map_kalem(desc, product=''):
    s = f'{desc} {product}'.lower()
    if 'inbound' in s:
        return 'pallet_in'
    if 'outbound' in s or 'outband' in s:
        return 'pallet_out'
    if 'storage' in s or 'strehim' in s or 'magazinim' in s:
        # ürün satırı warehouse ama açıklama inbound/outbound değilse storage
        if 'inbound' not in s and 'outbound' not in s:
            return 'storage'
    if 'loading' in s or 'unloading' in s:
        return 'handling'
    if 'transport' in s or 'mall' in s or 'return' in s or 'prisht' in s or 'prizren' in s or 'ferizaj' in s:
        return 'transport'
    if 'sherbime te transport' in s:
        return 'transport'
    if 'strehim' in s or 'magazinim' in s:
        return 'storage'
    return 'transport'


def parse_dardania_fature_pdf(path_or_bytes):
    """Dardania fature shitje → EUR net (Vlera pa Tvsh)."""
    text = _pdf_text(path_or_bytes)
    m = re.search(r'Fature\s*shitje\s*:\s*(DL/\d+/20\d{2})', text, re.I)
    fno = m.group(1).upper() if m else None
    m = re.search(r'Data\s*:\s*(\d{2}/\d{2}/\d{4})', text, re.I)
    tarih = _parse_date(m.group(1) if m else None)

    def _num(s):
        # 1,239.00 veya 1239.00
        return _parse_eur(s) or 0.0

    raw_lines = [re.sub(r'\s+', ' ', ln).strip() for ln in text.splitlines() if ln.strip()]
    kalemler = []
    i = 0
    # Son sütun VlMeTvsh bazen 1,239.00 (binlik virgül) → [\d.,]+ şart
    line_re = re.compile(
        r'^(\d+)\s+(\d+)\s+(.+?)\s+Cop[eë]\s+'
        r'([\d.,]+)\s+([\d.,]+)\s+([\d.,]+)\s+([\d.,]+)\s+(\d+)\s+([\d.,]+)\s+([\d.,]+)\s+([\d.,]+)$',
        re.I,
    )
    while i < len(raw_lines):
        line = raw_lines[i]
        m = line_re.match(line)
        if not m:
            i += 1
            continue
        product = m.group(3)
        qty = _num(m.group(4))
        unit_net = _num(m.group(7))  # ÇmPaTvsh
        line_gross = _num(m.group(11))  # VlMeTvsh
        desc = ''
        if i + 1 < len(raw_lines) and not re.match(r'^\d+\s+\d+\s+', raw_lines[i + 1]):
            nxt = raw_lines[i + 1]
            if not re.search(r'Vlera\s*(Bruto|pa)|P[eë]r\s*pages|Baza\s*tvsh|IBAN|Faqe', nxt, re.I):
                desc = nxt
                i += 1
        kod = _map_kalem(desc, product)
        # Net = qty * unit_net; fallback gross/1.18
        tutar = round(qty * unit_net, 2) if unit_net else round(line_gross / 1.18, 2)
        if tutar <= 0 and line_gross > 0:
            tutar = round(line_gross / 1.18, 2)
        kalemler.append({
            'aciklama': (desc or product)[:140],
            'kalem_kod': kod,
            'miktar': qty or 1,
            'birim_fiyat': unit_net or (tutar / (qty or 1)),
            'tutar': tutar,
        })
        i += 1

    m = re.search(r'Vlera\s*pa\s*Tvsh\s*:\s*([\d.,]+)', text, re.I)
    net_total = _parse_eur(m.group(1)) if m else None
    computed = round(sum(k['tutar'] for k in kalemler), 2) if kalemler else None
    if net_total is not None and computed and abs(net_total - computed) > 1.0:
        ratio = net_total / computed
        for k in kalemler:
            k['tutar'] = round(k['tutar'] * ratio, 2)
            k['birim_fiyat'] = round(k['tutar'] / (k['miktar'] or 1), 4)
        computed = net_total
    elif net_total is not None and not kalemler:
        computed = net_total
        kalemler = [{
            'aciklama': 'Logistics (net)', 'kalem_kod': 'transport',
            'miktar': 1, 'birim_fiyat': net_total, 'tutar': net_total,
        }]

    if tarih:
        d_bas = datetime.date(tarih.year, tarih.month, 1)
        # fatura tarihi çoğu zaman sonraki ay → dönem klasör ayına seed sırasında yazılır
        d_bit = datetime.date(tarih.year, tarih.month, calendar.monthrange(tarih.year, tarih.month)[1])
    else:
        d_bas = d_bit = None

    return {
        'ulke': 'xk',
        'fatura_no': fno,
        'fatura_tarihi': tarih.isoformat() if tarih else None,
        'donem_baslangic': d_bas.isoformat() if d_bas else None,
        'donem_bitis': d_bit.isoformat() if d_bit else None,
        'para_birimi': 'EUR',
        'tutar': computed if computed is not None else net_total,
        'kalemler': kalemler,
        'kaynak': 'dardania',
    }


def _scan_pdfs(root=None):
    root = root or _XK_DIR
    out = []
    if not os.path.isdir(root):
        return out
    for name in sorted(os.listdir(root)):
        path = os.path.join(root, name)
        if not os.path.isdir(path):
            continue
        ay = _folder_ay_no(name)
        if not ay:
            continue
        for f in sorted(os.listdir(path)):
            if f.lower().endswith('.pdf'):
                out.append({'folder': name, 'ay_no': ay, 'pdf': os.path.join(path, f)})
    return out


def _kalem_id_map(cur):
    cur.execute('SELECT id, kod FROM maliyet_kalemleri WHERE aktif')
    return {r[1]: r[0] for r in cur.fetchall()}


def _upsert_fatura(cur, draft, kalem_ids):
    fno = draft.get('fatura_no')
    if not fno:
        return None, 'fatura_no yok'
    kalemler = []
    for i, row in enumerate(draft.get('kalemler') or []):
        kod = row.get('kalem_kod')
        kid = kalem_ids.get(kod) or kalem_ids.get('handling')
        if not kid:
            continue
        tutar = float(row.get('tutar') or 0)
        kalemler.append({
            'kalem_id': kid, 'tarih': draft.get('donem_bitis'),
            'aciklama': row.get('aciklama') or kod, 'referans': None,
            'miktar': float(row.get('miktar') or 1),
            'birim_fiyat': float(row.get('birim_fiyat') or tutar),
            'tutar': tutar, 'sira': i,
        })
    if not kalemler:
        return None, 'kalem yok'
    tutar = float(draft.get('tutar') or sum(k['tutar'] for k in kalemler))
    cur.execute(
        'SELECT id FROM maliyet_faturalari WHERE ulke=%s AND fatura_no=%s LIMIT 1',
        ('xk', fno),
    )
    row = cur.fetchone()
    alanlar = {
        'ulke': 'xk', 'fatura_no': fno,
        'donem_baslangic': draft.get('donem_baslangic'),
        'donem_bitis': draft.get('donem_bitis'),
        'tutar': round(tutar, 2), 'para_birimi': 'EUR',
        'fatura_tarihi': draft.get('fatura_tarihi'),
        'notlar': f"XK {draft.get('kaynak') or 'import'}",
    }
    if row:
        fid = row[0]
        cur.execute('''UPDATE maliyet_faturalari SET
            donem_baslangic=%(donem_baslangic)s, donem_bitis=%(donem_bitis)s,
            tutar=%(tutar)s, para_birimi=%(para_birimi)s, fatura_tarihi=%(fatura_tarihi)s,
            notlar=%(notlar)s WHERE id=%(id)s''', {**alanlar, 'id': fid})
    else:
        cur.execute('''INSERT INTO maliyet_faturalari
            (ulke,fatura_no,donem_baslangic,donem_bitis,tutar,para_birimi,fatura_tarihi,notlar)
            VALUES (%(ulke)s,%(fatura_no)s,%(donem_baslangic)s,%(donem_bitis)s,%(tutar)s,
                    %(para_birimi)s,%(fatura_tarihi)s,%(notlar)s) RETURNING id''', alanlar)
        fid = cur.fetchone()[0]
    cur.execute('DELETE FROM maliyet_fatura_kalemleri WHERE fatura_id=%s', (fid,))
    for k in kalemler:
        cur.execute('''INSERT INTO maliyet_fatura_kalemleri
            (fatura_id,kalem_id,tarih,aciklama,referans,miktar,birim_fiyat,tutar,sira)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)''',
            (fid, k['kalem_id'], k['tarih'], k['aciklama'], k['referans'],
             k['miktar'], k['birim_fiyat'], k['tutar'], k['sira']))
    return fid, None


def _logistics_ciro_2026():
    """Kosovo Revenue Jan–Jun 2026."""
    wb = load_workbook(_LOGISTICS_XLSX, data_only=True)
    ws = wb['Kosovo']
    cols = [
        ('2026-01', 5), ('2026-02', 7), ('2026-03', 9),
        ('2026-04', 14), ('2026-05', 16), ('2026-06', 18),
    ]
    ciro_row = 38
    out = {}
    for ay, qc in cols:
        out[ay] = float(ws.cell(ciro_row, qc).value or 0)
    wb.close()
    return out


def _tablo_from_drafts(ay, drafts, ciro=0.0):
    bag = defaultdict(lambda: {'miktar': 0.0, 'tutar': 0.0})
    for d in drafts:
        for k in d.get('kalemler') or []:
            kod = k.get('kalem_kod')
            tutar = float(k.get('tutar') or 0)
            miktar = float(k.get('miktar') or 0)
            if kod == 'pallet_in':
                bag['inbound']['miktar'] += miktar
                bag['inbound']['tutar'] += tutar
            elif kod == 'pallet_out':
                bag['outbound']['miktar'] += miktar
                bag['outbound']['tutar'] += tutar
            elif kod == 'storage':
                bag['storage']['miktar'] += miktar
                bag['storage']['tutar'] += tutar
            elif kod in ('transport', 'handling', 'fuel_surcharge'):
                if kod == 'transport':
                    bag['transport']['miktar'] += miktar
                bag['transport']['tutar'] += tutar
    return {
        'ay': ay,
        'ciro_eur': float(ciro or 0),
        'inbound': {'miktar': round(bag['inbound']['miktar'], 4), 'tutar': round(bag['inbound']['tutar'], 2)},
        'storage': {'miktar': round(bag['storage']['miktar'], 4), 'tutar': round(bag['storage']['tutar'], 2)},
        'outbound': {'miktar': round(bag['outbound']['miktar'], 4), 'tutar': round(bag['outbound']['tutar'], 2)},
        'transport': {'miktar': round(bag['transport']['miktar'], 4), 'tutar': round(bag['transport']['tutar'], 2)},
    }


def seed_xk_from_folder(root=None):
    root = root or _XK_DIR
    fatura_ozet = []
    by_month = defaultdict(list)
    scanned = _scan_pdfs(root)
    ciro_map = _logistics_ciro_2026()

    conn = get_conn()
    cur = conn.cursor()
    try:
        kalem_ids = _kalem_id_map(cur)
        for item in scanned:
            draft = parse_dardania_fature_pdf(item['pdf'])
            if not draft.get('fatura_no'):
                fatura_ozet.append({'pdf': item['pdf'], 'hata': 'fatura_no yok'})
                continue
            yil = 2026
            ay_no = item['ay_no']
            last = calendar.monthrange(yil, ay_no)[1]
            draft['donem_baslangic'] = f'{yil}-{ay_no:02d}-01'
            draft['donem_bitis'] = f'{yil}-{ay_no:02d}-{last:02d}'
            ay = f'{yil}-{ay_no:02d}'
            fid, err = _upsert_fatura(cur, draft, kalem_ids)
            fatura_ozet.append({
                'ay': ay, 'fatura_no': draft['fatura_no'], 'id': fid,
                'tutar': draft.get('tutar'), 'hata': err, 'kaynak': 'dardania',
            })
            if not err:
                by_month[ay].append(draft)
        conn.commit()
    except Exception:
        conn.rollback()
        logging.exception('xk seed')
        return {'success': False, 'error': 'Kosova aktarımı başarısız'}
    finally:
        cur.close()
        conn.close()

    tablo = []
    for ay in sorted(by_month):
        tablo.append(_tablo_from_drafts(ay, by_month[ay], ciro_map.get(ay, 0)))
    tablo_kaydet('xk', tablo)
    ulke_ciro_kaydet('xk', {a['ay']: a['ciro_eur'] for a in tablo if a.get('ciro_eur')})

    return {
        'success': True,
        'klasor': root,
        'ay_sayisi': len(tablo),
        'faturalar': fatura_ozet,
        'aylar': [
            {
                'ay': a['ay'], 'ciro_eur': a.get('ciro_eur') or 0,
                'lojistik': round(
                    float((a.get('inbound') or {}).get('tutar') or 0)
                    + float((a.get('storage') or {}).get('tutar') or 0)
                    + float((a.get('outbound') or {}).get('tutar') or 0)
                    + float((a.get('transport') or {}).get('tutar') or 0), 2),
            }
            for a in tablo
        ],
        'not': 'Dardania PDF net (KDV hariç). Ciro LOGISTICS Kosovo Revenue. REPORT xlsx miktar amaçlı, para yok.',
    }


def maliyet_xk_seed_post():
    body = request.get_json(silent=True) or {}
    return jsonify(seed_xk_from_folder(body.get('klasor') or _XK_DIR))


def maliyet_xk_onizle_get():
    out = []
    for item in _scan_pdfs():
        d = parse_dardania_fature_pdf(item['pdf'])
        out.append({
            'folder': item['folder'], 'pdf': os.path.basename(item['pdf']),
            'fatura_no': d.get('fatura_no'), 'tutar': d.get('tutar'),
            'kalem_sayisi': len(d.get('kalemler') or []),
        })
    return jsonify({'success': True, 'pdf_aylar': out, 'ciro': _logistics_ciro_2026()})
