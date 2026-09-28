# Bosna fatura Excel'i (SATR / SP dışa aktarım).
# 1.xls ham döküm, 2.xls sınıflandırılmış hali:
#   güzergâh (BA)…(BA)…  → TRANSPORT
#   gümrük/spedisyon / ZATEZNE KAMATE → TAXES
#   UTOVAR ROBE…         → outbound / inbound / storage kırılımı
#
# Ay anahtarı fatura Datum sütunudur (DPU değil).
# Tutarlar BAM netodur; 1 EUR = 1,95583 BAM (KM sabiti) ile EUR'a çevrilir.
# KDV (PDV) maliyet tutarına dahil edilmez.

import datetime
import json
import os
import re
import unicodedata

from api.maliyet.ulke_tablo import donem_ay_keys

BAM_PER_EUR = 1.95583
KOLI_PER_PALET = 30

_CIRO_PATH = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
    'data', 'bosna_ciro.json',
)


def bosna_ciro_yukle():
    """Ay → ciro EUR sözlüğü. Örn. {'2026-01': 184075.0}."""
    try:
        from api.maliyet.ulke_tablo import ulke_ciro_yukle
        return dict(ulke_ciro_yukle().get('ba') or {})
    except Exception:
        pass
    try:
        with open(_CIRO_PATH, encoding='utf-8') as f:
            raw = json.load(f) or {}
    except (OSError, ValueError, TypeError):
        return {}
    out = {}
    for k, v in raw.items():
        key = str(k or '').strip()
        if not re.fullmatch(r'20\d{2}-\d{2}', key):
            continue
        try:
            out[key] = round(float(v or 0), 2)
        except (TypeError, ValueError):
            continue
    return out


def bosna_ciro_kaydet(guncelleme):
    """Verilen ay→ciro map'ini yazar. Tüm map döner."""
    try:
        from api.maliyet.ulke_tablo import ulke_ciro_kaydet
        return ulke_ciro_kaydet('ba', guncelleme)
    except Exception:
        pass
    cur = bosna_ciro_yukle()
    for k, v in (guncelleme or {}).items():
        key = str(k or '').strip()
        if not re.fullmatch(r'20\d{2}-\d{2}', key):
            continue
        try:
            n = round(float(v or 0), 2)
        except (TypeError, ValueError):
            continue
        if n <= 0:
            cur.pop(key, None)
        else:
            cur[key] = n
    os.makedirs(os.path.dirname(_CIRO_PATH), exist_ok=True)
    with open(_CIRO_PATH, 'w', encoding='utf-8') as f:
        json.dump(dict(sorted(cur.items())), f, ensure_ascii=False, indent=2)
    return cur


def bosna_ciro_uygula(rapor, overwrite=True):
    """Rapor aylarına kayıtlı ciro_eur yazar. overwrite=False ise yalnız boş olanları doldurur."""
    if not isinstance(rapor, dict):
        return rapor
    ciro = bosna_ciro_yukle()
    for a in rapor.get('aylar') or []:
        if not isinstance(a, dict):
            continue
        ay = a.get('ay')
        mevcut = a.get('ciro_eur')
        if overwrite or mevcut in (None, ''):
            a['ciro_eur'] = float(ciro.get(ay) or 0)
        else:
            try:
                a['ciro_eur'] = round(float(mevcut or 0), 2)
            except (TypeError, ValueError):
                a['ciro_eur'] = float(ciro.get(ay) or 0)
    return rapor


def maliyet_bosna_ciro_post():
    """POST /api/maliyet/bosna/ciro — {aylar:{'2026-01':184075,...}} veya {ciro:{...}}."""
    from flask import jsonify, request
    body = request.get_json(silent=True) or {}
    raw = body.get('aylar') if isinstance(body.get('aylar'), dict) else body.get('ciro')
    if not isinstance(raw, dict):
        return jsonify({'success': False, 'error': 'aylar map gerekli'}), 400
    kayit = bosna_ciro_kaydet(raw)
    return jsonify({'success': True, 'ciro': kayit})


def koli_to_palet(koli):
    """Koli / 30 → tam palet. Kesir >= 0.5 ise yukarı, değilse aşağı (kesir atılır)."""
    n = float(koli or 0)
    if n <= 0:
        return 0
    q = n / KOLI_PER_PALET
    whole = int(q)
    return whole + (1 if (q - whole) >= 0.5 else 0)


def bosna_rota_normalize(text):
    """Slobodan tekst / açıklamadan (BA) şehir - (BA) şehir güzergâhı."""
    s = str(text or '')
    s = re.sub(r'^(TRANSPORT|TAXES)\s*[·:\-]\s*', '', s, flags=re.I)
    m = re.search(
        r'\(\s*BA\s*\)\s*[-–]?\s*([^()]+?)\s*[-–]\s*\(\s*BA\s*\)\s*[-–]?\s*([^,+]+)',
        s, re.I,
    )
    if not m:
        return None
    a = re.sub(r'\s+', ' ', m.group(1)).strip(' -')
    b = re.sub(r'\s+', ' ', m.group(2)).strip(' -')
    if not a or not b:
        return None
    return f'(BA) - {a.upper()} - (BA) - {b.upper()}'


def _rota_kisa(rota):
    m = re.search(r'\(BA\)\s*-\s*(.+?)\s*-\s*\(BA\)\s*-\s*(.+)$', str(rota or ''), re.I)
    if not m:
        s = str(rota or '')
        return s[:32] + ('…' if len(s) > 32 else '')
    return f"{m.group(1).strip()} → {m.group(2).strip()}"


_AY_EN = ('Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec')


def _ay_ad_en(ay):
    s = str(ay or '')
    try:
        m = int(s[5:7])
        y = s[2:4]
        return f'{_AY_EN[m - 1]} {y}'
    except Exception:
        return s


def _eur2(n):
    return f"{float(n or 0):,.2f} €".replace(',', 'X').replace('.', ',').replace('X', '.')


def _q(n):
    return f"{int(round(float(n or 0))):,}".replace(',', '.')


def _bosna_rota_ozet(aylar, key='tutar'):
    bag = {}
    for a in aylar or []:
        for rt in a.get('rotalar') or []:
            nm = rt.get('rota')
            if not nm:
                continue
            rec = bag.setdefault(nm, {'rota': nm, 'miktar': 0, 'tutar': 0.0})
            rec['miktar'] += float(rt.get('miktar') or 0)
            rec['tutar'] += float(rt.get('tutar') or 0)
    rows = []
    for rec in bag.values():
        miktar = rec['miktar']
        tutar = rec['tutar']
        birim = (tutar / miktar) if miktar else 0.0
        rows.append({**rec, 'birim': birim})
    if key == 'miktar':
        rows.sort(key=lambda x: (-x['miktar'], -x['tutar'], x['rota']))
    elif key == 'birim':
        rows.sort(key=lambda x: (-x['birim'], -x['tutar'], x['rota']))
    else:
        rows.sort(key=lambda x: (-x['tutar'], -x['miktar'], x['rota']))
    return rows


def _bosna_yorum_satirlari(aylar):
    """Ekrandaki lojistik insight kartlarıyla aynı özet (tax yok)."""
    dolu = [
        a for a in (aylar or [])
        if float(a.get('lojistik') or 0) or float(a.get('genel') or 0)
        or _item(a, 'inbound')[0] or _item(a, 'transport')[0]
    ]
    if not dolu:
        dolu = list(aylar or [])

    def yil_top(key):
        return (
            sum(_item(a, key)[0] for a in (aylar or [])),
            sum(_item(a, key)[1] for a in (aylar or [])),
        )

    yil_loj = sum(float(a.get('lojistik') or 0) for a in (aylar or []))
    list_out = []

    zirve = sorted(dolu, key=lambda a: float(a.get('lojistik') or 0), reverse=True)
    if zirve and float(zirve[0].get('lojistik') or 0) > 0:
        z = zirve[0]
        list_out.append((
            f"{_ay_ad_en(z.get('ay'))} Highest Month",
            f"{_eur2(z.get('lojistik'))} Logistics",
        ))

    kamyon, nak = yil_top('transport')
    loj_pay = (nak / yil_loj * 100) if yil_loj else 0
    list_out.append((
        f"Transport Is {loj_pay:.0f}% Of Logistics",
        f"{_q(kamyon)} Trucks · Avg {_eur2(nak / kamyon if kamyon else 0)} / Trip",
    ))

    en_sefer = (_bosna_rota_ozet(aylar, 'miktar') or [None])[0]
    en_pahali = (_bosna_rota_ozet(aylar, 'tutar') or [None])[0]
    if en_sefer:
        list_out.append((
            _rota_kisa(en_sefer['rota']),
            f"Most Trips · {_q(en_sefer['miktar'])} Trucks · {_eur2(en_sefer['tutar'])}",
        ))
    if en_pahali and (not en_sefer or en_pahali['rota'] != en_sefer['rota']):
        list_out.append((
            _rota_kisa(en_pahali['rota']),
            f"Highest Cost · {_q(en_pahali['miktar'])} Trips · {_eur2(en_pahali['tutar'])}",
        ))
    en_birim = (_bosna_rota_ozet(aylar, 'birim') or [None])[0]
    if en_birim and en_birim.get('birim'):
        list_out.append((
            _rota_kisa(en_birim['rota']),
            f"Highest Unit Cost · {_eur2(en_birim['birim'])} / Trip · {_q(en_birim['miktar'])} Trucks",
        ))

    if len(dolu) >= 2:
        son, once = dolu[-1], dolu[-2]
        fark = float(son.get('lojistik') or 0) - float(once.get('lojistik') or 0)
        base = float(once.get('lojistik') or 0)
        pct = (fark / base * 100) if base else 0
        list_out.append((
            f"{_ay_ad_en(son.get('ay'))} / {_ay_ad_en(once.get('ay'))}",
            f"{'+' if fark >= 0 else ''}{_eur2(fark)} Logistics ({'+' if pct >= 0 else ''}{pct:.1f}%)",
        ))

    in_q, _ = yil_top('inbound')
    out_q, _ = yil_top('outbound')
    if abs(in_q - out_q) > 0.5:
        fark = abs(in_q - out_q)
        list_out.append((
            'Inbound ≠ Outbound',
            'More inbound than outbound' if in_q > out_q else 'More outbound than inbound',
        ))
        list_out.append((
            f"Diff {_q(fark)} Pallets",
            f"Inbound {_q(in_q)} Pallets · Outbound {_q(out_q)} Pallets",
        ))

    koli, st_tutar = yil_top('storage')
    if koli > 0 or st_tutar > 0:
        palet = koli_to_palet(koli)
        list_out.append((
            f"Billed Storage {_eur2(st_tutar)}",
            f"{_q(koli)} cartons ≈ {_q(palet)} pallet-eq (÷30) · not stock on hand",
        ))
    return list_out[:12]

HEADER_ALIASES = {
    'fatura_no': {'broj fakture', 'broj', 'fatura no', 'invoice', 'invoice no'},
    'datum': {'datum', 'date', 'fatura tarihi', 'invoice date'},
    'dpu': {'dpu', 'datum isporuke'},
    'neto': {'neto', 'net', 'iznos'},
    'pdv': {'pdv iznos', 'pdv', 'vat'},
    'bruto': {'bruto', 'gross'},
    'tekst': {'slobodan tekst', 'tekst', 'opis', 'description', 'napomena'},
}

WAREHOUSE_KALEM = {
    'UTOVAR ROBE': ('pallet_out', 'Outbound'),
    'ISTOVAR ROBE': ('pallet_in', 'Inbound'),
    'SKLADISTENJE': ('storage', 'Storage'),
}

WAREHOUSE_RE = re.compile(
    r'(UTOVAR\s+ROBE|ISTOVAR\s+ROBE|SKLADISTENJE)\s*[-:]?\s*'
    r'(\d+(?:[.,]\d+)?)\s*(?:PALETA)?\s*[xX×]\s*'
    r'(\d+(?:[.,]\d+)?)\s*EUR\s*\(\s*(\d+(?:[.,]\d+)?)\s*BAM\s*\)\s*=\s*'
    r'(\d+(?:[.,]\d+)?)\s*BAM',
    re.I,
)

TAX_HINTS = (
    'carinsko', 'spediter', 'spediterka', 'taksa', 'takse',
    'obracun carinskog', 'sanitarne', 'euro inspekt', 'euroinspekta',
    'platnog prometa', 'obrasci', 'zatezne', 'kamate',
)


def _fold(value):
    s = str(value or '').strip().lower().replace('ı', 'i')
    s = s.translate(str.maketrans('çğıöşü', 'cgiosu'))
    s = unicodedata.normalize('NFKD', s)
    s = ''.join(ch for ch in s if not unicodedata.combining(ch))
    s = re.sub(r'[^a-z0-9]+', ' ', s)
    return re.sub(r'\s+', ' ', s).strip()


def _fold_upper(value):
    s = str(value or '').upper()
    s = s.translate(str.maketrans('ŠĆČŽĐİ', 'SCCZDI'))
    s = unicodedata.normalize('NFKD', s)
    s = ''.join(ch for ch in s if not unicodedata.combining(ch))
    return re.sub(r'\s+', ' ', s)


def _num(value):
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    s = str(value).strip().replace(' ', '').replace('\xa0', '')
    if not s:
        return None
    if re.fullmatch(r'\d+,\d+', s):
        s = s.replace(',', '.')
    elif s.count(',') == 1 and s.count('.') >= 1 and s.rfind(',') > s.rfind('.'):
        s = s.replace('.', '').replace(',', '.')
    elif ',' in s and '.' not in s:
        s = s.replace(',', '.')
    try:
        return float(s)
    except ValueError:
        return None


def bam_to_eur(bam):
    if bam is None:
        return None
    return round(float(bam) / BAM_PER_EUR, 2)


def eur_to_bam(eur):
    if eur is None:
        return 0.0
    return round(float(eur) * BAM_PER_EUR, 2)


def satir_bam(row):
    """Excel Neto (BAM). Yoksa kayıtlı EUR × 1,95583."""
    v = (row or {}).get('tutar_bam')
    if v not in (None, ''):
        try:
            return float(v)
        except (TypeError, ValueError):
            pass
    return eur_to_bam((row or {}).get('tutar'))


def ensure_tutar_bam_kolon(cur):
    """BAM kolonunu yoksa ekler. ALTER yetkisi yoksa işlemi bozmaz (savepoint)."""
    cur.execute('''
        SELECT 1 FROM information_schema.columns
        WHERE table_schema = 'public'
          AND table_name = 'maliyet_fatura_kalemleri'
          AND column_name = 'tutar_bam'
    ''')
    if cur.fetchone():
        return
    try:
        cur.execute('SAVEPOINT tutar_bam_kolon')
        cur.execute('''
            ALTER TABLE maliyet_fatura_kalemleri
            ADD COLUMN IF NOT EXISTS tutar_bam NUMERIC(14,2)
        ''')
        cur.execute('RELEASE SAVEPOINT tutar_bam_kolon')
    except Exception:
        try:
            cur.execute('ROLLBACK TO SAVEPOINT tutar_bam_kolon')
        except Exception:
            pass


def _parse_date(value):
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    if isinstance(value, (int, float)):
        try:
            from openpyxl.utils.datetime import from_excel
            dt = from_excel(value)
            return dt.date() if hasattr(dt, 'date') else dt
        except Exception:
            return None
    s = str(value).strip()
    for fmt in ('%Y-%m-%d', '%d.%m.%Y', '%d/%m/%Y', '%d-%m-%Y'):
        try:
            return datetime.datetime.strptime(s[:10], fmt).date()
        except ValueError:
            continue
    return None


def _row_vals(row):
    vals = []
    for v in row:
        if v is None or (isinstance(v, float) and pd.isna(v)):
            vals.append('')
        else:
            vals.append(v)
    return vals


def _header_map(headers):
    mapping = {}
    folded = [_fold(h) for h in headers]
    for alan, adaylar in HEADER_ALIASES.items():
        for i, h in enumerate(folded):
            if h in adaylar and alan not in mapping:
                mapping[alan] = i
    return mapping


def is_bosna_fatura_sheet(df):
    if df is None or df.empty:
        return False
    for raw in df.head(8).itertuples(index=False, name=None):
        mapping = _header_map(_row_vals(raw))
        if {'fatura_no', 'neto', 'tekst'} <= set(mapping):
            return True
    return False


def _kalem_by_kod(kalem_map, kod):
    if not kalem_map:
        return None
    if kod in kalem_map:
        return kalem_map[kod]
    key = _fold(kod)
    if key in kalem_map:
        return kalem_map[key]
    for item in kalem_map.values():
        if isinstance(item, dict) and item.get('kod') == kod:
            return item
    return None


def _satir(ulke, fatura_no, tarih, donem, kalem, aciklama, miktar, birim_fiyat, tutar, kaynak,
           referans=None, fatura_tarihi=None, tutar_bam=None, tarih_kaynak=None):
    etkin = tarih or fatura_tarihi
    eur = round(float(tutar or 0), 2)
    bam = round(float(tutar_bam), 2) if tutar_bam not in (None, '') else eur_to_bam(eur)
    return {
        'ulke': ulke,
        'fatura_no': fatura_no,
        'tarih': etkin.isoformat() if etkin else None,
        'donem': donem,
        'donem_baslangic': (tarih or fatura_tarihi).isoformat() if (tarih or fatura_tarihi) else None,
        'donem_bitis': (tarih or fatura_tarihi).isoformat() if (tarih or fatura_tarihi) else None,
        'fatura_tarihi': (fatura_tarihi or tarih).isoformat() if (fatura_tarihi or tarih) else None,
        'kalem_id': kalem['id'] if kalem else None,
        'kalem_ad': kalem['ad'] if kalem else aciklama,
        'kalem_kod': kalem['kod'] if kalem else None,
        'aciklama': aciklama,
        'referans': referans,
        'miktar': miktar,
        'birim_fiyat': round(float(birim_fiyat or 0), 4),
        'tutar': eur,
        'tutar_bam': bam,
        'para_birimi': 'EUR',
        'kaynak': kaynak,
        'tarih_kaynak': tarih_kaynak,
    }


def _sinif(tekst):
    folded = _fold(tekst)
    upper = _fold_upper(tekst)
    if upper.startswith('UTOVAR ROBE') or 'UTOVAR ROBE' in upper:
        return 'warehouse'
    if folded.startswith('zatezne') or 'zatezne kamate' in folded:
        return 'taxes'
    if folded == 'transport' or folded.startswith('transport '):
        return 'transport'
    if folded == 'taxes' or folded.startswith('taxes '):
        return 'taxes'
    if '(ba)' in folded or re.search(r'\bba\b.+\bba\b', folded):
        return 'transport'
    if any(h in folded for h in TAX_HINTS):
        return 'taxes'
    return 'taxes' if folded else None


def _warehouse_parts(tekst):
    parts = []
    for label, qty, eur, bam_unit, bam in WAREHOUSE_RE.findall(_fold_upper(tekst)):
        key = re.sub(r'\s+', ' ', label.upper().strip())
        if key not in WAREHOUSE_KALEM:
            continue
        kod, ad = WAREHOUSE_KALEM[key]
        parts.append({
            'kod': kod,
            'ad': ad,
            'miktar': _num(qty),
            'birim_eur': _num(eur),
            'birim_bam': _num(bam_unit),
            'tutar_bam': _num(bam),
        })
    return parts


def parse_bosna_fatura_sheet(name, df, kalem_map):
    """Bosna fatura dökümünü fatura satırlarına çevirir. Uygun değilse None."""
    rows = [_row_vals(r) for r in df.itertuples(index=False, name=None)]
    header_idx = None
    mapping = None
    for i, raw in enumerate(rows[:8]):
        mp = _header_map(raw)
        if {'fatura_no', 'neto', 'tekst'} <= set(mp):
            header_idx, mapping = i, mp
            break
    if mapping is None:
        return None

    satirlar, hatalar = [], []
    for i, row in enumerate(rows[header_idx + 1:], header_idx + 2):
        def col(alan):
            idx = mapping.get(alan)
            return row[idx] if idx is not None and idx < len(row) else ''

        fatura_no = str(col('fatura_no') or '').strip()
        if not fatura_no:
            continue
        # Toplam satırı: "103" gibi sayı
        if re.fullmatch(r'\d+(\.0+)?', fatura_no):
            continue
        neto = _num(col('neto'))
        if neto is None:
            hatalar.append(f'{name} satır {i}: neto okunamadı ({fatura_no})')
            continue
        tekst = str(col('tekst') or '').strip()
        fatura_tarihi = _parse_date(col('datum'))
        dpu = _parse_date(col('dpu'))
        if fatura_tarihi:
            tarih, tarih_kaynak = fatura_tarihi, 'datum'
        elif dpu:
            tarih, tarih_kaynak = dpu, 'dpu'
        else:
            tarih, tarih_kaynak = None, None
            hatalar.append(f'{name} satır {i}: Datum ve DPU boş ({fatura_no})')
        donem = tarih.strftime('%Y-%m') if tarih else None
        sinif = _sinif(tekst) or 'taxes'
        kaynak = f'{name}:{i}'
        ortak = dict(fatura_tarihi=fatura_tarihi, tarih_kaynak=tarih_kaynak)

        if sinif == 'warehouse':
            parts = _warehouse_parts(tekst)
            if not parts:
                hatalar.append(f'{name} satır {i}: UTOVAR kırılımı okunamadı ({fatura_no})')
                kalem = _kalem_by_kod(kalem_map, 'storage')
                satirlar.append(_satir(
                    'ba', fatura_no, tarih, donem, kalem,
                    tekst or 'Warehouse', 1, bam_to_eur(neto), bam_to_eur(neto), kaynak,
                    tutar_bam=neto, **ortak,
                ))
                continue
            for part in parts:
                kalem = _kalem_by_kod(kalem_map, part['kod'])
                tutar_eur = bam_to_eur(part['tutar_bam'])
                birim = part['birim_eur'] if part['birim_eur'] is not None else (
                    tutar_eur / part['miktar'] if part['miktar'] else tutar_eur
                )
                aciklama = (
                    f"{part['ad']} · {part['miktar']:g} palet × {part['birim_eur']} EUR "
                    f"({part['tutar_bam']} BAM)"
                )
                satirlar.append(_satir(
                    'ba', fatura_no, tarih, donem, kalem, aciklama,
                    part['miktar'] or 1, birim, tutar_eur, kaynak,
                    tutar_bam=part['tutar_bam'], **ortak,
                ))
            continue

        kod = 'transport' if sinif == 'transport' else 'taxes'
        kalem = _kalem_by_kod(kalem_map, kod)
        etiket = 'TRANSPORT' if sinif == 'transport' else 'TAXES'
        kisa = tekst if len(tekst) <= 180 else tekst[:177] + '…'
        aciklama = etiket if not kisa else f'{etiket} · {kisa}'
        tutar_eur = bam_to_eur(neto)
        satirlar.append(_satir(
            'ba', fatura_no, tarih, donem, kalem, aciklama,
            1, tutar_eur, tutar_eur, kaynak, referans=fatura_no, tutar_bam=neto, **ortak,
        ))

    if not satirlar:
        return None
    return {
        'tip': 'fatura',
        'sheet': name,
        'kaynak': 'bosna',
        'satirlar': satirlar,
        'hatalar': hatalar,
        'dpu_yedek': _dpu_yedek_ozet(satirlar),
    }


def _bosna_ay_bos(donem):
    z = {'miktar': 0.0, 'tutar': 0.0, 'tutar_bam': 0.0}
    return {
        'ay': donem,
        'inbound': dict(z), 'outbound': dict(z), 'storage': dict(z),
        'transport': dict(z), 'taxes': dict(z),
        'rotalar': [],
        'ciro_eur': 0.0,
    }


def bosna_aylik_rapor(satirlar):
    """Fatura Datum ayına göre LOGISTICS COSTS Bosnia kırılımı (EUR, KDV hariç)."""
    kod_map = {
        'pallet_in': 'inbound', 'pallet_out': 'outbound', 'storage': 'storage',
        'transport': 'transport', 'taxes': 'taxes',
    }
    by = {}
    fatura_adet = {}
    rota_seen = {}
    for r in satirlar or []:
        donem = r.get('donem') or str(r.get('tarih') or '')[:7]
        if not donem or len(donem) < 7:
            continue
        key = kod_map.get(r.get('kalem_kod'))
        if not key:
            continue
        bag = by.setdefault(donem, _bosna_ay_bos(donem))
        item = bag[key]
        item['tutar'] += float(r.get('tutar') or 0)
        item['tutar_bam'] += satir_bam(r)
        if key in ('transport', 'taxes'):
            fn = r.get('fatura_no') or r.get('kaynak')
            seen = fatura_adet.setdefault((donem, key), set())
            if fn not in seen:
                seen.add(fn)
                item['miktar'] += 1
        else:
            item['miktar'] += float(r.get('miktar') or 0)
        if key == 'transport':
            rota = bosna_rota_normalize(r.get('aciklama') or r.get('referans') or '')
            if rota:
                rh = bag.setdefault('_rotalar', {})
                rec = rh.setdefault(rota, {'rota': rota, 'miktar': 0, 'tutar': 0.0, 'tutar_bam': 0.0})
                rec['tutar'] += float(r.get('tutar') or 0)
                rec['tutar_bam'] += satir_bam(r)
                seen_r = rota_seen.setdefault((donem, rota), set())
                fn = r.get('fatura_no') or r.get('kaynak')
                if fn not in seen_r:
                    seen_r.add(fn)
                    rec['miktar'] += 1
    aylar = []
    for donem in sorted(by):
        a = by[donem]
        ham = a.pop('_rotalar', {})
        a['rotalar'] = [
            {
                'rota': v['rota'], 'miktar': int(v['miktar']),
                'tutar': round(v['tutar'], 2), 'tutar_bam': round(v['tutar_bam'], 2),
            }
            for v in sorted(ham.values(), key=lambda x: (-x['tutar'], x['rota']))
        ]
        for k in ('inbound', 'outbound', 'storage', 'transport', 'taxes'):
            a[k]['tutar'] = round(a[k]['tutar'], 2)
            a[k]['tutar_bam'] = round(a[k]['tutar_bam'], 2)
            if k in ('transport', 'taxes'):
                a[k]['miktar'] = int(a[k]['miktar'])
            else:
                a[k]['miktar'] = round(a[k]['miktar'], 2)
            if k == 'storage':
                a[k]['miktar_palet'] = koli_to_palet(a[k]['miktar'])
        a['lojistik'] = round(
            a['inbound']['tutar'] + a['outbound']['tutar'] + a['storage']['tutar'] + a['transport']['tutar'], 2)
        a['lojistik_bam'] = round(
            a['inbound']['tutar_bam'] + a['outbound']['tutar_bam'] + a['storage']['tutar_bam'] + a['transport']['tutar_bam'], 2)
        a['genel'] = round(a['lojistik'] + a['taxes']['tutar'], 2)
        a['genel_bam'] = round(a['lojistik_bam'] + a['taxes']['tutar_bam'], 2)
        aylar.append(a)

    def _top(field):
        return {
            'miktar': round(sum(a[field]['miktar'] for a in aylar), 2),
            'tutar': round(sum(a[field]['tutar'] for a in aylar), 2),
            'tutar_bam': round(sum(a[field]['tutar_bam'] for a in aylar), 2),
        }

    toplam = {k: _top(k) for k in ('inbound', 'outbound', 'storage', 'transport', 'taxes')}
    toplam['transport']['miktar'] = int(toplam['transport']['miktar'])
    toplam['taxes']['miktar'] = int(toplam['taxes']['miktar'])
    toplam['storage']['miktar_palet'] = koli_to_palet(toplam['storage']['miktar'])
    toplam['lojistik'] = round(sum(a['lojistik'] for a in aylar), 2)
    toplam['lojistik_bam'] = round(sum(a.get('lojistik_bam') or 0 for a in aylar), 2)
    toplam['genel'] = round(sum(a['genel'] for a in aylar), 2)
    toplam['genel_bam'] = round(sum(a.get('genel_bam') or 0 for a in aylar), 2)
    return {
        'ulke': 'ba',
        'kur': BAM_PER_EUR,
        'aylar': aylar,
        'toplam': toplam,
        'yillar': sorted({a['ay'][:4] for a in aylar}),
        'dpu_yedek': _dpu_yedek_ozet(satirlar),
        'faturalar': bosna_fatura_listesi(satirlar),
    }


def bosna_fatura_listesi(satirlar):
    """SATR/SP satırlarından fatura no · Datum · Neto (fatura başına tek satır)."""
    tip_ad = {
        'warehouse': 'Warehouse',
        'transport': 'Transport',
        'taxes': 'Taxes',
        'other': 'Other',
    }
    bag = {}
    for r in satirlar or []:
        key = _fno_key(r.get('fatura_no'))
        if not key:
            continue
        kod = r.get('kalem_kod') or ''
        if kod in ('pallet_in', 'pallet_out', 'storage'):
            tip_kod = 'warehouse'
        elif kod == 'transport':
            tip_kod = 'transport'
        elif kod == 'taxes':
            tip_kod = 'taxes'
        else:
            tip_kod = 'other'
        rec = bag.get(key)
        if not rec:
            bag[key] = {
                'fatura_no': r.get('fatura_no'),
                'datum': r.get('tarih') or r.get('fatura_tarihi'),
                'tarih_kaynak': r.get('tarih_kaynak') or '',
                'donem': r.get('donem') or str(r.get('tarih') or r.get('fatura_tarihi') or '')[:7],
                'neto_bam': 0.0,
                'neto_eur': 0.0,
                '_tipler': set(),
            }
            rec = bag[key]
        elif not rec.get('datum'):
            rec['datum'] = r.get('tarih') or r.get('fatura_tarihi')
        if r.get('tarih_kaynak') and not rec.get('tarih_kaynak'):
            rec['tarih_kaynak'] = r.get('tarih_kaynak')
        rec['neto_bam'] += satir_bam(r)
        rec['neto_eur'] += float(r.get('tutar') or 0)
        rec['_tipler'].add(tip_kod)
    out = []
    for rec in bag.values():
        tipler = rec.pop('_tipler', set())
        rec['tip'] = ' + '.join(tip_ad[t] for t in ('warehouse', 'transport', 'taxes', 'other') if t in tipler) or 'Other'
        rec['neto_bam'] = round(rec['neto_bam'], 2)
        rec['neto_eur'] = round(rec['neto_eur'], 2)
        out.append(rec)
    out.sort(key=lambda x: (str(x.get('datum') or ''), str(x.get('fatura_no') or '')))
    return out


def _fno_key(value):
    return re.sub(r'\s+', ' ', str(value or '').strip())


def _dpu_yedek_ozet(satirlar):
    """Datum boş olduğu için DPU yazılan faturalar (tekrarsız)."""
    seen = set()
    out = []
    for r in satirlar or []:
        if r.get('tarih_kaynak') != 'dpu':
            continue
        key = _fno_key(r.get('fatura_no'))
        if not key or key in seen:
            continue
        seen.add(key)
        out.append({
            'fatura_no': r.get('fatura_no'),
            'donem': r.get('donem'),
            'tarih': r.get('tarih'),
        })
    return out


def bosna_kayitli_veri():
    """Kayıtlı Bosna faturaları: {fatura_no_key: {fatura_no, tutar, id}} + rapor satırları."""
    from api.db import get_conn
    conn = get_conn()
    cur = conn.cursor()
    try:
        ensure_tutar_bam_kolon(cur)
        conn.commit()
        cur.execute('''
            SELECT id, fatura_no, tutar FROM maliyet_faturalari WHERE ulke = 'ba'
        ''')
        ozet = {}
        for fid, fno, tutar in cur.fetchall():
            ozet[_fno_key(fno)] = {
                'id': fid, 'fatura_no': fno, 'tutar': float(tutar or 0),
            }
        cur.execute('''
            SELECT f.fatura_no, COALESCE(fk.tarih, f.donem_baslangic), k.kod, fk.miktar, fk.tutar, fk.aciklama, fk.tutar_bam
            FROM maliyet_faturalari f
            JOIN maliyet_fatura_kalemleri fk ON fk.fatura_id = f.id
            JOIN maliyet_kalemleri k ON k.id = fk.kalem_id
            WHERE f.ulke = 'ba'
        ''')
        satirlar = []
        for fno, gun, kod, miktar, tutar, aciklama, tutar_bam in cur.fetchall():
            satirlar.append({
                'fatura_no': fno,
                'donem': gun.strftime('%Y-%m') if gun else None,
                'tarih': gun.isoformat() if gun else None,
                'kalem_kod': kod,
                'miktar': float(miktar or 0),
                'tutar': float(tutar or 0),
                'tutar_bam': float(tutar_bam) if tutar_bam is not None else None,
                'aciklama': aciklama or '',
            })
        return ozet, satirlar
    finally:
        cur.close()
        conn.close()


def _satir_grup(satirlar):
    bag = {}
    for r in satirlar or []:
        key = _fno_key(r.get('fatura_no'))
        if key:
            bag.setdefault(key, []).append(r)
    return bag


def _kalem_ozet(rows):
    bag = {}
    tot = 0.0
    donem = None
    for r in rows or []:
        kod = r.get('kalem_kod') or ''
        bag.setdefault(kod, {'miktar': 0.0, 'tutar': 0.0})
        bag[kod]['miktar'] += float(r.get('miktar') or 0)
        bag[kod]['tutar'] += float(r.get('tutar') or 0)
        tot += float(r.get('tutar') or 0)
        donem = donem or r.get('donem')
    return bag, round(tot, 2), donem


def _ozet_degisti(eski, yeni):
    for kod in set(eski) | set(yeni):
        a = eski.get(kod) or {'miktar': 0.0, 'tutar': 0.0}
        b = yeni.get(kod) or {'miktar': 0.0, 'tutar': 0.0}
        if abs(a['tutar'] - b['tutar']) > 0.05 or abs(a['miktar'] - b['miktar']) > 0.05:
            return True
    return False


def bosna_excel_aktarim(excel_satirlar):
    """Excel'i kayıtlı Bosna faturalarıyla karşılaştırır.

    Yeni fatura no → eklenir.
    Aynı no, aynı kırılım → dokunulmaz.
    Aynı no, tutar/miktar/ay değişmiş → güncelleme adayı (onayla yazılır).
    Önizleme raporu: korunan kayıt + yeni + değişenlerin Excel hali.
    """
    kayitli_map, kayitli_satirlar = bosna_kayitli_veri()
    excel_tutar = {}
    for r in excel_satirlar or []:
        key = _fno_key(r.get('fatura_no'))
        if not key:
            continue
        excel_tutar[key] = excel_tutar.get(key, 0) + float(r.get('tutar') or 0)

    kayitli_keys = set(kayitli_map)
    excel_keys = set(excel_tutar)
    yeni = sorted(excel_keys - kayitli_keys)
    ortak = sorted(excel_keys & kayitli_keys)
    sadece_db = sorted(kayitli_keys - excel_keys)

    excel_grup = _satir_grup(excel_satirlar)
    db_grup = _satir_grup(kayitli_satirlar)
    cakisma = []
    for key in ortak:
        e_ozet, e_tot, e_donem = _kalem_ozet(excel_grup.get(key))
        d_ozet, d_tot, d_donem = _kalem_ozet(db_grup.get(key))
        db_header = round(kayitli_map[key]['tutar'], 2)
        if (
            not _ozet_degisti(d_ozet, e_ozet)
            and abs(e_tot - db_header) <= 0.05
            and (e_donem or '') == (d_donem or '')
        ):
            continue
        detay = []
        for kod in sorted(set(e_ozet) | set(d_ozet)):
            a = d_ozet.get(kod) or {'miktar': 0.0, 'tutar': 0.0}
            b = e_ozet.get(kod) or {'miktar': 0.0, 'tutar': 0.0}
            if abs(a['tutar'] - b['tutar']) > 0.05 or abs(a['miktar'] - b['miktar']) > 0.05:
                detay.append({
                    'kod': kod,
                    'kayitli_miktar': round(a['miktar'], 2),
                    'excel_miktar': round(b['miktar'], 2),
                    'kayitli_tutar': round(a['tutar'], 2),
                    'excel_tutar': round(b['tutar'], 2),
                })
        cakisma.append({
            'id': kayitli_map[key]['id'],
            'fatura_no': kayitli_map[key]['fatura_no'],
            'donem': e_donem or d_donem,
            'kayitli_donem': d_donem,
            'excel_donem': e_donem,
            'kayitli_tutar': db_header,
            'excel_tutar': e_tot,
            'fark': round(e_tot - db_header, 2),
            'detay': detay[:8],
        })

    yeni_set = set(yeni)
    degisen_set = {_fno_key(c['fatura_no']) for c in cakisma}
    korunan = [s for s in kayitli_satirlar if _fno_key(s.get('fatura_no')) not in degisen_set]
    excel_guncel = [
        r for r in (excel_satirlar or [])
        if _fno_key(r.get('fatura_no')) in yeni_set or _fno_key(r.get('fatura_no')) in degisen_set
    ]
    rapor = bosna_aylik_rapor(list(korunan) + excel_guncel)
    rapor['kaynak'] = 'birlesik' if kayitli_map else 'excel'
    dpu_yedek = _dpu_yedek_ozet(excel_satirlar)
    rapor['dpu_yedek'] = dpu_yedek

    excel_rapor = bosna_aylik_rapor(excel_satirlar)
    db_rapor = bosna_aylik_rapor(kayitli_satirlar)
    excel_ay = {a['ay']: a for a in excel_rapor.get('aylar') or []}
    db_ay = {a['ay']: a for a in db_rapor.get('aylar') or []}
    ay_fark = []
    alanlar = (
        ('transport', 'Nakliye (kamyon)'),
        ('taxes', 'Taxes'),
        ('inbound', 'Inbound'),
        ('outbound', 'Outbound'),
        ('storage', 'Storage'),
    )
    for ay in sorted(set(excel_ay) | set(db_ay)):
        for kod, ad in alanlar:
            e = (excel_ay.get(ay) or {}).get(kod) or {'miktar': 0, 'tutar': 0}
            d = (db_ay.get(ay) or {}).get(kod) or {'miktar': 0, 'tutar': 0}
            if abs(float(e['miktar']) - float(d['miktar'])) < 0.05 and abs(float(e['tutar']) - float(d['tutar'])) <= 0.05:
                continue
            ay_fark.append({
                'ay': ay, 'alan': ad, 'kod': kod,
                'kayitli_miktar': d['miktar'], 'excel_miktar': e['miktar'],
                'kayitli_tutar': d['tutar'], 'excel_tutar': e['tutar'],
            })

    rota_fark = []
    for ay in sorted(set(excel_ay) | set(db_ay)):
        e_rotalar = {x.get('rota'): x for x in (excel_ay.get(ay) or {}).get('rotalar') or []}
        d_rotalar = {x.get('rota'): x for x in (db_ay.get(ay) or {}).get('rotalar') or []}
        for rota in sorted(set(e_rotalar) | set(d_rotalar)):
            if not rota:
                continue
            e = e_rotalar.get(rota) or {'miktar': 0, 'tutar': 0}
            d = d_rotalar.get(rota) or {'miktar': 0, 'tutar': 0}
            if abs(float(e['miktar']) - float(d['miktar'])) < 0.05 and abs(float(e['tutar']) - float(d['tutar'])) <= 0.05:
                continue
            rota_fark.append({
                'ay': ay, 'rota': rota, 'kod': 'rota',
                'kayitli_miktar': d['miktar'], 'excel_miktar': e['miktar'],
                'kayitli_tutar': d['tutar'], 'excel_tutar': e['tutar'],
            })

    aktarim = {
        'ayni_sayisi': len(ortak) - len(cakisma),
        'kayitli_sayisi': len(ortak),
        'yeni_sayisi': len(yeni),
        'guncelle_sayisi': len(cakisma),
        'excel_fatura': len(excel_keys),
        'db_fatura': len(kayitli_keys),
        'excel_disinda_kayitli': len(sadece_db),
        'yeni_nolar': yeni[:80],
        'kayitli_nolar': [k for k in ortak if k not in degisen_set],
        'cakisma': cakisma[:40],
        'cakisma_sayisi': len(cakisma),
        'ay_fark': ay_fark[:24],
        'rota_fark': rota_fark[:60],
        'dpu_yedek': dpu_yedek,
        'dpu_yedek_sayisi': len(dpu_yedek),
    }
    bosna_ciro_uygula(rapor)
    return aktarim, rapor


def maliyet_bosna_rapor_get():
    """GET /api/maliyet/bosna/rapor?yil=2026 — kayıtlı Bosna faturalarının aylık raporu."""
    from flask import jsonify, request
    from api.db import get_conn

    yil = str(request.args.get('yil') or datetime.date.today().year)
    if not re.fullmatch(r'20\d{2}', yil):
        return jsonify({'success': False, 'error': 'Geçerli bir yıl girin'}), 400
    start = datetime.date(int(yil), 1, 1)
    end = datetime.date(int(yil), 12, 31)
    conn = get_conn()
    cur = conn.cursor()
    try:
        ensure_tutar_bam_kolon(cur)
        conn.commit()
        cur.execute('''
            SELECT f.fatura_no, f.donem_baslangic, k.kod, fk.miktar, fk.tutar, fk.tarih, fk.aciklama, fk.tutar_bam
            FROM maliyet_faturalari f
            JOIN maliyet_fatura_kalemleri fk ON fk.fatura_id = f.id
            JOIN maliyet_kalemleri k ON k.id = fk.kalem_id
            WHERE f.ulke = 'ba'
              AND COALESCE(fk.tarih, f.donem_baslangic) BETWEEN %s AND %s
        ''', (start, end))
        satirlar = []
        for fno, donem_bas, kod, miktar, tutar, tarih, aciklama, tutar_bam in cur.fetchall():
            gun = tarih or donem_bas
            satirlar.append({
                'fatura_no': fno,
                'donem': gun.strftime('%Y-%m') if gun else None,
                'tarih': gun.isoformat() if gun else None,
                'kalem_kod': kod,
                'miktar': float(miktar or 0),
                'tutar': float(tutar or 0),
                'tutar_bam': float(tutar_bam) if tutar_bam is not None else None,
                'aciklama': aciklama or '',
            })
    finally:
        cur.close()
        conn.close()
    rapor = bosna_aylik_rapor(satirlar)
    rapor['yil'] = yil
    rapor['kaynak'] = 'kayit'
    bosna_ciro_uygula(rapor)
    # Ciro girilmiş ama fatura olmayan ayları da ekle
    by = {a.get('ay'): a for a in (rapor.get('aylar') or [])}
    for ay, v in bosna_ciro_yukle().items():
        if not str(ay).startswith(yil) or ay in by:
            continue
        a = _bosna_ay_bos(ay)
        a['ciro_eur'] = float(v or 0)
        a['lojistik'] = 0.0
        a['lojistik_bam'] = 0.0
        a['genel'] = 0.0
        a['genel_bam'] = 0.0
        rapor['aylar'].append(a)
        by[ay] = a
    rapor['aylar'] = sorted(rapor.get('aylar') or [], key=lambda x: x.get('ay') or '')
    return jsonify({'success': True, 'rapor': rapor})


# ── Excel dışa aktarım ────────────────────────────────────────────────────────

_AY_KISA = ('Oca', 'Şub', 'Mar', 'Nis', 'May', 'Haz', 'Tem', 'Ağu', 'Eyl', 'Eki', 'Kas', 'Ara')

_INK = '1A1916'
_TEAL = '0F766E'
_TEAL_BG = 'F0FDFA'
_CREAM = 'FBFAF7'
_LINE = 'EDE8DF'
_RED = 'DC2626'
_BLUE = '2563EB'
_CYAN = '0891B2'
_ROTA_COLORS = (
    '2563EB', '0F766E', 'EA580C', '7C3AED', 'DC2626', '0891B2',
    'CA8A04', 'DB2777', '4F46E5', '16A34A',
)


def _xlsx_fill(hex_color):
    from openpyxl.styles import PatternFill
    return PatternFill('solid', fgColor=hex_color)


def _xlsx_font(**kwargs):
    from openpyxl.styles import Font
    return Font(name='Calibri', **kwargs)


def _xlsx_align(h='center', v='center', wrap=False):
    from openpyxl.styles import Alignment
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


def _xlsx_border():
    from openpyxl.styles import Border, Side
    s = Side(style='thin', color=_LINE)
    return Border(left=s, right=s, top=s, bottom=s)


def _bosna_yil_aylar(rapor, yil):
    z = {'miktar': 0, 'tutar': 0, 'tutar_bam': 0}
    by = {a.get('ay'): a for a in (rapor or {}).get('aylar') or []}
    aylar = []
    for m in range(1, 13):
        key = f'{yil}-{m:02d}'
        a = by.get(key) or {
            'ay': key,
            'inbound': dict(z), 'outbound': dict(z), 'storage': dict(z),
            'transport': dict(z), 'taxes': dict(z),
            'lojistik': 0, 'genel': 0, 'lojistik_bam': 0, 'genel_bam': 0,
        }
        st = a.get('storage') or dict(z)
        if st.get('miktar_palet') is None:
            st = dict(st)
            st['miktar_palet'] = koli_to_palet(st.get('miktar'))
            a = dict(a)
            a['storage'] = st
        aylar.append(a)
    return aylar


def _item(a, key):
    x = (a or {}).get(key) or {}
    return float(x.get('miktar') or 0), float(x.get('tutar') or 0)


def _item_bam(a, key):
    x = (a or {}).get(key) or {}
    if x.get('tutar_bam') not in (None, ''):
        return float(x.get('tutar_bam') or 0)
    return eur_to_bam(x.get('tutar'))


def _ay_bam(a, field):
    v = (a or {}).get(field)
    if v not in (None, ''):
        return float(v)
    if field == 'lojistik_bam':
        return round(sum(_item_bam(a, k) for k in ('inbound', 'outbound', 'storage', 'transport')), 2)
    if field == 'genel_bam':
        return round(_ay_bam(a, 'lojistik_bam') + _item_bam(a, 'taxes'), 2)
    return 0.0


def _style_range(ws, row, col1, col2, font=None, fill=None, align=None):
    b = _xlsx_border()
    for c in range(col1, col2 + 1):
        cell = ws.cell(row=row, column=c)
        cell.border = b
        if font:
            cell.font = font
        if fill:
            cell.fill = fill
        if align:
            cell.alignment = align


def _bosna_tablo_sheet(wb, yil, aylar, donem_etiket=None):
    from openpyxl.utils import get_column_letter
    ws = wb.create_sheet('Bosnia', 0)
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = '1:5'
    ws.freeze_panes = 'B6'

    n = max(len(aylar or []), 1)
    last_col = 1 + n * 3 + 3
    baslik = donem_etiket or yil
    ink = _xlsx_font(bold=True, color='FFFFFF', size=11)
    sub = _xlsx_font(bold=True, color='FFFFFF', size=9)
    lab = _xlsx_font(bold=True, size=10, color=_INK)
    pad = _xlsx_font(size=10, color='6F6B64')
    numf = _xlsx_font(size=10)
    totf = _xlsx_font(bold=True, size=10, color=_TEAL)
    grandf = _xlsx_font(bold=True, size=10, color='FFFFFF')
    title = _xlsx_font(bold=True, size=16, color=_INK)
    kicker = _xlsx_font(bold=True, size=10, color=_TEAL)
    muted = _xlsx_font(size=9, color='6F6B64')
    center = _xlsx_align('center', 'center', True)
    left = _xlsx_align('left', 'center')
    right = _xlsx_align('right', 'center')
    fill_ink = _xlsx_fill(_INK)
    fill_teal = _xlsx_fill(_TEAL)
    fill_cream = _xlsx_fill(_CREAM)
    fill_tot = _xlsx_fill(_TEAL_BG)
    fill_white = _xlsx_fill('FFFFFF')

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.cell(1, 1, 'Bosnia Herzegovina').font = kicker
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws.cell(2, 1, f'Lojistik maliyet · {baslik}').font = title
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_col)
    ws.cell(3, 1, (
        'Cost € = Excel Neto / 1,95583. Cost BAM = Excel Neto (KDV hariç). '
        'Storage miktarı koli; Billed Storage (Pallet-eq) = koli ÷ 30 (0,5 ve üzeri yukarı). '
        'Inbound/outbound faturadaki palet. Taxes ayrı satır. '
        'Transportation altında güzergâh kırılımı.'
    )).font = muted
    ws.row_dimensions[1].height = 16
    ws.row_dimensions[2].height = 24
    ws.row_dimensions[3].height = 28

    ws.cell(4, 1, 'Kalem').font = ink
    ws.cell(4, 1).fill = fill_ink
    ws.cell(4, 1).alignment = center
    ws.merge_cells(start_row=4, start_column=1, end_row=5, end_column=1)
    _style_range(ws, 4, 1, 1, ink, fill_ink, center)
    _style_range(ws, 5, 1, 1, ink, fill_ink, center)

    for i, a in enumerate(aylar):
        c = 2 + i * 3
        m = int(str(a['ay'])[5:7])
        ws.merge_cells(start_row=4, start_column=c, end_row=4, end_column=c + 2)
        cell = ws.cell(4, c, f"{_AY_KISA[m - 1]} {str(yil)[2:]}")
        cell.font = ink
        cell.fill = fill_ink
        cell.alignment = center
        ws.cell(4, c + 1).fill = fill_ink
        ws.cell(4, c + 1).border = _xlsx_border()
        ws.cell(4, c + 2).fill = fill_ink
        ws.cell(4, c + 2).border = _xlsx_border()
        ws.cell(5, c, 'Miktar').font = sub
        ws.cell(5, c + 1, 'Cost €').font = sub
        ws.cell(5, c + 2, 'Cost BAM').font = sub
        _style_range(ws, 4, c, c + 2, ink, fill_ink, center)
        _style_range(ws, 5, c, c + 2, sub, fill_teal, center)

    tc = 2 + n * 3
    tot_label = f'{baslik} toplam'
    ws.merge_cells(start_row=4, start_column=tc, end_row=4, end_column=tc + 2)
    ws.cell(4, tc, tot_label).font = ink
    _style_range(ws, 4, tc, tc + 2, ink, fill_ink, center)
    ws.cell(5, tc, 'Miktar').font = sub
    ws.cell(5, tc + 1, 'Cost €').font = sub
    ws.cell(5, tc + 2, 'Cost BAM').font = sub
    _style_range(ws, 5, tc, tc + 2, sub, fill_teal, center)
    ws.row_dimensions[4].height = 22
    ws.row_dimensions[5].height = 18

    def _write_triple(row, c, q, eur, bam, font, fill, qty_fmt='#,##0', skip_qty=False, skip_cost=False):
        qcell = ws.cell(row, c, None if skip_qty else (q or 0))
        ecell = ws.cell(row, c + 1, None if skip_cost else (eur or 0))
        bcell = ws.cell(row, c + 2, None if skip_cost else (bam or 0))
        for cell in (qcell, ecell, bcell):
            cell.font = font
            cell.fill = fill
            cell.alignment = right
            cell.border = _xlsx_border()
        if not skip_qty:
            qcell.number_format = qty_fmt
        if not skip_cost:
            ecell.number_format = '#,##0.00'
            bcell.number_format = '#,##0.00'
        return (0 if skip_qty else (q or 0), 0 if skip_cost else (eur or 0), 0 if skip_cost else (bam or 0))

    def qty_cost(row, key, font, fill, qty_fmt='#,##0', cost_only=False, palet=False):
        q_sum, t_sum, b_sum = 0.0, 0.0, 0.0
        for i, a in enumerate(aylar):
            c = 2 + i * 3
            if palet:
                q = float((a.get('storage') or {}).get('miktar_palet') or 0)
                qq, _, _ = _write_triple(row, c, q, None, None, font, fill, qty_fmt, skip_cost=True)
                q_sum += qq
            else:
                q, t = _item(a, key) if key else (0, float(a.get('lojistik') or 0))
                b = _item_bam(a, key) if key else _ay_bam(a, 'lojistik_bam')
                qq, tt, bb = _write_triple(row, c, q, t, b, font, fill, qty_fmt, skip_qty=cost_only)
                q_sum += qq
                t_sum += tt
                b_sum += bb
        _write_triple(row, tc, q_sum, t_sum, b_sum, font, fill, qty_fmt, skip_qty=cost_only, skip_cost=palet)
        return q_sum, t_sum, b_sum

    def label_row(row, text, font, fill, indent=False):
        cell = ws.cell(row, 1, text)
        cell.font = font
        cell.fill = fill
        cell.alignment = left
        cell.border = _xlsx_border()
        if indent:
            cell.alignment = _xlsx_align('left', 'center')

    # row 6 warehouse banner
    r = 6
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=last_col)
    ws.cell(r, 1, 'Warehouse').font = lab
    _style_range(ws, r, 1, last_col, lab, fill_cream, left)

    r = 7
    label_row(r, '    Inbound', pad, fill_white, True)
    qty_cost(r, 'inbound', numf, fill_white)

    r = 8
    label_row(r, '    Storage (koli)', pad, fill_white, True)
    qty_cost(r, 'storage', numf, fill_white)

    r = 9
    label_row(r, '    Billed Storage (Pallet-eq)', pad, fill_white, True)
    qty_cost(r, 'storage', numf, fill_white, palet=True)

    r = 10
    label_row(r, '    Outbound', pad, fill_white, True)
    qty_cost(r, 'outbound', numf, fill_white)

    r = 11
    label_row(r, 'Transportation (Truck)', lab, fill_cream)
    qty_cost(r, 'transport', numf, fill_cream)

    rota_adlari = []
    seen_rota = set()
    for a in aylar:
        for rt in a.get('rotalar') or []:
            nm = rt.get('rota')
            if nm and nm not in seen_rota:
                seen_rota.add(nm)
                rota_adlari.append(nm)
    rota_adlari.sort()
    fill_rota = _xlsx_fill('F8FAFC')
    for rota in rota_adlari:
        r += 1
        label_row(r, '    ' + rota, pad, fill_rota, True)
        q_sum, t_sum, b_sum = 0.0, 0.0, 0.0
        for i, a in enumerate(aylar):
            found = next((x for x in (a.get('rotalar') or []) if x.get('rota') == rota), None)
            q = float((found or {}).get('miktar') or 0)
            t = float((found or {}).get('tutar') or 0)
            b = float((found or {}).get('tutar_bam') or 0) or eur_to_bam(t)
            q_sum += q
            t_sum += t
            b_sum += b
            _write_triple(r, 2 + i * 3, q, t, b, pad, fill_rota)
        _write_triple(r, tc, q_sum, t_sum, b_sum, pad, fill_rota)

    r += 1
    label_row(r, 'Total Cost', totf, fill_tot)
    t_sum, b_sum = 0.0, 0.0
    for i, a in enumerate(aylar):
        t = float(a.get('lojistik') or 0)
        b = _ay_bam(a, 'lojistik_bam')
        t_sum += t
        b_sum += b
        _write_triple(r, 2 + i * 3, None, t, b, totf, fill_tot, skip_qty=True)
    _write_triple(r, tc, None, t_sum, b_sum, totf, fill_tot, skip_qty=True)

    r += 1
    label_row(r, 'Taxes', lab, fill_white)
    qty_cost(r, 'taxes', numf, fill_white)

    r += 1
    label_row(r, 'Genel toplam', grandf, fill_ink)
    t_sum, b_sum = 0.0, 0.0
    for i, a in enumerate(aylar):
        t = float(a.get('genel') or 0)
        b = _ay_bam(a, 'genel_bam')
        t_sum += t
        b_sum += b
        _write_triple(r, 2 + i * 3, None, t, b, grandf, fill_ink, skip_qty=True)
    _write_triple(r, tc, None, t_sum, b_sum, grandf, fill_ink, skip_qty=True)

    # Turnover (Ciro) + Logistics / Turnover — EUR only under Cost € column
    fill_ciro = _xlsx_fill('FFFBEB')
    fill_oran = _xlsx_fill('F8FAFC')
    cirof = _xlsx_font(italic=True, size=10, color=_INK)
    oranf = _xlsx_font(bold=True, size=10, color=_INK)
    loj_sum = sum(float(a.get('lojistik') or 0) for a in aylar)
    ciro_sum = sum(float(a.get('ciro_eur') or 0) for a in aylar)

    r += 1
    label_row(r, 'Turnover (Ciro)', cirof, fill_ciro)
    for i, a in enumerate(aylar):
        c = 2 + i * 3
        _write_triple(r, c, None, None, None, cirof, fill_ciro, skip_qty=True, skip_cost=True)
        v = float(a.get('ciro_eur') or 0)
        ecell = ws.cell(r, c + 1, v if v else None)
        ecell.font = cirof
        ecell.fill = fill_ciro
        ecell.alignment = right
        ecell.border = _xlsx_border()
        if v:
            ecell.number_format = '#,##0.00'
        bcell = ws.cell(r, c + 2, None)
        bcell.fill = fill_ciro
        bcell.border = _xlsx_border()
    _write_triple(r, tc, None, None, None, cirof, fill_ciro, skip_qty=True, skip_cost=True)
    ecell = ws.cell(r, tc + 1, ciro_sum if ciro_sum else None)
    ecell.font = _xlsx_font(italic=True, bold=True, size=10, color=_INK)
    ecell.fill = fill_ciro
    ecell.alignment = right
    ecell.border = _xlsx_border()
    if ciro_sum:
        ecell.number_format = '#,##0.00'
    ws.cell(r, tc + 2).fill = fill_ciro
    ws.cell(r, tc + 2).border = _xlsx_border()

    def _oran_txt(loj, ciro):
        if not ciro or ciro <= 0:
            return None
        pct = loj / ciro * 100.0
        return f'{pct:.2f}%'.replace('.', ',')

    r += 1
    label_row(r, 'Logistics / Turnover', oranf, fill_oran)
    for i, a in enumerate(aylar):
        c = 2 + i * 3
        _write_triple(r, c, None, None, None, oranf, fill_oran, skip_qty=True, skip_cost=True)
        txt = _oran_txt(float(a.get('lojistik') or 0), float(a.get('ciro_eur') or 0))
        ecell = ws.cell(r, c + 1, txt)
        ecell.font = oranf
        ecell.fill = fill_oran
        ecell.alignment = right
        ecell.border = _xlsx_border()
        bcell = ws.cell(r, c + 2, None)
        bcell.fill = fill_oran
        bcell.border = _xlsx_border()
    _write_triple(r, tc, None, None, None, oranf, fill_oran, skip_qty=True, skip_cost=True)
    ecell = ws.cell(r, tc + 1, _oran_txt(loj_sum, ciro_sum))
    ecell.font = oranf
    ecell.fill = fill_oran
    ecell.alignment = right
    ecell.border = _xlsx_border()
    ws.cell(r, tc + 2).fill = fill_oran
    ws.cell(r, tc + 2).border = _xlsx_border()

    yorumlar = _bosna_yorum_satirlari(aylar)
    if yorumlar:
        r += 2
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=min(6, last_col))
        ws.cell(r, 1, 'Logistics Insights').font = _xlsx_font(bold=True, size=12, color=_TEAL)
        r += 1
        ws.cell(r, 1, 'Title').font = ink
        ws.cell(r, 1).fill = fill_ink
        ws.cell(r, 2, 'Detail').font = ink
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=min(6, last_col))
        for c in range(1, min(6, last_col) + 1):
            cell = ws.cell(r, c)
            cell.fill = fill_ink
            cell.font = ink
            cell.border = _xlsx_border()
        for baslik, detay in yorumlar:
            r += 1
            ws.cell(r, 1, baslik).font = lab
            ws.cell(r, 1).alignment = left
            ws.cell(r, 1).border = _xlsx_border()
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=min(6, last_col))
            ws.cell(r, 2, detay).font = muted
            ws.cell(r, 2).alignment = left
            for c in range(2, min(6, last_col) + 1):
                ws.cell(r, c).border = _xlsx_border()
                ws.cell(r, c).fill = fill_white

    r = _bosna_ekran_grafiklerini_ekle(ws, r + 2, aylar)

    ws.column_dimensions['A'].width = 48
    for col in range(2, last_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = 11
    ws.auto_filter.ref = None
    return ws


def _bosna_ekran_grafiklerini_ekle(ws, start_row, aylar):
    """Ekrandaki iki grafiği yan yana ekler. Kaynak veri veryHidden ChartData sayfasında."""
    from openpyxl.chart import BarChart, LineChart, Reference

    if not aylar:
        return start_row

    title_f = _xlsx_font(bold=True, size=12, color=_TEAL)
    wb = ws.parent

    # Kaynak veri: gizli sayfa (sütun gizleyince Excel grafiği boş çizer)
    data_name = 'ChartData'
    if data_name in wb.sheetnames:
        data_ws = wb[data_name]
        wb.remove(data_ws)
    data_ws = wb.create_sheet(data_name)
    data_ws.sheet_state = 'veryHidden'
    data_ws.cell(1, 1, 'Month')
    data_ws.cell(1, 2, 'Trucks')
    data_ws.cell(1, 3, 'Transport EUR')
    data_ws.cell(1, 4, 'Billed Storage (Pallet-eq)')
    for i, a in enumerate(aylar, start=2):
        tr_q, tr_t = _item(a, 'transport')
        st_q, _st_t = _item(a, 'storage')
        palet = float((a.get('storage') or {}).get('miktar_palet') or koli_to_palet(st_q))
        data_ws.cell(i, 1, _ay_ad_en(a.get('ay')))
        data_ws.cell(i, 2, tr_q)
        data_ws.cell(i, 3, tr_t)
        data_ws.cell(i, 4, palet)
    last_row = 1 + len(aylar)
    cats = Reference(data_ws, min_col=1, min_row=2, max_row=last_row)

    ws.cell(start_row, 1, 'Charts').font = title_f
    chart_row = start_row + 1

    bar = BarChart()
    bar.type = 'col'
    bar.title = 'Transport: Trucks × Cost'
    bar.y_axis.title = 'Qty'
    bar.y_axis.majorGridlines = None
    trucks_data = Reference(data_ws, min_col=2, min_row=1, max_row=last_row)
    bar.add_data(trucks_data, titles_from_data=True)
    bar.set_categories(cats)
    bar.shape = 4
    bar.legend.position = 'b'

    line = LineChart()
    line.y_axis.axId = 200
    line.y_axis.title = 'EUR'
    cost_data = Reference(data_ws, min_col=3, min_row=1, max_row=last_row)
    line.add_data(cost_data, titles_from_data=True)
    line.set_categories(cats)
    line.y_axis.crosses = 'max'
    bar.y_axis.crosses = 'min'
    bar += line
    bar.width = 10
    bar.height = 8
    ws.add_chart(bar, f'A{chart_row}')

    st = BarChart()
    st.type = 'col'
    st.title = 'Billed Storage (Pallet-eq)'
    st.y_axis.title = 'Pallet-eq'
    pal_data = Reference(data_ws, min_col=4, min_row=1, max_row=last_row)
    st.add_data(pal_data, titles_from_data=True)
    st.set_categories(cats)
    st.legend.position = 'b'
    st.width = 10
    st.height = 8
    from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
    from openpyxl.drawing.xdr import XDRPositiveSize2D
    from openpyxl.utils.units import pixels_to_EMU, cm_to_EMU
    # G sütunundan 250px sola
    st.anchor = OneCellAnchor(
        _from=AnchorMarker(
            col=6,
            colOff=-pixels_to_EMU(250),
            row=chart_row - 1,
            rowOff=0,
        ),
        ext=XDRPositiveSize2D(cm_to_EMU(st.width), cm_to_EMU(st.height)),
    )
    ws.add_chart(st)

    return chart_row + 18


def _bosna_grafik_sheet(wb, yil, aylar):
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.chart.series import DataPoint

    ws = wb.create_sheet('Grafikler')
    ws.sheet_view.showGridLines = False
    title = _xlsx_font(bold=True, size=14, color=_INK)
    hdr = _xlsx_font(bold=True, color='FFFFFF', size=10)
    numf = _xlsx_font(size=10)
    muted = _xlsx_font(size=9, color='6F6B64')
    fill_ink = _xlsx_fill(_INK)
    right = _xlsx_align('right', 'center')
    center = _xlsx_align('center', 'center')

    dolu = [a for a in aylar if (
        float(a.get('genel') or 0) or float(a.get('lojistik') or 0)
        or _item(a, 'inbound')[0] or _item(a, 'transport')[0]
    )]
    if not dolu:
        dolu = aylar

    ws.cell(1, 1, f'Analiz verisi · {yil}').font = title
    ws.merge_cells('A1:H1')
    ws.cell(2, 1, 'Üstte veri tabloları, altta grafikler. Grafikler hücrelerin üstüne binmez.').font = muted
    ws.merge_cells('A2:H2')

    headers = ['Ay', 'Depo €', 'Nakliye €', 'Taxes €', 'Kamyon', 'Palet', 'Koli', 'Genel €']
    for i, h in enumerate(headers, 1):
        cell = ws.cell(4, i, h)
        cell.font = hdr
        cell.fill = fill_ink
        cell.alignment = center
        cell.border = _xlsx_border()

    depo_y = nak_y = tax_y = 0.0
    for i, a in enumerate(dolu):
        row = 5 + i
        m = int(str(a['ay'])[5:7])
        in_q, in_t = _item(a, 'inbound')
        out_q, out_t = _item(a, 'outbound')
        st_q, st_t = _item(a, 'storage')
        tr_q, tr_t = _item(a, 'transport')
        tx_q, tx_t = _item(a, 'taxes')
        depo = in_t + out_t + st_t
        depo_y += depo
        nak_y += tr_t
        tax_y += tx_t
        palet = float((a.get('storage') or {}).get('miktar_palet') or koli_to_palet(st_q))
        vals = [
            f"{_AY_KISA[m - 1]} {str(yil)[2:]}",
            depo, tr_t, tx_t, tr_q, palet, st_q, float(a.get('genel') or 0),
        ]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row, c, v)
            cell.font = numf
            cell.alignment = right if c > 1 else _xlsx_align('left', 'center')
            cell.border = _xlsx_border()
            if c > 1 and c != 5 and c != 6 and c != 7:
                cell.number_format = '#,##0.00'
            elif c in (5, 6, 7):
                cell.number_format = '#,##0'
    n = len(dolu)
    last = 4 + n

    # pie source
    ws.cell(4, 10, 'Kalem').font = hdr
    ws.cell(4, 11, 'Tutar €').font = hdr
    _style_range(ws, 4, 10, 11, hdr, fill_ink, center)
    pie_rows = [('Depo', depo_y), ('Nakliye', nak_y), ('Taxes', tax_y)]
    for i, (ad, tut) in enumerate(pie_rows):
        ws.cell(5 + i, 10, ad).font = numf
        ws.cell(5 + i, 10).border = _xlsx_border()
        cell = ws.cell(5 + i, 11, tut)
        cell.font = numf
        cell.number_format = '#,##0.00'
        cell.border = _xlsx_border()

    rota_tot = {}
    rota_qty = {}
    rota_adlari = []
    for a in dolu:
        for rt in a.get('rotalar') or []:
            nm = rt.get('rota')
            if not nm:
                continue
            rota_tot[nm] = rota_tot.get(nm, 0.0) + float(rt.get('tutar') or 0)
            rota_qty[nm] = rota_qty.get(nm, 0) + int(rt.get('miktar') or 0)
            if nm not in rota_adlari:
                rota_adlari.append(nm)
    rota_adlari.sort(key=lambda r: (-rota_tot.get(r, 0), r))
    n_rota = len(rota_adlari)
    p0 = last + 2
    b0 = p0 + n_rota + 3
    table_end = (b0 + 1 + n_rota) if n_rota else last
    chart0 = table_end + 3
    chart_gap = 20

    # stacked column — lojistik (taxes yok)
    stack_loj = BarChart()
    stack_loj.type = 'col'
    stack_loj.grouping = 'stacked'
    stack_loj.overlap = 100
    stack_loj.title = 'Aylık kırılım · lojistik (Depo / Nakliye)'
    stack_loj.y_axis.title = 'EUR'
    stack_loj.y_axis.numFmt = '#,##0'
    data_loj = Reference(ws, min_col=2, min_row=4, max_col=3, max_row=last)
    cats = Reference(ws, min_col=1, min_row=5, max_row=last)
    stack_loj.add_data(data_loj, titles_from_data=True)
    stack_loj.set_categories(cats)
    stack_loj.shape = 4
    stack_loj.width = 15
    stack_loj.height = 9
    for i, ser in enumerate(stack_loj.series):
        ser.graphicalProperties.solidFill = (_TEAL, _BLUE)[i]
    ws.add_chart(stack_loj, 'A' + str(chart0))

    pie_loj = PieChart()
    pie_loj.title = 'Yıl payı · lojistik'
    labels_loj = Reference(ws, min_col=10, min_row=5, max_row=6)
    pdata_loj = Reference(ws, min_col=11, min_row=4, max_row=6)
    pie_loj.add_data(pdata_loj, titles_from_data=True)
    pie_loj.set_categories(labels_loj)
    pie_loj.dataLabels = DataLabelList()
    pie_loj.dataLabels.showPercent = True
    pie_loj.dataLabels.showVal = False
    pie_loj.dataLabels.showCatName = True
    pie_loj.width = 12
    pie_loj.height = 9
    for i, ser in enumerate(pie_loj.series):
        for j, col in enumerate((_TEAL, _BLUE)):
            pt = DataPoint(idx=j)
            pt.graphicalProperties.solidFill = col
            ser.data_points.append(pt)
    ws.add_chart(pie_loj, 'J' + str(chart0))

    # stacked column
    stack = BarChart()
    stack.type = 'col'
    stack.grouping = 'stacked'
    stack.overlap = 100
    stack.title = 'Aylık kırılım (Depo / Nakliye / Taxes)'
    stack.y_axis.title = 'EUR'
    stack.y_axis.numFmt = '#,##0'
    data = Reference(ws, min_col=2, min_row=4, max_col=4, max_row=last)
    stack.add_data(data, titles_from_data=True)
    stack.set_categories(cats)
    stack.shape = 4
    stack.width = 15
    stack.height = 9
    colors = (_TEAL, _BLUE, _RED)
    for i, ser in enumerate(stack.series):
        ser.graphicalProperties.solidFill = colors[i]
    ws.add_chart(stack, 'A' + str(chart0 + chart_gap))

    pie = PieChart()
    pie.title = 'Yıl payı'
    labels = Reference(ws, min_col=10, min_row=5, max_row=7)
    pdata = Reference(ws, min_col=11, min_row=4, max_row=7)
    pie.add_data(pdata, titles_from_data=True)
    pie.set_categories(labels)
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    pie.dataLabels.showVal = False
    pie.dataLabels.showCatName = True
    pie.width = 12
    pie.height = 9
    for i, ser in enumerate(pie.series):
        for j, col in enumerate((_TEAL, _BLUE, _RED)):
            pt = DataPoint(idx=j)
            pt.graphicalProperties.solidFill = col
            ser.data_points.append(pt)
    ws.add_chart(pie, 'J' + str(chart0 + chart_gap))

    bar = BarChart()
    bar.type = 'col'
    bar.title = 'Nakliye: kamyon × maliyet'
    bar.y_axis.title = 'Kamyon'
    bar.y_axis.axId = 100
    kamyon = Reference(ws, min_col=5, min_row=4, max_row=last)
    bar.add_data(kamyon, titles_from_data=True)
    bar.set_categories(cats)
    bar.shape = 4
    bar.y_axis.crosses = 'min'
    if bar.series:
        bar.series[0].graphicalProperties.solidFill = _BLUE

    line = LineChart()
    line.y_axis.axId = 200
    line.y_axis.title = 'EUR'
    line.y_axis.crosses = 'max'
    nak = Reference(ws, min_col=3, min_row=4, max_row=last)
    line.add_data(nak, titles_from_data=True)
    if line.series:
        line.series[0].graphicalProperties.line.solidFill = _TEAL
        line.series[0].graphicalProperties.line.width = 25000
    bar += line
    bar.width = 15
    bar.height = 9
    ws.add_chart(bar, 'A' + str(chart0 + chart_gap * 2))

    palet = BarChart()
    palet.type = 'col'
    palet.title = 'Depolama palet (koli ÷ 30)'
    palet.y_axis.title = 'Palet'
    pdata = Reference(ws, min_col=6, min_row=4, max_row=last)
    palet.add_data(pdata, titles_from_data=True)
    palet.set_categories(cats)
    palet.shape = 4
    palet.width = 12
    palet.height = 9
    if palet.series:
        palet.series[0].graphicalProperties.solidFill = _CYAN
    ws.add_chart(palet, 'J' + str(chart0 + chart_gap * 2))

    if rota_adlari:
        sc = 13
        cell = ws.cell(4, sc, 'Ay')
        cell.font = hdr
        cell.fill = fill_ink
        cell.alignment = center
        cell.border = _xlsx_border()
        for j, rota in enumerate(rota_adlari):
            cell = ws.cell(4, sc + 1 + j, _rota_kisa(rota))
            cell.font = hdr
            cell.fill = fill_ink
            cell.alignment = center
            cell.border = _xlsx_border()
        for i, a in enumerate(dolu):
            row = 5 + i
            m = int(str(a['ay'])[5:7])
            ws.cell(row, sc, f"{_AY_KISA[m - 1]} {str(yil)[2:]}").font = numf
            ws.cell(row, sc).border = _xlsx_border()
            for j, rota in enumerate(rota_adlari):
                found = next((x for x in (a.get('rotalar') or []) if x.get('rota') == rota), None)
                t = float((found or {}).get('tutar') or 0)
                cell = ws.cell(row, sc + 1 + j, t)
                cell.font = numf
                cell.number_format = '#,##0.00'
                cell.alignment = right
                cell.border = _xlsx_border()
        rota_stack = BarChart()
        rota_stack.type = 'col'
        rota_stack.grouping = 'stacked'
        rota_stack.overlap = 100
        rota_stack.title = 'Güzergâh: aylık nakliye'
        rota_stack.y_axis.title = 'EUR'
        rota_stack.y_axis.numFmt = '#,##0'
        rdata = Reference(ws, min_col=sc + 1, min_row=4, max_col=sc + len(rota_adlari), max_row=last)
        rcats = Reference(ws, min_col=sc, min_row=5, max_row=last)
        rota_stack.add_data(rdata, titles_from_data=True)
        rota_stack.set_categories(rcats)
        rota_stack.shape = 4
        rota_stack.width = 15
        rota_stack.height = 10
        for i, ser in enumerate(rota_stack.series):
            ser.graphicalProperties.solidFill = _ROTA_COLORS[i % len(_ROTA_COLORS)]
        ws.add_chart(rota_stack, 'A' + str(chart0 + chart_gap * 3))

        p0 = last + 2
        ws.cell(p0, sc, 'Güzergâh').font = hdr
        ws.cell(p0, sc).fill = fill_ink
        ws.cell(p0, sc).border = _xlsx_border()
        ws.cell(p0, sc + 1, 'Tutar €').font = hdr
        ws.cell(p0, sc + 1).fill = fill_ink
        ws.cell(p0, sc + 1).border = _xlsx_border()
        ws.cell(p0, sc + 2, 'Sefer').font = hdr
        ws.cell(p0, sc + 2).fill = fill_ink
        ws.cell(p0, sc + 2).border = _xlsx_border()
        ws.cell(p0, sc + 3, '€ / sefer').font = hdr
        ws.cell(p0, sc + 3).fill = fill_ink
        ws.cell(p0, sc + 3).border = _xlsx_border()
        for i, rota in enumerate(rota_adlari):
            q = rota_qty.get(rota, 0)
            t = rota_tot[rota]
            birim = (t / q) if q else 0.0
            ws.cell(p0 + 1 + i, sc, _rota_kisa(rota)).font = numf
            ws.cell(p0 + 1 + i, sc).border = _xlsx_border()
            cell = ws.cell(p0 + 1 + i, sc + 1, t)
            cell.font = numf
            cell.number_format = '#,##0.00'
            cell.border = _xlsx_border()
            cell = ws.cell(p0 + 1 + i, sc + 2, q)
            cell.font = numf
            cell.number_format = '#,##0'
            cell.border = _xlsx_border()
            cell = ws.cell(p0 + 1 + i, sc + 3, birim)
            cell.font = numf
            cell.number_format = '#,##0.00'
            cell.border = _xlsx_border()
        rota_pie = PieChart()
        rota_pie.title = 'Güzergâh payı'
        rlabels = Reference(ws, min_col=sc, min_row=p0 + 1, max_row=p0 + len(rota_adlari))
        rpdata = Reference(ws, min_col=sc + 1, min_row=p0, max_row=p0 + len(rota_adlari))
        rota_pie.add_data(rpdata, titles_from_data=True)
        rota_pie.set_categories(rlabels)
        rota_pie.dataLabels = DataLabelList()
        rota_pie.dataLabels.showPercent = True
        rota_pie.dataLabels.showVal = False
        rota_pie.dataLabels.showCatName = True
        rota_pie.width = 12
        rota_pie.height = 10
        if rota_pie.series:
            for j in range(len(rota_adlari)):
                pt = DataPoint(idx=j)
                pt.graphicalProperties.solidFill = _ROTA_COLORS[j % len(_ROTA_COLORS)]
                rota_pie.series[0].data_points.append(pt)
        ws.add_chart(rota_pie, 'J' + str(chart0 + chart_gap * 3))

        birim_adlari = sorted(
            rota_adlari,
            key=lambda r: (
                -(rota_tot[r] / rota_qty[r] if rota_qty.get(r) else 0),
                -rota_tot[r],
                r,
            ),
        )
        b0 = p0 + len(rota_adlari) + 3
        ws.cell(b0, sc, 'Sıralama: birim maliyet').font = hdr
        ws.cell(b0, sc).fill = fill_ink
        ws.cell(b0, sc).border = _xlsx_border()
        ws.merge_cells(start_row=b0, start_column=sc, end_row=b0, end_column=sc + 3)
        ws.cell(b0 + 1, sc, 'Güzergâh').font = hdr
        ws.cell(b0 + 1, sc).fill = fill_ink
        ws.cell(b0 + 1, sc).border = _xlsx_border()
        ws.cell(b0 + 1, sc + 1, 'Sefer').font = hdr
        ws.cell(b0 + 1, sc + 1).fill = fill_ink
        ws.cell(b0 + 1, sc + 1).border = _xlsx_border()
        ws.cell(b0 + 1, sc + 2, 'Tutar €').font = hdr
        ws.cell(b0 + 1, sc + 2).fill = fill_ink
        ws.cell(b0 + 1, sc + 2).border = _xlsx_border()
        ws.cell(b0 + 1, sc + 3, '€ / sefer').font = hdr
        ws.cell(b0 + 1, sc + 3).fill = fill_ink
        ws.cell(b0 + 1, sc + 3).border = _xlsx_border()
        for i, rota in enumerate(birim_adlari):
            q = rota_qty.get(rota, 0)
            t = rota_tot[rota]
            birim = (t / q) if q else 0.0
            ws.cell(b0 + 2 + i, sc, _rota_kisa(rota)).font = numf
            ws.cell(b0 + 2 + i, sc).border = _xlsx_border()
            cell = ws.cell(b0 + 2 + i, sc + 1, q)
            cell.font = numf
            cell.number_format = '#,##0'
            cell.border = _xlsx_border()
            cell = ws.cell(b0 + 2 + i, sc + 2, t)
            cell.font = numf
            cell.number_format = '#,##0.00'
            cell.border = _xlsx_border()
            cell = ws.cell(b0 + 2 + i, sc + 3, birim)
            cell.font = numf
            cell.number_format = '#,##0.00'
            cell.border = _xlsx_border()
        birim_bar = BarChart()
        birim_bar.type = 'bar'
        birim_bar.title = 'Sıralama: birim maliyet (€ / sefer)'
        birim_bar.y_axis.title = None
        birim_bar.x_axis.title = 'EUR / sefer'
        birim_bar.x_axis.numFmt = '#,##0'
        bdata = Reference(ws, min_col=sc + 3, min_row=b0 + 1, max_row=b0 + 1 + len(birim_adlari))
        bcats = Reference(ws, min_col=sc, min_row=b0 + 2, max_row=b0 + 1 + len(birim_adlari))
        birim_bar.add_data(bdata, titles_from_data=True)
        birim_bar.set_categories(bcats)
        birim_bar.shape = 4
        birim_bar.width = 15
        birim_bar.height = max(8, min(16, 1.1 * len(birim_adlari)))
        if birim_bar.series:
            birim_bar.series[0].graphicalProperties.solidFill = _BLUE
        ws.add_chart(birim_bar, 'A' + str(chart0 + chart_gap * 4))

    for col, w in enumerate((14, 12, 12, 12, 10, 10, 10, 12, 4, 14, 14), 1):
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(col)].width = w
    return ws


def bosna_rapor_xlsx(rapor, yil, ay_keys=None, donem=None, donem_etiket=None):
    """Ekrandaki Logistics Cost tablosu (Bosnia sheet). Invoices ve Grafikler yok."""
    import io
    from openpyxl import Workbook
    yil = str(yil)
    aylar = _bosna_yil_aylar(rapor, yil)
    want = set(donem_ay_keys(yil, donem, ay_keys))
    if want:
        aylar = [a for a in aylar if a.get('ay') in want]
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    _bosna_tablo_sheet(wb, yil, aylar, donem_etiket)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _bosna_yil_faturalar(rapor, yil):
    yil = str(yil)
    rows = (rapor or {}).get('faturalar') or []
    return [f for f in rows if str(f.get('donem') or f.get('datum') or '').startswith(yil)]


def _bosna_fatura_sheet(wb, yil, faturalar, donem_etiket=None):
    from openpyxl.utils import get_column_letter
    ws = wb.create_sheet('Invoices', 1)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'A5'
    n = len(faturalar or [])
    ws.auto_filter.ref = f'A4:E{4 + n}' if n else 'A4:E4'
    ink = _xlsx_font(bold=True, color='FFFFFF', size=11)
    lab = _xlsx_font(size=10, color=_INK)
    numf = _xlsx_font(size=10)
    title = _xlsx_font(bold=True, size=16, color=_INK)
    kicker = _xlsx_font(bold=True, size=10, color=_TEAL)
    muted = _xlsx_font(size=9, color='6F6B64')
    totf = _xlsx_font(bold=True, size=10, color=_TEAL)
    center = _xlsx_align('center', 'center')
    left = _xlsx_align('left', 'center')
    right = _xlsx_align('right', 'center')
    fill_ink = _xlsx_fill(_INK)
    fill_cream = _xlsx_fill(_CREAM)
    fill_tot = _xlsx_fill(_TEAL_BG)
    fill_white = _xlsx_fill('FFFFFF')

    ws.merge_cells('A1:E1')
    ws.cell(1, 1, 'Bosnia Herzegovina').font = kicker
    ws.merge_cells('A2:E2')
    ws.cell(2, 1, f'Invoices · {donem_etiket or yil}').font = title
    ws.merge_cells('A3:E3')
    ws.cell(3, 1, 'Invoice No · Datum · Neto (excl. VAT). Neto BAM = Excel Neto. Neto EUR = Neto / 1,95583.').font = muted
    ws.row_dimensions[1].height = 16
    ws.row_dimensions[2].height = 24
    headers = ('Invoice No', 'Datum', 'Neto BAM', 'Neto EUR', 'Type')
    for i, h in enumerate(headers, 1):
        cell = ws.cell(4, i, h)
        cell.font = ink
        cell.fill = fill_ink
        cell.alignment = center
        cell.border = _xlsx_border()
    bam_tot = 0.0
    eur_tot = 0.0
    for i, f in enumerate(faturalar or []):
        r = 5 + i
        fill = fill_cream if i % 2 else fill_white
        datum = f.get('datum') or ''
        if len(str(datum)) >= 10:
            datum = str(datum)[:10]
        if (f.get('tarih_kaynak') or '') == 'dpu' and datum:
            datum = f'{datum} · DPU'
        ws.cell(r, 1, f.get('fatura_no') or '').font = lab
        ws.cell(r, 2, datum).font = numf
        bam = round(float(f.get('neto_bam') or 0), 2)
        eur = round(float(f.get('neto_eur') or 0), 2)
        bam_tot += bam
        eur_tot += eur
        c_bam = ws.cell(r, 3, bam)
        c_eur = ws.cell(r, 4, eur)
        c_bam.number_format = '#,##0.00'
        c_eur.number_format = '#,##0.00'
        c_bam.font = numf
        c_eur.font = numf
        ws.cell(r, 5, f.get('tip') or '').font = numf
        for c in range(1, 6):
            ws.cell(r, c).fill = fill
            ws.cell(r, c).border = _xlsx_border()
            ws.cell(r, c).alignment = left if c in (1, 2, 5) else right
    tot_row = 5 + len(faturalar or [])
    ws.cell(tot_row, 1, f'Total · {len(faturalar or [])} Invoices').font = totf
    ws.cell(tot_row, 2, '').font = totf
    ws.cell(tot_row, 3, round(bam_tot, 2)).font = totf
    ws.cell(tot_row, 4, round(eur_tot, 2)).font = totf
    ws.cell(tot_row, 5, '').font = totf
    ws.cell(tot_row, 3).number_format = '#,##0.00'
    ws.cell(tot_row, 4).number_format = '#,##0.00'
    for c in range(1, 6):
        ws.cell(tot_row, c).fill = fill_tot
        ws.cell(tot_row, c).border = _xlsx_border()
        ws.cell(tot_row, c).alignment = left if c in (1, 2, 5) else right
    for col, w in enumerate((28, 16, 14, 14, 18), 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    return ws


def bosna_fatura_xlsx(rapor, yil, faturalar=None, donem_etiket=None):
    import io
    from openpyxl import Workbook
    yil = str(yil)
    if faturalar is None:
        faturalar = _bosna_yil_faturalar(rapor, yil)
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    _bosna_fatura_sheet(wb, yil, faturalar, donem_etiket)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def maliyet_bosna_export():
    """GET/POST /api/maliyet/bosna/export — ekrandaki Bosnia tablosu; faturalar için sadece=faturalar."""
    from flask import jsonify, request, send_file

    body = request.get_json(silent=True) or {}
    yil = str(body.get('yil') or request.args.get('yil') or datetime.date.today().year)
    if not re.fullmatch(r'20\d{2}', yil):
        return jsonify({'success': False, 'error': 'Geçerli bir yıl girin'}), 400
    sadece = str(body.get('sadece') or request.args.get('sadece') or '').strip().lower()
    ay_keys = donem_ay_keys(yil, body.get('donem'), body.get('aylar'))
    donem_etiket = str(body.get('donem_etiket') or '').strip() or yil
    rapor = body.get('rapor')
    if not isinstance(rapor, dict) or not rapor.get('aylar'):
        _ozet, satirlar = bosna_kayitli_veri()
        satirlar = [s for s in satirlar if (s.get('donem') or '').startswith(yil)]
        rapor = bosna_aylik_rapor(satirlar)
        rapor['yil'] = yil
        bosna_ciro_uygula(rapor)
    else:
        bosna_ciro_uygula(rapor, overwrite=False)
        if not isinstance(rapor.get('faturalar'), list):
            _ozet, satirlar = bosna_kayitli_veri()
            satirlar = [s for s in satirlar if (s.get('donem') or '').startswith(yil)]
            rapor = dict(rapor)
            rapor['faturalar'] = bosna_fatura_listesi(satirlar)
    if sadece == 'faturalar':
        faturalar = _bosna_yil_faturalar(rapor, yil)
        if ay_keys:
            want = set(ay_keys)
            faturalar = [f for f in faturalar
                         if str(f.get('donem') or f.get('datum') or '')[:7] in want]
        if not faturalar:
            return jsonify({'success': False, 'error': 'Bu dönem için fatura listesi yok'}), 400
        buf = bosna_fatura_xlsx(rapor, yil, faturalar=faturalar, donem_etiket=donem_etiket)
        slug = '' if donem_etiket == yil else '_' + re.sub(r'[^A-Za-z0-9]+', '', donem_etiket)[:16]
        return send_file(
            buf,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=f'bosna_invoices_{yil}{slug}.xlsx',
        )
    aylar = [a for a in (rapor.get('aylar') or []) if str(a.get('ay') or '').startswith(yil)]
    if ay_keys:
        want = set(ay_keys)
        aylar = [a for a in aylar if a.get('ay') in want]
    if not aylar:
        return jsonify({'success': False, 'error': 'Bu dönem için Bosna raporu yok'}), 400
    buf = bosna_rapor_xlsx(rapor, yil, ay_keys=ay_keys, donem=body.get('donem'), donem_etiket=donem_etiket)
    slug = '' if donem_etiket == yil else '_' + re.sub(r'[^A-Za-z0-9]+', '', donem_etiket)[:16]
    return send_file(
        buf,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f'bosna_lojistik_{yil}{slug}.xlsx',
    )
