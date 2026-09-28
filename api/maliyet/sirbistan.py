# api/maliyet/sirbistan.py
# Sırbistan: M&M DEPO Excel (EUR) + BGDT nakliye PDF (RSD, KDV hariç) +
# LOGISTICS COSTS Serbia 2026. Depolama (koli stok) lojistiğin büyük kısmı.

import calendar
import datetime
import logging
import os
import re
import unicodedata
from collections import defaultdict

import pdfplumber
from flask import jsonify, request
from openpyxl import load_workbook

from api.db import get_conn
from api.kur import get_tcmb_kurlar
from api.maliyet.hesap import to_eur
from api.maliyet.ulke_tablo import tablo_kaydet, ulke_ciro_kaydet

_BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
_RS_DIR = os.path.join(_BASE, 'SIRBİSTAN MALİYET', '2026')
_LOGISTICS_XLSX = os.path.join(_BASE, 'Logistic Cost', 'LOGISTICS COSTS.xlsx')

_MONTH_TR = {
    'ocak': 1, 'januar': 1, 'january': 1, 'jan': 1,
    'subat': 2, 'şubat': 2, 'februar': 2, 'february': 2, 'feb': 2,
    'mart': 3, 'march': 3, 'mar': 3,
    'nisan': 4, 'april': 4, 'apr': 4,
    'mayis': 5, 'mayıs': 5, 'maj': 5, 'may': 5,
    'haziran': 6, 'jun': 6, 'june': 6,
    'temmuz': 7, 'jul': 7, 'july': 7,
    'agustos': 8, 'ağustos': 8, 'avgust': 8, 'august': 8,
    'eylul': 9, 'eylül': 9, 'septembar': 9, 'september': 9,
    'ekim': 10, 'oktobar': 10, 'october': 10,
    'kasim': 11, 'kasım': 11, 'novembar': 11, 'november': 11,
    'aralik': 12, 'aralık': 12, 'decembar': 12, 'december': 12,
}


def _fold(s):
    s = unicodedata.normalize('NFKD', str(s or ''))
    s = ''.join(c for c in s if not unicodedata.combining(c)).lower()
    return s.translate(str.maketrans('çğıöşüı', 'cgiosui'))


def _num(value):
    if value is None or value == '':
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace('\xa0', '').replace(' ', '').replace('€', '')
    s = s.replace('RSD', '').replace('rsd', '').replace('EUR', '')
    if not s:
        return None
    if ',' in s and '.' in s:
        if s.rfind(',') > s.rfind('.'):
            s = s.replace('.', '').replace(',', '.')
        else:
            s = s.replace(',', '')
    elif ',' in s:
        if re.search(r',\d{1,2}$', s):
            s = s.replace('.', '').replace(',', '.')
        else:
            s = s.replace(',', '')
    try:
        return float(s)
    except ValueError:
        return None


def _folder_ay_no(name):
    n = _fold(name)
    m = re.match(r'^(\d{1,2})\b', n)
    if m:
        v = int(m.group(1))
        if 1 <= v <= 12:
            return v
    for key, mon in _MONTH_TR.items():
        if key in n:
            return mon
    return None


def _sheet_ay_no(name):
    n = _fold(name).replace("'", '').replace(' ', '')
    for key, mon in sorted(_MONTH_TR.items(), key=lambda x: -len(x[0])):
        if key in n:
            return mon
    return None


def _is_depo_name(name):
    n = _fold(name)
    if 'nakliye' in n or 'bgdt' in n:
        return False
    return any(k in n for k in ('depo', 'depolama', 'storage', 'bglg'))


def _is_nakliye_name(name):
    n = _fold(name)
    return 'nakliye' in n or 'bgdt' in n or 'fakturisanje' in n or n.startswith('2026') or 'madame coco' in n


def parse_rs_depo_xlsx(path, ay_no=None, yil=2026):
    """DEPO Excel 'for invoicing' bloğu → EUR kalemler + ortalama koli."""
    try:
        wb = load_workbook(path, data_only=True)
    except Exception:
        logging.exception('rs depo open %s', path)
        return None
    try:
        ws = None
        if ay_no:
            for sn in wb.sheetnames:
                if _sheet_ay_no(sn) == ay_no:
                    ws = wb[sn]
                    break
        if ws is None:
            ws = wb[wb.sheetnames[0]]
            ay_no = _sheet_ay_no(wb.sheetnames[0]) or ay_no

        kur = None
        period_start = period_end = None
        fatura_ref = None
        for row in ws.iter_rows(min_row=1, max_row=12, values_only=True):
            cells = list(row)
            joined = ' '.join(str(c or '') for c in cells)
            for i, c in enumerate(cells):
                low = str(c or '').strip().lower()
                if low.startswith('exchange') and i + 2 < len(cells):
                    kur = _num(cells[i + 2]) or _num(cells[i + 1])
                if low == 'order' and i + 1 < len(cells) and cells[i + 1]:
                    fatura_ref = str(cells[i + 1]).strip()
            dates = []
            for c in cells:
                s = str(c or '').strip().rstrip('.')
                for fmt in ('%d.%m.%Y', '%d/%m/%Y'):
                    try:
                        dates.append(datetime.datetime.strptime(s, fmt).date())
                    except ValueError:
                        pass
            if len(dates) >= 2 and 'period' in joined.lower():
                period_start, period_end = dates[0], dates[1]

        inbound = {'miktar': 0.0, 'tutar': 0.0}
        outbound = {'miktar': 0.0, 'tutar': 0.0}
        storage_tutar = 0.0
        invoicing = False
        for row in ws.iter_rows(min_row=1, max_row=80, values_only=True):
            cells = list(row)
            joined = ' '.join(str(c or '') for c in cells).lower()
            if 'for invoicing' in joined:
                invoicing = True
                continue
            if not invoicing:
                continue
            label = ' '.join(str(c or '').strip() for c in cells[3:6] if c).strip().lower()
            if not label:
                label = ' '.join(str(c or '').strip() for c in cells if c).lower()
            pcs = None
            eur = None
            # pcs | price | amount EUR — genelde H, I, J
            for idx in range(len(cells) - 1, 4, -1):
                v = _num(cells[idx])
                if v is None:
                    continue
                nxt = str(cells[idx - 1] or '').lower() if idx else ''
                if eur is None and ('amount eur' in joined or idx >= 8):
                    # son sayısal EUR, bir önceki RSD olabilir
                    pass
            # Sabit kolon: H=pcs(8), I=price(9), J=EUR(10) — 1-based 8,9,10 → idx 7,8,9
            if len(cells) > 9:
                pcs = _num(cells[7])
                eur = _num(cells[9])
            if 'inbound box' in label:
                inbound['miktar'] += float(pcs or 0)
                inbound['tutar'] += float(eur or 0)
            elif 'outbound box' in label or 'ooubound box' in label:
                outbound['miktar'] += float(pcs or 0)
                outbound['tutar'] += float(eur or 0)
            elif 'inbound pallet' in label:
                inbound['tutar'] += float(eur or 0)
            elif 'outbound pallet' in label:
                outbound['tutar'] += float(eur or 0)
            elif label.strip() == 'storage' or label.endswith('storage'):
                storage_tutar = float(eur or 0)
            if 'sum' in joined and eur:
                break

        # Günlük STORAGE BOX ortalaması
        box_vals = []
        for row in ws.iter_rows(min_row=8, max_row=45, values_only=True):
            cells = list(row)
            day = str(cells[0] or '').strip()
            if not re.match(r'^\d+\.?$', day):
                continue
            # J kolonu (idx 9) STORAGE BOX
            box = _num(cells[9]) if len(cells) > 9 else None
            if box and box > 0:
                box_vals.append(box)
        avg_box = (sum(box_vals) / len(box_vals)) if box_vals else 0.0

        if inbound['tutar'] < 0.01 and outbound['tutar'] < 0.01 and storage_tutar < 0.01:
            return None

        last = calendar.monthrange(yil, ay_no or 1)[1]
        ay_key = f'{yil}-{(ay_no or 1):02d}'
        if not period_start:
            period_start = datetime.date(yil, ay_no or 1, 1)
            period_end = datetime.date(yil, ay_no or 1, last)
        return {
            'ay': ay_key,
            'ay_no': ay_no,
            'kur_rsd': kur,
            'order': fatura_ref,
            'donem_baslangic': period_start.isoformat(),
            'donem_bitis': period_end.isoformat(),
            'inbound': {'miktar': round(inbound['miktar'], 2), 'tutar': round(inbound['tutar'], 2)},
            'outbound': {'miktar': round(outbound['miktar'], 2), 'tutar': round(outbound['tutar'], 2)},
            'storage': {'miktar': round(avg_box, 2), 'tutar': round(storage_tutar, 2)},
        }
    finally:
        wb.close()


def parse_rs_bgdt_pdf(path):
    """BGDT nakliye faturası — Osnovica RSD (KDV hariç)."""
    try:
        with pdfplumber.open(path) as pdf:
            text = '\n'.join((p.extract_text(x_tolerance=2, y_tolerance=3) or '') for p in pdf.pages)
    except Exception:
        logging.exception('rs bgdt %s', path)
        return None
    if len((text or '').strip()) < 40:
        return None
    m_no = re.search(r'RA[ČC]UN\s+(BGDT\s*[-–]\s*\d+(?:/\d+)?)', text, re.I)
    fatura_no = (m_no.group(1) if m_no else '').replace('  ', ' ').strip()
    if not fatura_no:
        m_no = re.search(r'(BGDT\s*[-–]\s*\d+(?:/\d+)?)', text, re.I)
        fatura_no = (m_no.group(1) if m_no else '').strip()
    if not fatura_no:
        return None
    fatura_no = re.sub(r'\s+', ' ', fatura_no)

    def _d(s):
        for fmt in ('%d.%m.%Y', '%d/%m/%Y'):
            try:
                return datetime.datetime.strptime(s, fmt).date()
            except ValueError:
                pass
        return None

    m_prom = re.search(r'Datum prometa:\s*(\d{1,2}[./]\d{1,2}[./]\d{4})', text, re.I)
    m_izd = re.search(r'Datum izdavanja ra[čc]una:\s*(\d{1,2}[./]\d{1,2}[./]\d{4})', text, re.I)
    donem = _d(m_prom.group(1)) if m_prom else None
    fatura_tarihi = _d(m_izd.group(1)) if m_izd else donem
    m_net = re.search(r'Osnovica:\s*([\d.]+,\d{2})', text)
    net_rsd = _num(m_net.group(1)) if m_net else None
    if net_rsd is None:
        return None
    yil = (donem or fatura_tarihi or datetime.date.today()).year
    ay = (donem or fatura_tarihi or datetime.date.today()).month
    last = calendar.monthrange(yil, ay)[1]
    return {
        'fatura_no': fatura_no,
        'tutar_rsd': round(net_rsd, 2),
        'donem_baslangic': f'{yil}-{ay:02d}-01',
        'donem_bitis': f'{yil}-{ay:02d}-{last:02d}',
        'fatura_tarihi': (fatura_tarihi or donem).isoformat() if (fatura_tarihi or donem) else None,
        'ay': f'{yil}-{ay:02d}',
        'kaynak': 'bgdt',
    }


def parse_rs_nakliye_ukupno(path, ay_no):
    """Nakliye Excel UKUPNO satırı (RSD) — PDF metni yoksa yedek."""
    try:
        wb = load_workbook(path, data_only=True)
    except Exception:
        return None
    try:
        ws = None
        for sn in wb.sheetnames:
            if _sheet_ay_no(sn) == ay_no:
                ws = wb[sn]
                break
        if ws is None:
            return None
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 90, 90), values_only=True):
            cells = list(row)
            labels = [str(c or '').strip().lower() for c in cells if c]
            if not any(l.startswith('ukupno') for l in labels):
                continue
            nums = [_num(c) for c in cells]
            nums = [n for n in nums if n and n > 100000]
            if nums:
                return max(nums)
        return None
    finally:
        wb.close()


def _logistics_2026_aylar():
    wb = load_workbook(_LOGISTICS_XLSX, data_only=True)
    ws = wb['Serbia']
    cols = [
        ('2026-01', 5, 6), ('2026-02', 7, 8), ('2026-03', 9, 10),
        ('2026-04', 13, 14),
    ]
    row_map = {'inbound': 21, 'storage': 22, 'outbound': 23, 'transport': 24}
    ciro_row = 27
    aylar = []
    for ay, qc, cc in cols:
        a = {'ay': ay, 'ciro_eur': float(ws.cell(ciro_row, qc).value or 0)}
        for kod, rr in row_map.items():
            a[kod] = {
                'miktar': float(ws.cell(rr, qc).value or 0),
                'tutar': float(ws.cell(rr, cc).value or 0),
            }
        if sum(a[k]['tutar'] for k in row_map) > 0.005 or a['ciro_eur'] > 0:
            aylar.append(a)
    wb.close()
    return aylar


def _ciro_sayfa1():
    """Sayfa1 2026 Q2: Nisan–Haziran."""
    out = {}
    try:
        wb = load_workbook(_LOGISTICS_XLSX, data_only=True)
        ws = wb['Sayfa1']

        def _eu(v):
            if isinstance(v, (int, float)):
                return float(v)
            s = str(v or '').replace('\xa0', '').replace('€', '').strip()
            if not s:
                return 0.0
            if ',' in s:
                s = s.replace('.', '').replace(',', '.')
            try:
                return float(re.sub(r'[^\d.-]', '', s) or 0)
            except ValueError:
                return 0.0

        for row in ws.iter_rows(min_row=4, max_row=16, values_only=True):
            ad = _fold(row[2] if len(row) > 2 else '')
            if 'sirbistan' not in ad and 'serbia' not in ad:
                continue
            for ay, col in (('2026-04', 3), ('2026-05', 4), ('2026-06', 5)):
                n = _eu(row[col] if len(row) > col else None)
                if n:
                    out[ay] = n
            break
        wb.close()
    except Exception:
        logging.exception('rs sayfa1 ciro')
    return out


def _scan_month_dirs(root=None):
    root = root or _RS_DIR
    if not os.path.isdir(root):
        return []
    out = []
    for name in sorted(os.listdir(root)):
        path = os.path.join(root, name)
        if not os.path.isdir(path):
            continue
        ay = _folder_ay_no(name)
        if not ay:
            continue
        depo = nakliye = None
        pdfs = []
        for f in os.listdir(path):
            full = os.path.join(path, f)
            low = f.lower()
            if low.endswith(('.xlsx', '.xls')):
                if _is_depo_name(f) and not depo:
                    depo = full
                elif _is_nakliye_name(f) and not nakliye:
                    nakliye = full
            elif low.endswith('.pdf'):
                pdfs.append(full)
        out.append({'folder': name, 'ay_no': ay, 'depo': depo, 'nakliye': nakliye, 'pdfs': pdfs})
    return out


def _kalem_id_map(cur):
    cur.execute('SELECT id, kod FROM maliyet_kalemleri WHERE aktif')
    return {r[1]: r[0] for r in cur.fetchall()}


def _upsert_fatura(cur, draft, kalem_ids, ulke='rs'):
    fno = draft.get('fatura_no')
    if not fno:
        return None, 'fatura_no yok'
    kalemler = []
    for i, row in enumerate(draft.get('kalemler') or []):
        kod = row.get('kalem_kod')
        kid = kalem_ids.get(kod) or kalem_ids.get('handling')
        if not kid:
            continue
        tutar = float(row.get('tutar') or 0)
        kalemler.append({
            'kalem_id': kid, 'tarih': draft.get('donem_bitis'),
            'aciklama': row.get('aciklama') or kod, 'referans': None,
            'miktar': float(row.get('miktar') or 1),
            'birim_fiyat': float(row.get('birim_fiyat') or tutar),
            'tutar': tutar, 'sira': i,
        })
    if not kalemler:
        return None, 'kalem yok'
    tutar = float(draft.get('tutar') or sum(k['tutar'] for k in kalemler))
    cur.execute(
        'SELECT id FROM maliyet_faturalari WHERE ulke=%s AND fatura_no=%s LIMIT 1',
        (ulke, fno),
    )
    row = cur.fetchone()
    alanlar = {
        'ulke': ulke, 'fatura_no': fno,
        'donem_baslangic': draft.get('donem_baslangic'),
        'donem_bitis': draft.get('donem_bitis'),
        'tutar': round(tutar, 2),
        'para_birimi': draft.get('para_birimi') or 'EUR',
        'fatura_tarihi': draft.get('fatura_tarihi'),
        'notlar': draft.get('notlar') or 'RS M&M depo/nakliye',
    }
    if row:
        fid = row[0]
        cur.execute('''UPDATE maliyet_faturalari SET
            donem_baslangic=%(donem_baslangic)s, donem_bitis=%(donem_bitis)s,
            tutar=%(tutar)s, para_birimi=%(para_birimi)s, fatura_tarihi=%(fatura_tarihi)s,
            notlar=%(notlar)s WHERE id=%(id)s''', {**alanlar, 'id': fid})
    else:
        cur.execute('''INSERT INTO maliyet_faturalari
            (ulke,fatura_no,donem_baslangic,donem_bitis,tutar,para_birimi,fatura_tarihi,notlar)
            VALUES (%(ulke)s,%(fatura_no)s,%(donem_baslangic)s,%(donem_bitis)s,%(tutar)s,
                    %(para_birimi)s,%(fatura_tarihi)s,%(notlar)s) RETURNING id''', alanlar)
        fid = cur.fetchone()[0]
    cur.execute('DELETE FROM maliyet_fatura_kalemleri WHERE fatura_id=%s', (fid,))
    for k in kalemler:
        cur.execute('''INSERT INTO maliyet_fatura_kalemleri
            (fatura_id,kalem_id,tarih,aciklama,referans,miktar,birim_fiyat,tutar,sira)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)''',
            (fid, k['kalem_id'], k['tarih'], k['aciklama'], k['referans'],
             k['miktar'], k['birim_fiyat'], k['tutar'], k['sira']))
    return fid, None


def _draft_from_depo(det):
    ay = det['ay']
    kalemler = []
    for kod, tablo, ad in (
        ('box_in', 'inbound', 'Inbound boxes'),
        ('storage', 'storage', 'Storage (avr. box)'),
        ('box_out', 'outbound', 'Outbound boxes'),
    ):
        src = det.get(tablo) or {}
        tutar = float(src.get('tutar') or 0)
        miktar = float(src.get('miktar') or 0)
        if tutar <= 0 and miktar <= 0:
            continue
        kalemler.append({
            'aciklama': ad, 'kalem_kod': kod,
            'miktar': miktar or 1,
            'birim_fiyat': (tutar / miktar) if miktar else tutar,
            'tutar': round(tutar, 2),
        })
    if not kalemler:
        return None
    return {
        'ulke': 'rs',
        'fatura_no': f"RS-DEPO-{ay}",
        'fatura_tarihi': det.get('donem_bitis'),
        'donem_baslangic': det.get('donem_baslangic'),
        'donem_bitis': det.get('donem_bitis'),
        'para_birimi': 'EUR',
        'tutar': round(sum(k['tutar'] for k in kalemler), 2),
        'kalemler': kalemler,
        'notlar': 'RS DEPO Excel (KDV hariç EUR)',
    }


def _draft_from_bgdt(parsed, kurlar, depo_kur=None):
    rsd = float(parsed['tutar_rsd'])
    kur = depo_kur or float((kurlar or {}).get('RSD') or 0)
    eur = (rsd / kur) if kur else to_eur(rsd, 'RSD', kurlar)
    if eur is None:
        return None
    eur = round(eur, 2)
    return {
        'ulke': 'rs',
        'fatura_no': parsed['fatura_no'],
        'fatura_tarihi': parsed.get('fatura_tarihi'),
        'donem_baslangic': parsed['donem_baslangic'],
        'donem_bitis': parsed['donem_bitis'],
        'para_birimi': 'EUR',
        'tutar': eur,
        'kalemler': [{
            'aciklama': 'Delivery in Serbia (Osnovica)',
            'kalem_kod': 'transport',
            'miktar': 1,
            'birim_fiyat': eur,
            'tutar': eur,
        }],
        'notlar': f"BGDT RSD {parsed['tutar_rsd']:.2f} / {kur:.4f}",
        'ay': parsed['ay'],
        '_tutar_eur': eur,
    }


def seed_rs_from_folder(root=None):
    root = root or _RS_DIR
    months = _scan_month_dirs(root)
    log_aylar = _logistics_2026_aylar()
    ciro_map = {a['ay']: a['ciro_eur'] for a in log_aylar}
    ciro_map.update(_ciro_sayfa1())
    kurlar = get_tcmb_kurlar()

    fatura_ozet = []
    folder_tablo = {}
    transport_by_ay = defaultdict(lambda: {'miktar': 0.0, 'tutar': 0.0})

    conn = get_conn()
    cur = conn.cursor()
    try:
        kalem_ids = _kalem_id_map(cur)
        for item in months:
            yil = 2026
            ay_key = f'{yil}-{item["ay_no"]:02d}'
            det = parse_rs_depo_xlsx(item['depo'], item['ay_no'], yil) if item.get('depo') else None
            if det:
                folder_tablo[ay_key] = det
                draft = _draft_from_depo(det)
                if draft:
                    fid, err = _upsert_fatura(cur, draft, kalem_ids)
                    fatura_ozet.append({
                        'ay': ay_key, 'fatura_no': draft['fatura_no'], 'id': fid,
                        'tutar': draft['tutar'], 'hata': err, 'kaynak': 'depo',
                    })
            depo_kur = (det or {}).get('kur_rsd')
            bgdt_ok = False
            for pdf in item.get('pdfs') or []:
                parsed = parse_rs_bgdt_pdf(pdf)
                if not parsed:
                    continue
                draft = _draft_from_bgdt(parsed, kurlar, depo_kur)
                if not draft:
                    fatura_ozet.append({'pdf': os.path.basename(pdf), 'hata': 'RSD kur yok'})
                    continue
                fid, err = _upsert_fatura(cur, draft, kalem_ids)
                fatura_ozet.append({
                    'ay': draft['ay'], 'fatura_no': draft['fatura_no'], 'id': fid,
                    'tutar': draft['tutar'], 'hata': err, 'kaynak': 'bgdt',
                })
                if not err:
                    transport_by_ay[draft['ay']]['tutar'] += draft['tutar']
                    bgdt_ok = True
            if not bgdt_ok and item.get('nakliye'):
                rsd = parse_rs_nakliye_ukupno(item['nakliye'], item['ay_no'])
                if rsd:
                    last = calendar.monthrange(2026, item['ay_no'])[1]
                    parsed = {
                        'fatura_no': f'RS-NAK-{ay_key}',
                        'tutar_rsd': rsd,
                        'donem_baslangic': f'{ay_key}-01',
                        'donem_bitis': f'{ay_key}-{last:02d}',
                        'fatura_tarihi': f'{ay_key}-{last:02d}',
                        'ay': ay_key,
                    }
                    draft = _draft_from_bgdt(parsed, kurlar, depo_kur)
                    if draft:
                        fid, err = _upsert_fatura(cur, draft, kalem_ids)
                        fatura_ozet.append({
                            'ay': ay_key, 'fatura_no': draft['fatura_no'], 'id': fid,
                            'tutar': draft['tutar'], 'hata': err, 'kaynak': 'nakliye-xlsx',
                        })
                        if not err:
                            transport_by_ay[ay_key]['tutar'] += draft['tutar']
        conn.commit()
    except Exception:
        conn.rollback()
        logging.exception('rs seed')
        return {'success': False, 'error': 'Sırbistan aktarımı başarısız'}
    finally:
        cur.close()
        conn.close()

    tablo = []
    seen = set()
    for a in log_aylar:
        row = dict(a)
        row['ciro_eur'] = ciro_map.get(a['ay'], a.get('ciro_eur') or 0)
        tablo.append(row)
        seen.add(a['ay'])
    for ay, det in sorted(folder_tablo.items()):
        if ay in seen:
            continue
        tr = transport_by_ay.get(ay) or {'miktar': 0.0, 'tutar': 0.0}
        tablo.append({
            'ay': ay,
            'ciro_eur': float(ciro_map.get(ay) or 0),
            'inbound': det['inbound'],
            'storage': det['storage'],
            'outbound': det['outbound'],
            'transport': {
                'miktar': round(tr['tutar'] / 15.0, 2) if tr['tutar'] else 0.0,
                'tutar': round(tr['tutar'], 2),
            },
        })
        seen.add(ay)
    tablo.sort(key=lambda x: x['ay'])
    tablo_kaydet('rs', tablo)
    ulke_ciro_kaydet('rs', {a['ay']: a['ciro_eur'] for a in tablo if a.get('ciro_eur')})
    return {
        'success': True,
        'klasor': root,
        'ay_sayisi': len(tablo),
        'faturalar': fatura_ozet,
        'aylar': [
            {
                'ay': a['ay'], 'ciro_eur': a.get('ciro_eur') or 0,
                'inbound': a.get('inbound'), 'storage': a.get('storage'),
                'outbound': a.get('outbound'), 'transport': a.get('transport'),
                'lojistik': round(
                    float((a.get('inbound') or {}).get('tutar') or 0)
                    + float((a.get('storage') or {}).get('tutar') or 0)
                    + float((a.get('outbound') or {}).get('tutar') or 0)
                    + float((a.get('transport') or {}).get('tutar') or 0), 2),
            }
            for a in tablo
        ],
        'not': 'Oca–Nis LOGISTICS; May–Tem DEPO Excel + BGDT PDF. Depolama koli stoku lojistiği yükseltir.',
    }


def maliyet_rs_seed_post():
    body = request.get_json(silent=True) or {}
    return jsonify(seed_rs_from_folder(body.get('klasor') or _RS_DIR))


def maliyet_rs_onizle_get():
    months = _scan_month_dirs()
    out = []
    for item in months:
        det = parse_rs_depo_xlsx(item['depo'], item['ay_no']) if item.get('depo') else None
        pdfs = []
        for p in item.get('pdfs') or []:
            if 'bgdt' in os.path.basename(p).lower():
                pdfs.append(parse_rs_bgdt_pdf(p))
        out.append({
            'folder': item['folder'], 'ay_no': item['ay_no'],
            'depo': det, 'bgdt': [x for x in pdfs if x],
        })
    return jsonify({
        'success': True,
        'aylar': out,
        'logistics': _logistics_2026_aylar(),
        'ciro_q2': _ciro_sayfa1(),
    })
