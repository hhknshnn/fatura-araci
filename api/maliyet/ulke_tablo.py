# api/maliyet/ulke_tablo.py
# Ülke bazlı aylık lojistik tablosu (Bosna benzeri, EUR).
# Maliyet satırları + ciro JSON'da; Lojistik/Ciro oran UI/export'ta hesaplanır.

import json
import os
import re
from copy import deepcopy

_BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
_TABLO_DIR = os.path.join(_BASE, 'data', 'maliyet_tablo')
_CIRO_PATH = os.path.join(_BASE, 'data', 'ulke_ciro.json')
_BOSNA_CIRO_LEGACY = os.path.join(_BASE, 'data', 'bosna_ciro.json')

# Varsayılan (boş) şablon — Bosna EUR yapısına yakın
_DEFAULT_ITEMS = [
    {'kod': 'inbound', 'label': 'Giriş', 'group': 'Depo'},
    {'kod': 'storage', 'label': 'Depolama', 'group': 'Depo'},
    {'kod': 'outbound', 'label': 'Çıkış', 'group': 'Depo'},
    {'kod': 'transport', 'label': 'Nakliye', 'group': None},
]

ULKE_TABLO_SCHEMA = {
    'nl': {
        'title': 'Hollanda',
        'subtitle': 'Lojistik Maliyet',
        'items': [
            {'kod': 'inbound', 'label': 'Giriş (Palet)', 'group': 'Depo'},
            {'kod': 'storage', 'label': 'Stok (Ort. Palet)', 'group': 'Depo'},
            {'kod': 'outbound', 'label': 'Çıkış (Palet)', 'group': 'Depo'},
            {'kod': 'transport', 'label': 'Nakliye (Palet)', 'group': None},
        ],
    },
    'be': {
        'title': 'Belçika',
        'subtitle': 'Lojistik Maliyet',
        'items': [
            {'kod': 'inbound', 'label': 'Giriş', 'group': 'Depo'},
            {'kod': 'storage', 'label': 'Stok Ve Elleçleme', 'group': 'Depo'},
            {'kod': 'outbound', 'label': 'Çıkış', 'group': 'Depo'},
            {'kod': 'transport', 'label': 'Nakliye (Teslimat)', 'group': None},
        ],
    },
    'mk': {
        'title': 'Makedonya',
        'subtitle': 'Lojistik Maliyet',
        'items': [
            {'kod': 'inbound', 'label': 'Giriş', 'group': 'Depo'},
            {'kod': 'storage', 'label': 'Stok (Kamyon)', 'group': 'Depo'},
            {'kod': 'outbound', 'label': 'Çıkış', 'group': 'Depo'},
            {'kod': 'transport', 'label': 'Nakliye (Teslimat)', 'group': None},
        ],
    },
    'xk': {
        'title': 'Kosova',
        'subtitle': 'Lojistik Maliyet',
        'items': [
            {'kod': 'inbound', 'label': 'Giriş', 'group': 'Depo'},
            {'kod': 'storage', 'label': 'Stok (Kamyon)', 'group': 'Depo'},
            {'kod': 'outbound', 'label': 'Çıkış', 'group': 'Depo'},
            {'kod': 'transport', 'label': 'Nakliye (Teslimat)', 'group': None},
        ],
    },
    'ge': {
        'title': 'Gürcistan',
        'subtitle': 'Lojistik Maliyet',
        'items': [
            {'kod': 'inbound', 'label': 'Giriş (Palet)', 'group': 'Depo'},
            {'kod': 'storage', 'label': 'Stok (Ort. Palet)', 'group': 'Depo'},
            {'kod': 'outbound', 'label': 'Çıkış (Palet)', 'group': 'Depo'},
            {'kod': 'transport', 'label': 'Nakliye (Palet)', 'group': None},
        ],
    },
    'kz': {
        'title': 'Kazakistan',
        'subtitle': 'Lojistik Maliyet · Depolama',
        'items': [
            {'kod': 'inbound', 'label': 'Giriş', 'group': 'Depo'},
            {'kod': 'storage', 'label': 'Stok (Ort. Koli)', 'group': 'Depo'},
            {'kod': 'outbound', 'label': 'Çıkış', 'group': 'Depo'},
            # Depolama = depo kalemleri toplamı (LOGISTICS satırı); lojistikte iki kez sayılmaz.
            {'kod': 'depolama', 'label': 'Depolama (Ortalama Koli)', 'group': None, 'skip_total': True},
            {'kod': 'transport', 'label': 'İç Nakliye', 'group': None},
        ],
    },
    'rs': {
        'title': 'Sırbistan',
        'subtitle': 'Lojistik Maliyet · Depolama',
        'items': [
            {'kod': 'inbound', 'label': 'Giriş', 'group': 'Depo'},
            {'kod': 'storage', 'label': 'Stok (Ort. Koli)', 'group': 'Depo'},
            {'kod': 'outbound', 'label': 'Çıkış', 'group': 'Depo'},
            {'kod': 'transport', 'label': 'Nakliye (Palet)', 'group': None},
        ],
    },
}


def _z():
    return {'miktar': 0.0, 'tutar': 0.0}


_DONEM_AYLAR = {
    'year': tuple(range(1, 13)),
    'h1': (1, 2, 3, 4, 5, 6),
    'h2': (7, 8, 9, 10, 11, 12),
    'q1': (1, 2, 3),
    'q2': (4, 5, 6),
    'q3': (7, 8, 9),
    'q4': (10, 11, 12),
}


def donem_ay_keys(yil, donem=None, ay_keys=None):
    """Seçilen dönem → ['YYYY-MM', ...]. Excel ve tablo filtresi ortak."""
    yil = str(yil or '').strip()
    if isinstance(ay_keys, list) and ay_keys:
        keys = [str(x) for x in ay_keys if re.fullmatch(r'20\d{2}-\d{2}', str(x or ''))]
        if keys:
            return keys
    d = str(donem or 'year').strip().lower()
    if re.fullmatch(r'0[1-9]|1[0-2]', d) and re.fullmatch(r'20\d{2}', yil):
        return [f'{yil}-{d}']
    months = _DONEM_AYLAR.get(d) or _DONEM_AYLAR['year']
    if re.fullmatch(r'20\d{2}', yil):
        return [f'{yil}-{m:02d}' for m in months]
    return []


def _schema(ulke):
    base = ULKE_TABLO_SCHEMA.get(ulke) or {
        'title': ulke.upper(),
        'subtitle': 'Lojistik Maliyet',
        'items': list(_DEFAULT_ITEMS),
    }
    return deepcopy(base)


def _tablo_path(ulke):
    return os.path.join(_TABLO_DIR, f'{ulke}.json')


def _ay_bos(ay, items):
    a = {'ay': ay, 'ciro_eur': 0.0, 'lojistik': 0.0}
    for it in items:
        a[it['kod']] = _z()
    return a


def _ay_hesap(a, items):
    loj = 0.0
    for it in items:
        kod = it['kod']
        if not a.get(kod):
            a[kod] = _z()
        a[kod]['miktar'] = round(float(a[kod].get('miktar') or 0), 4)
        a[kod]['tutar'] = round(float(a[kod].get('tutar') or 0), 2)
        if it.get('skip_total'):
            continue
        loj += float(a[kod]['tutar'] or 0)
    a['lojistik'] = round(loj, 2)
    a['ciro_eur'] = round(float(a.get('ciro_eur') or 0), 2)
    return a


def ulke_ciro_yukle():
    """{ulke: {YYYY-MM: eur}} — legacy bosna_ciro.json birleştirilir."""
    out = {}
    try:
        with open(_CIRO_PATH, encoding='utf-8') as f:
            raw = json.load(f) or {}
        if isinstance(raw, dict):
            for uk, mp in raw.items():
                if not isinstance(mp, dict):
                    continue
                bag = {}
                for k, v in mp.items():
                    if re.fullmatch(r'20\d{2}-\d{2}', str(k or '')):
                        try:
                            bag[str(k)] = round(float(v or 0), 2)
                        except (TypeError, ValueError):
                            pass
                out[str(uk)] = bag
    except (OSError, ValueError, TypeError):
        pass
    if 'ba' not in out:
        try:
            with open(_BOSNA_CIRO_LEGACY, encoding='utf-8') as f:
                legacy = json.load(f) or {}
            bag = {}
            for k, v in legacy.items():
                if re.fullmatch(r'20\d{2}-\d{2}', str(k or '')):
                    try:
                        bag[str(k)] = round(float(v or 0), 2)
                    except (TypeError, ValueError):
                        pass
            if bag:
                out['ba'] = bag
        except (OSError, ValueError, TypeError):
            pass
    return out


def ulke_ciro_kaydet(ulke, guncelleme):
    ulke = str(ulke or '').strip().lower()
    if not re.fullmatch(r'[a-z]{2}', ulke):
        raise ValueError('Geçersiz ülke')
    cur = ulke_ciro_yukle()
    bag = dict(cur.get(ulke) or {})
    for k, v in (guncelleme or {}).items():
        key = str(k or '').strip()
        if not re.fullmatch(r'20\d{2}-\d{2}', key):
            continue
        try:
            n = round(float(v or 0), 2)
        except (TypeError, ValueError):
            continue
        if n <= 0:
            bag.pop(key, None)
        else:
            bag[key] = n
    cur[ulke] = dict(sorted(bag.items()))
    os.makedirs(os.path.dirname(_CIRO_PATH), exist_ok=True)
    with open(_CIRO_PATH, 'w', encoding='utf-8') as f:
        json.dump(cur, f, ensure_ascii=False, indent=2)
    # Bosna legacy dosyasını da senkron tut
    if ulke == 'ba':
        os.makedirs(os.path.dirname(_BOSNA_CIRO_LEGACY), exist_ok=True)
        with open(_BOSNA_CIRO_LEGACY, 'w', encoding='utf-8') as f:
            json.dump(cur['ba'], f, ensure_ascii=False, indent=2)
    return cur[ulke]


def tablo_yukle(ulke):
    path = _tablo_path(ulke)
    try:
        with open(path, encoding='utf-8') as f:
            raw = json.load(f) or {}
    except (OSError, ValueError, TypeError):
        raw = {}
    schema = _schema(ulke)
    items = schema['items']
    by = {}
    for a in (raw.get('aylar') or []):
        if not isinstance(a, dict) or not a.get('ay'):
            continue
        by[a['ay']] = _ay_hesap(a, items)
    # ciro overlay
    for ay, v in (ulke_ciro_yukle().get(ulke) or {}).items():
        if ay not in by:
            by[ay] = _ay_bos(ay, items)
        by[ay]['ciro_eur'] = float(v or 0)
        _ay_hesap(by[ay], items)
    return {
        'ulke': ulke,
        'schema': schema,
        'aylar': sorted(by.values(), key=lambda x: x.get('ay') or ''),
        'kaynak': 'tablo',
    }


def tablo_kaydet(ulke, aylar):
    schema = _schema(ulke)
    items = schema['items']
    kodlar = {it['kod'] for it in items}
    # Mevcut diğer yılları koru
    mevcut = {a.get('ay'): a for a in (tablo_yukle(ulke).get('aylar') or []) if a.get('ay')}
    ciro_map = {}
    for raw in aylar or []:
        if not isinstance(raw, dict):
            continue
        ay = str(raw.get('ay') or '').strip()
        if not re.fullmatch(r'20\d{2}-\d{2}', ay):
            continue
        a = _ay_bos(ay, items)
        for kod in kodlar:
            src = raw.get(kod) or {}
            if isinstance(src, dict):
                a[kod]['miktar'] = float(src.get('miktar') or 0)
                a[kod]['tutar'] = float(src.get('tutar') or 0)
        a['ciro_eur'] = float(raw.get('ciro_eur') or 0)
        _ay_hesap(a, items)
        mevcut[ay] = a
        ciro_map[ay] = a['ciro_eur']
    out = sorted(mevcut.values(), key=lambda x: x.get('ay') or '')
    os.makedirs(_TABLO_DIR, exist_ok=True)
    path = _tablo_path(ulke)
    with open(path, 'w', encoding='utf-8') as f:
        json.dump({'ulke': ulke, 'aylar': out}, f, ensure_ascii=False, indent=2)
    if ciro_map:
        ulke_ciro_kaydet(ulke, ciro_map)
    return tablo_yukle(ulke)


def tablo_yil_aylar(rapor, yil):
    schema = (rapor or {}).get('schema') or _schema((rapor or {}).get('ulke'))
    items = schema['items']
    by = {a.get('ay'): a for a in (rapor or {}).get('aylar') or []}
    aylar = []
    for m in range(1, 13):
        key = f'{yil}-{m:02d}'
        a = by.get(key) or _ay_bos(key, items)
        aylar.append(_ay_hesap(deepcopy(a), items))
    return aylar, schema


def maliyet_ulke_tablo_get():
    from flask import jsonify, request
    ulke = str(request.args.get('ulke') or '').strip().lower()
    yil = str(request.args.get('yil') or '').strip()
    if not re.fullmatch(r'[a-z]{2}', ulke) or ulke == 'ba':
        return jsonify({'success': False, 'error': 'Geçerli ülke kodu girin (ba hariç)'}), 400
    if yil and not re.fullmatch(r'20\d{2}', yil):
        return jsonify({'success': False, 'error': 'Geçerli yıl girin'}), 400
    rapor = tablo_yukle(ulke)
    if yil:
        aylar, schema = tablo_yil_aylar(rapor, yil)
        rapor = {
            'ulke': ulke,
            'yil': yil,
            'schema': schema,
            'aylar': aylar,
            'kaynak': 'tablo',
        }
    return jsonify({'success': True, 'rapor': rapor})


def maliyet_ulke_tablo_post():
    from flask import jsonify, request
    body = request.get_json(silent=True) or {}
    ulke = str(body.get('ulke') or '').strip().lower()
    if not re.fullmatch(r'[a-z]{2}', ulke) or ulke == 'ba':
        return jsonify({'success': False, 'error': 'Geçerli ülke kodu girin (ba hariç)'}), 400
    aylar = body.get('aylar')
    if not isinstance(aylar, list):
        return jsonify({'success': False, 'error': 'aylar listesi gerekli'}), 400
    rapor = tablo_kaydet(ulke, aylar)
    return jsonify({'success': True, 'rapor': rapor})


def maliyet_ulke_ciro_post():
    """POST /api/maliyet/ciro — {ulke, aylar:{YYYY-MM: eur}}."""
    from flask import jsonify, request
    body = request.get_json(silent=True) or {}
    ulke = str(body.get('ulke') or '').strip().lower()
    if not re.fullmatch(r'[a-z]{2}', ulke):
        return jsonify({'success': False, 'error': 'ulke gerekli'}), 400
    raw = body.get('aylar') if isinstance(body.get('aylar'), dict) else body.get('ciro')
    if not isinstance(raw, dict):
        return jsonify({'success': False, 'error': 'aylar map gerekli'}), 400
    try:
        kayit = ulke_ciro_kaydet(ulke, raw)
    except ValueError as e:
        return jsonify({'success': False, 'error': str(e)}), 400
    return jsonify({'success': True, 'ulke': ulke, 'ciro': kayit})


def maliyet_ulke_tablo_export():
    """POST /api/maliyet/tablo/export — ülke aylık tabloyu xlsx indir."""
    import io
    from flask import jsonify, request, send_file
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

    body = request.get_json(silent=True) or {}
    ulke = str(body.get('ulke') or '').strip().lower()
    yil = str(body.get('yil') or '').strip()
    if not re.fullmatch(r'[a-z]{2}', ulke) or ulke == 'ba':
        return jsonify({'success': False, 'error': 'Geçerli ülke kodu girin'}), 400
    if not re.fullmatch(r'20\d{2}', yil):
        return jsonify({'success': False, 'error': 'Geçerli yıl girin'}), 400
    rapor = body.get('rapor') if isinstance(body.get('rapor'), dict) else None
    if not rapor or not rapor.get('aylar'):
        rapor = tablo_yukle(ulke)
    aylar, schema = tablo_yil_aylar(rapor, yil)
    want = set(donem_ay_keys(yil, body.get('donem'), body.get('aylar')))
    if want:
        aylar = [a for a in aylar if a.get('ay') in want]
    if not aylar:
        return jsonify({'success': False, 'error': 'Seçilen dönemde ay yok'}), 400
    items = schema.get('items') or []

    wb = Workbook()
    ws = wb.active
    ws.title = (schema.get('title') or ulke)[:31]
    thin = Border(
        left=Side(style='thin', color='EDE8DF'),
        right=Side(style='thin', color='EDE8DF'),
        top=Side(style='thin', color='EDE8DF'),
        bottom=Side(style='thin', color='EDE8DF'),
    )
    ink = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
    lab = Font(name='Calibri', bold=True, size=10)
    fill_ink = PatternFill('solid', fgColor='1A1916')
    fill_teal = PatternFill('solid', fgColor='0F766E')
    fill_tot = PatternFill('solid', fgColor='F0FDFA')
    fill_ciro = PatternFill('solid', fgColor='FFFBEB')
    right = Alignment(horizontal='right', vertical='center')

    ws.cell(1, 1, schema.get('title') or ulke).font = Font(name='Calibri', bold=True, color='0F766E', size=10)
    ws.cell(2, 1, f"Logistics cost · {yil}").font = Font(name='Calibri', bold=True, size=16)
    donem_etiket = str(body.get('donem_etiket') or '').strip()
    if donem_etiket and donem_etiket != yil:
        ws.cell(2, 1, f"Logistics cost · {donem_etiket}").font = Font(name='Calibri', bold=True, size=16)
    ws.cell(4, 1, 'Item').font = ink
    ws.cell(4, 1).fill = fill_ink
    ay_ad = ('Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec')
    n = len(aylar)
    for i, a in enumerate(aylar):
        c = 2 + i * 2
        try:
            mon = int(str(a.get('ay') or '')[5:7])
        except ValueError:
            mon = i + 1
        label = ay_ad[mon - 1] if 1 <= mon <= 12 else str(a.get('ay') or '')
        ws.merge_cells(start_row=4, start_column=c, end_row=4, end_column=c + 1)
        ws.cell(4, c, label).font = ink
        ws.cell(4, c).fill = fill_ink
        ws.cell(4, c).alignment = Alignment(horizontal='center')
        ws.cell(5, c, 'Qty').font = Font(name='Calibri', bold=True, color='FFFFFF', size=9)
        ws.cell(5, c + 1, 'Cost €').font = Font(name='Calibri', bold=True, color='FFFFFF', size=9)
        ws.cell(5, c).fill = fill_teal
        ws.cell(5, c + 1).fill = fill_teal
    tc = 2 + n * 2
    tot_label = f'{donem_etiket} Total' if donem_etiket else f'{yil} Total'
    ws.merge_cells(start_row=4, start_column=tc, end_row=4, end_column=tc + 1)
    ws.cell(4, tc, tot_label).font = ink
    ws.cell(4, tc).fill = fill_ink
    ws.cell(5, tc, 'Qty').font = Font(name='Calibri', bold=True, color='FFFFFF', size=9)
    ws.cell(5, tc + 1, 'Cost €').font = Font(name='Calibri', bold=True, color='FFFFFF', size=9)
    ws.cell(5, tc).fill = fill_teal
    ws.cell(5, tc + 1).fill = fill_teal

    r = 6
    last_g = object()
    for it in items:
        g = it.get('group') or ''
        if g and g != last_g:
            ws.cell(r, 1, g).font = lab
            r += 1
            last_g = g
        ws.cell(r, 1, ('    ' if g else '') + it['label']).font = Font(name='Calibri', size=10)
        q_sum = t_sum = 0.0
        for i, a in enumerate(aylar):
            c = 2 + i * 2
            item = a.get(it['kod']) or {}
            qq = float(item.get('miktar') or 0)
            tt = float(item.get('tutar') or 0)
            q_sum += qq
            t_sum += tt
            ws.cell(r, c, qq).number_format = '#,##0.##'
            ws.cell(r, c + 1, tt).number_format = '#,##0.00'
            ws.cell(r, c).alignment = right
            ws.cell(r, c + 1).alignment = right
            for cc in (c, c + 1):
                ws.cell(r, cc).border = thin
        ws.cell(r, tc, q_sum).number_format = '#,##0.##'
        ws.cell(r, tc + 1, t_sum).number_format = '#,##0.00'
        ws.cell(r, tc).alignment = right
        ws.cell(r, tc + 1).alignment = right
        r += 1

    # Total / Ciro / Oran
    def _write_cost_row(row, label, vals, fill, italic=False, bold=False):
        cell = ws.cell(row, 1, label)
        cell.font = Font(name='Calibri', bold=True, italic=italic, size=10)
        cell.fill = fill
        for i, v in enumerate(vals):
            c = 2 + i * 2
            ws.cell(row, c).fill = fill
            ws.cell(row, c).border = thin
            e = ws.cell(row, c + 1, v)
            e.fill = fill
            e.border = thin
            e.alignment = right
            e.font = Font(name='Calibri', bold=bold or True, italic=italic, size=10)
            if isinstance(v, (int, float)):
                e.number_format = '#,##0.00'
        ws.cell(row, tc).fill = fill
        ws.cell(row, tc).border = thin
        e = ws.cell(row, tc + 1, sum(vals) if all(isinstance(x, (int, float)) for x in vals) else None)
        e.fill = fill
        e.border = thin
        e.alignment = right
        e.font = Font(name='Calibri', bold=True, italic=italic, size=10)
        if isinstance(e.value, (int, float)):
            e.number_format = '#,##0.00'

    loj_vals = [float(a.get('lojistik') or 0) for a in aylar]
    _write_cost_row(r, 'Total Cost', loj_vals, fill_tot, bold=True)
    r += 1
    ciro_vals = [float(a.get('ciro_eur') or 0) for a in aylar]
    _write_cost_row(r, 'Turnover (Ciro)', ciro_vals, fill_ciro, italic=True)
    r += 1
    oran_vals = []
    for loj, ciro in zip(loj_vals, ciro_vals):
        if ciro > 0:
            oran_vals.append(f'{(loj / ciro * 100):.2f}%'.replace('.', ','))
        else:
            oran_vals.append(None)
    ws.cell(r, 1, 'Logistics / Turnover').font = Font(name='Calibri', bold=True, size=10)
    for i, txt in enumerate(oran_vals):
        c = 2 + i * 2
        ws.cell(r, c + 1, txt).alignment = right
        ws.cell(r, c + 1).font = Font(name='Calibri', bold=True, size=10)
    loj_sum, ciro_sum = sum(loj_vals), sum(ciro_vals)
    ws.cell(r, tc + 1, f'{(loj_sum / ciro_sum * 100):.2f}%'.replace('.', ',') if ciro_sum > 0 else None).font = Font(name='Calibri', bold=True, size=10)
    ws.cell(r, tc + 1).alignment = right

    ws.column_dimensions['A'].width = 28
    for i in range(2, tc + 2):
        ws.column_dimensions[chr(64 + i) if i < 27 else 'A'].width = 11
    # simpler widths
    from openpyxl.utils import get_column_letter
    for i in range(2, tc + 2):
        ws.column_dimensions[get_column_letter(i)].width = 11

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    slug = ''
    if donem_etiket and donem_etiket != yil:
        slug = '_' + re.sub(r'[^A-Za-z0-9]+', '', donem_etiket)[:16]
    return send_file(
        buf,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f'{ulke}_lojistik_{yil}{slug}.xlsx',
    )


def seed_nl_from_excel(xlsx_path=None):
    """LOGISTICS COSTS.xlsx Netherland 2026 bloğundan tablo + ciro doldur."""
    from openpyxl import load_workbook
    path = xlsx_path or os.path.join(_BASE, 'Logistic Cost', 'LOGISTICS COSTS.xlsx')
    wb = load_workbook(path, data_only=True)
    ws = wb['Netherland']
    # qty/cost kolon çiftleri: Jan..Jun
    cols = [
        ('2026-01', 5, 6),
        ('2026-02', 7, 8),
        ('2026-03', 9, 10),
        ('2026-04', 14, 15),
        ('2026-05', 16, 17),
        ('2026-06', 18, 19),
    ]
    row_map = {
        'inbound': 22,
        'storage': 23,
        'outbound': 24,
        'transport': 25,
    }
    ciro_row = 28
    aylar = []
    ciro = {}
    for ay, qc, cc in cols:
        a = {'ay': ay, 'ciro_eur': 0.0}
        for kod, rr in row_map.items():
            q = ws.cell(rr, qc).value
            c = ws.cell(rr, cc).value
            a[kod] = {
                'miktar': float(q or 0),
                'tutar': float(c or 0),
            }
        # Total/Ciro Excel'de Quantity kolonunda
        a['ciro_eur'] = float(ws.cell(ciro_row, qc).value or 0)
        ciro[ay] = a['ciro_eur']
        aylar.append(a)
    tablo_kaydet('nl', aylar)
    return {'aylar': len(aylar), 'ciro': ciro}
