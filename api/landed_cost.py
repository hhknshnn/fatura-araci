import io
import re

from flask import jsonify, request, send_file

from api.db import get_conn


KURUMSAL_ULKELER = {
    'SIRBİSTAN', 'BOSNA', 'GÜRCİSTAN', 'KOSOVA', 'MAKEDONYA',
    'BELÇİKA', 'ALMANYA', 'HOLLANDA', 'KAZAKİSTAN',
}


def _to_float(value):
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _norm(value):
    return str(value or '').strip().upper()


def _depo_from_fatura(fatura_no):
    return 'ANT' if str(fatura_no or '').startswith('ANT') else 'IHR'


def _natural_sort_key(value):
    parts = re.split(r'(\d+)', str(value or ''))
    return [int(part) if part.isdigit() else part.lower() for part in parts]


def _is_kurumsal(row):
    tip = str(row.get('musteri_tipi') or '').strip().lower()
    return tip == 'kurumsal' or (not tip and _norm(row.get('ulke')) in KURUMSAL_ULKELER)


def _cost_parts(rows):
    fatura = sum(_to_float(r.get('fatura_bedeli_eur')) for r in rows)
    operasyon = sum(
        _to_float(r.get('ihracat_beyanname_eur')) +
        _to_float(r.get('arac_bekleme')) +
        _to_float(r.get('brokerage_eur'))
        for r in rows
    )
    navlun = sum(_to_float(r.get('navlun_eur')) for r in rows)
    vergi = sum(_to_float(r.get('gumruk_vergisi_eur')) for r in rows)
    sigorta = sum(_to_float(r.get('sigorta_eur')) for r in rows)
    landed = operasyon + navlun + vergi + sigorta
    return {
        'fatura_eur': fatura,
        'operasyon_eur': operasyon,
        'navlun_eur': navlun,
        'vergi_eur': vergi,
        'sigorta_eur': sigorta,
        'landed_cost_eur': landed,
        'oran': round((landed / fatura) * 100, 2) if fatura else 0,
    }


def _sefer_count(rows):
    groups = set()
    singles = 0
    for row in rows:
        if row.get('sefer_id'):
            groups.add(row.get('sefer_id'))
        else:
            singles += 1
    return singles + len(groups)


def _has_landed_cost_ready(row):
    return _to_float(row.get('brokerage_eur')) > 0


def _split_landed_cost_rows(rows):
    ready = []
    pending = []
    for row in rows:
        if _has_landed_cost_ready(row):
            ready.append(row)
        else:
            pending.append(row)
    return ready, pending


def _query_rows():
    countries = [c.strip().upper() for c in request.args.get('countries', '').split(',') if c.strip()]
    date_from = request.args.get('date_from')
    date_to = request.args.get('date_to')
    depo = request.args.get('depo')
    group_type = request.args.get('group_type', 'all')

    conn = get_conn()
    cur = conn.cursor()
    query = '''
        SELECT id, ihracat_dosya_no, fatura_no, ulke, nakliye_firmasi, plaka,
               fatura_bedeli_eur, navlun_eur, sigorta_eur, yukleme_tarihi,
               ihracat_beyanname_eur, arac_bekleme, brokerage_eur,
               gumruk_vergisi_eur, kdv_eur, musteri_tipi, sefer_id, durum
        FROM shipments
        WHERE 1=1
    '''
    params = []
    if date_from:
        query += ' AND yukleme_tarihi >= %s'
        params.append(date_from)
    if date_to:
        query += ' AND yukleme_tarihi <= %s'
        params.append(date_to)
    if group_type == 'single':
        query += ' AND sefer_id IS NULL'
    elif group_type == 'grouped':
        query += ' AND sefer_id IS NOT NULL'
    query += ' ORDER BY ulke, yukleme_tarihi NULLS LAST, id'

    cur.execute(query, params)
    cols = [d[0] for d in cur.description]
    rows = [dict(zip(cols, r)) for r in cur.fetchall()]
    cur.close()
    conn.close()

    filtered = []
    for row in rows:
        if not _is_kurumsal(row):
            continue
        if countries and _norm(row.get('ulke')) not in countries:
            continue
        if depo and _depo_from_fatura(row.get('fatura_no')) != depo:
            continue
        filtered.append(row)
    return filtered


def _pending_payload(rows):
    by_status = {}
    by_country = {}
    detail = []

    for row in rows:
        status = row.get('durum') or 'Belirsiz'
        country = row.get('ulke') or 'Belirsiz'
        by_status[status] = by_status.get(status, 0) + 1
        by_country[country] = by_country.get(country, 0) + 1
        detail.append({
            'id': row.get('id'),
            'ihracat_dosya_no': row.get('ihracat_dosya_no'),
            'fatura_no': row.get('fatura_no'),
            'ulke': country,
            'depo': _depo_from_fatura(row.get('fatura_no')),
            'yukleme_tarihi': str(row.get('yukleme_tarihi') or ''),
            'durum': status,
            'brokerage_eur': _to_float(row.get('brokerage_eur')),
            'fatura_eur': _to_float(row.get('fatura_bedeli_eur')),
            'sefer_id': row.get('sefer_id'),
        })
    detail.sort(key=lambda item: _natural_sort_key(item.get('ihracat_dosya_no')), reverse=True)

    return {
        'summary': {
            'fatura_sayisi': len(rows),
            'sefer_sayisi': _sefer_count(rows),
            'fatura_eur': sum(_to_float(r.get('fatura_bedeli_eur')) for r in rows),
            'by_status': [
                {'durum': status, 'sayi': count}
                for status, count in sorted(by_status.items(), key=lambda item: item[0])
            ],
            'by_country': [
                {'ulke': country, 'sayi': count}
                for country, count in sorted(by_country.items(), key=lambda item: item[0])
            ],
        },
        'detail': detail,
    }


def _build_payload(rows, pending_rows=None):
    pending_rows = pending_rows or []
    summary = _cost_parts(rows)
    summary['fatura_sayisi'] = len(rows)
    summary['sefer_sayisi'] = _sefer_count(rows)
    summary['ortalama_sefer_maliyeti_eur'] = (
        summary['landed_cost_eur'] / summary['sefer_sayisi']
        if summary['sefer_sayisi'] else 0
    )

    by_country = {}
    for row in rows:
        by_country.setdefault(row.get('ulke') or 'Belirsiz', []).append(row)

    countries = []
    for country, country_rows in by_country.items():
        data = _cost_parts(country_rows)
        data.update({
            'ulke': country,
            'fatura_sayisi': len(country_rows),
            'sefer_sayisi': _sefer_count(country_rows),
        })
        countries.append(data)
    countries.sort(key=lambda item: item['landed_cost_eur'], reverse=True)

    monthly = {}
    for row in rows:
        month = str(row.get('yukleme_tarihi') or '')[:7] or 'Tarihsiz'
        monthly.setdefault(month, []).append(row)
    months = []
    for month, month_rows in monthly.items():
        data = _cost_parts(month_rows)
        data['month'] = month
        months.append(data)
    months.sort(key=lambda item: item['month'])

    detail = []
    for row in rows:
        costs = _cost_parts([row])
        detail.append({
            'id': row.get('id'),
            'ihracat_dosya_no': row.get('ihracat_dosya_no'),
            'fatura_no': row.get('fatura_no'),
            'ulke': row.get('ulke'),
            'depo': _depo_from_fatura(row.get('fatura_no')),
            'yukleme_tarihi': str(row.get('yukleme_tarihi') or ''),
            'sefer_id': row.get('sefer_id'),
            **costs,
        })

    return {
        'summary': summary,
        'countries': countries,
        'months': months,
        'detail': detail,
        'pending': _pending_payload(pending_rows),
        'cost_labels': {
            'operasyon_eur': 'Operasyon',
            'navlun_eur': 'Navlun',
            'vergi_eur': 'Vergi',
            'sigorta_eur': 'Sigorta',
        },
    }


def landed_cost_get():
    ready_rows, pending_rows = _split_landed_cost_rows(_query_rows())
    return jsonify({'success': True, **_build_payload(ready_rows, pending_rows)})


def landed_cost_export():
    ready_rows, pending_rows = _split_landed_cost_rows(_query_rows())
    payload = _build_payload(ready_rows, pending_rows)

    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        return jsonify({'success': False, 'error': 'openpyxl kurulu değil'}), 500

    wb = openpyxl.Workbook()
    header_fill = PatternFill('solid', fgColor='1F3864')
    header_font = Font(bold=True, color='FFFFFF')
    money_fmt = '#,##0.00 €'
    pct_fmt = '0.00%'

    def write_table(ws, headers, rows_data):
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=header[0])
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')
        for r_idx, item in enumerate(rows_data, start=2):
            for c_idx, header in enumerate(headers, start=1):
                key = header[1]
                value = item.get(key, '')
                if key == 'oran':
                    value = (value or 0) / 100
                cell = ws.cell(row=r_idx, column=c_idx, value=value)
                if key.endswith('_eur') or key == 'landed_cost_eur':
                    cell.number_format = money_fmt
                if key == 'oran':
                    cell.number_format = pct_fmt
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 18

    ws = wb.active
    ws.title = 'Ozet'
    summary_rows = [{**payload['summary'], 'ulke': 'Toplam Kurumsal'}]
    write_table(ws, [
        ('Kapsam', 'ulke'),
        ('Fatura EUR', 'fatura_eur'),
        ('Landed Cost EUR', 'landed_cost_eur'),
        ('Oran', 'oran'),
        ('Operasyon EUR', 'operasyon_eur'),
        ('Navlun EUR', 'navlun_eur'),
        ('Vergi EUR', 'vergi_eur'),
        ('Sigorta EUR', 'sigorta_eur'),
        ('Fatura Sayısı', 'fatura_sayisi'),
        ('Sefer Sayısı', 'sefer_sayisi'),
    ], summary_rows)

    ws = wb.create_sheet('Ulke Bazli')
    write_table(ws, [
        ('Ülke', 'ulke'),
        ('Fatura EUR', 'fatura_eur'),
        ('Landed Cost EUR', 'landed_cost_eur'),
        ('Oran', 'oran'),
        ('Operasyon EUR', 'operasyon_eur'),
        ('Navlun EUR', 'navlun_eur'),
        ('Vergi EUR', 'vergi_eur'),
        ('Sigorta EUR', 'sigorta_eur'),
        ('Fatura Sayısı', 'fatura_sayisi'),
        ('Sefer Sayısı', 'sefer_sayisi'),
    ], payload['countries'])

    ws = wb.create_sheet('Detay')
    write_table(ws, [
        ('Dosya No', 'ihracat_dosya_no'),
        ('Fatura No', 'fatura_no'),
        ('Ülke', 'ulke'),
        ('Depo', 'depo'),
        ('Yükleme', 'yukleme_tarihi'),
        ('Grup', 'sefer_id'),
        ('Fatura EUR', 'fatura_eur'),
        ('Landed Cost EUR', 'landed_cost_eur'),
        ('Oran', 'oran'),
        ('Operasyon EUR', 'operasyon_eur'),
        ('Navlun EUR', 'navlun_eur'),
        ('Vergi EUR', 'vergi_eur'),
        ('Sigorta EUR', 'sigorta_eur'),
    ], payload['detail'])

    ws = wb.create_sheet('Bekleyenler')
    write_table(ws, [
        ('Dosya No', 'ihracat_dosya_no'),
        ('Fatura No', 'fatura_no'),
        ('Ülke', 'ulke'),
        ('Depo', 'depo'),
        ('Yükleme', 'yukleme_tarihi'),
        ('Durum', 'durum'),
        ('Grup', 'sefer_id'),
        ('Fatura EUR', 'fatura_eur'),
        ('Brokerage EUR', 'brokerage_eur'),
    ], payload['pending']['detail'])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='landed_cost_raporu.xlsx',
    )
