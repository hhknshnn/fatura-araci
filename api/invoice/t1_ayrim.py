"""T1 Ayrımı — INV+PL dosyasının M–R sütunlarını Masterfile + T1 PDF'inden doldurur.

Mantık:
  • Masterfile'da ağırlık koli bazında yazılıdır. Asorti kolide toplam brüt/net
    ilk satırda durur, devam satırlarında 0'dır. Bir koli grubunun ağırlığı
    içindeki satırlara ADET oranında paylaştırılır.
  • Masterfile'daki her (Document Number, Item Number) grubu = T1'de bir kalem.
    Eşleştirme GTİP-6 + brüt + net + kap adedi üzerinden yapılır.
  • Fatura satırı ürün koduyla (ITEM CODE = Model Code) masterfile'a bağlanır.

Doldurulan sütunlar:
  M Declarion No | N Index | O Total Gross | P Total Net | Q Unit Gross | R Unit Net

Kurallar bu modüle gömülüdür; kullanıcının referans örnek fatura yüklemesi
gerekmez. Kuralların çıkarıldığı referans dosya:
  templates/ornek_invoice_t1_ayrimi.xlsx
"""

import io
import re
from collections import defaultdict

import openpyxl
import pdfplumber

try:
    import xlrd  # eski .XLS talep formları için
except ImportError:
    xlrd = None

# INV sayfası sütun indeksleri (1 tabanlı)
COL_ITEM_CODE = 3    # C — ITEM CODE
COL_QTY       = 7    # G — QTY
COL_DECL      = 13   # M — Declarion No
COL_INDEX     = 14   # N — Index
COL_TGROSS    = 15   # O — Total Gross
COL_TNET      = 16   # P — Total Net
COL_UGROSS    = 17   # Q — Unit Gross
COL_UNET      = 18   # R — Unit Net

# Masterfile sütunları (1 tabanlı)
MF_MODEL_CODE = 4    # D — Model Code
MF_BOX_QTY    = 11   # K — Box Quantity
MF_TOTAL_QTY  = 13   # M — Total Quantity
MF_DOC_NO     = 15   # O — Document Number
MF_ITEM_NO    = 17   # Q — Item Number
MF_GTIP       = 18   # R — Gtip
MF_GROSS      = 21   # U — Total Gross KG
MF_NET        = 22   # V — Total Net KG

_TOL = 0.02          # ağırlık eşleştirme toleransı (kg)


def _norm_code(v):
    return str(v).strip() if v is not None else ''


def _num(v):
    if v is None or v == '':
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(',', '.').strip())
    except ValueError:
        return 0.0


def _gtip6(v):
    """'\\n7013.28.90.00.00' → '701328'"""
    return re.sub(r'\D', '', str(v or ''))[:6]


# ── T1 PDF ───────────────────────────────────────────────────────────────────

def parse_t1_pdf(pdf_bytes):
    """T1 transit beyannamesinden kalem listesini çıkarır."""
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        text = '\n'.join((p.extract_text() or '') for p in pdf.pages)

    lines = text.split('\n')
    items = []
    for i, satir in enumerate(lines):
        if not satir.startswith('Eşya Kalem Numarası'):
            continue
        if i + 1 >= len(lines):
            continue
        head = lines[i + 1]
        blok = '\n'.join(lines[i + 1:i + 22])

        m_no  = re.match(r'\s*(\d+)\s+(\d+)\s', head)
        m_pkg = re.search(r'/\w+,\s*(\d+)\s*,', head)
        m_gr  = re.search(r'Brüt Ağırlık \(Kg\)[^\n]*\n\s*([\d.,]+)', blok)
        m_net = re.search(r'\n\s*\d+\s+\S+\s+\w{2}\s+(\d{6})\s+([\d.,]+)', blok)
        if not (m_no and m_gr and m_net):
            continue

        items.append({
            'kalem': int(m_no.group(1)),
            'kap':   int(m_pkg.group(1)) if m_pkg else 0,
            'gtip6': m_net.group(1),
            'brut':  _num(m_gr.group(1)),
            'net':   _num(m_net.group(2)),
        })
    return items


# ── T1 Talep Formu (opsiyonel) ───────────────────────────────────────────────

_TALEP_BASLIK = {
    'kalem':     ('KALEM',),
    'beyanname': ('BEYANNAME NO', 'BEYANNAME'),
    'satir':     ('SATIR', 'SIRA'),
    'gtip':      ('GTIP', 'GTİP'),
    'kap':       ('KAP',),
    'brut':      ('BRUT KG', 'BRÜT KG', 'BRUT', 'BRÜT'),
    'net':       ('NET KG', 'NET'),
    'adet':      ('ADET',),
}


def _talep_satirlari(dosya_bytes):
    """Talep formunu (.xls veya .xlsx) satır listesi olarak okur."""
    if xlrd is not None:
        try:
            sh = xlrd.open_workbook(file_contents=dosya_bytes).sheet_by_index(0)
            return [[sh.cell_value(r, c) for c in range(sh.ncols)]
                    for r in range(sh.nrows)]
        except Exception:
            pass
    wb = openpyxl.load_workbook(io.BytesIO(dosya_bytes), data_only=True)
    ws = wb.worksheets[0]
    return [[ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            for r in range(1, ws.max_row + 1)]


def parse_talep_formu(dosya_bytes):
    """T1 talep formundan (beyanname no, satır) → kalem no eşlemesini çıkarır.

    Dönen: {(beyanname, satir): {kalem, gtip, kap, brut, net, adet}}
    """
    satirlar = _talep_satirlari(dosya_bytes)

    # Başlık satırını ve sütun konumlarını bul
    sutun = None
    baslangic = None
    for i, satir in enumerate(satirlar):
        ustler = [str(v or '').strip().upper() for v in satir]
        bulunan = {}
        for alan, adaylar in _TALEP_BASLIK.items():
            for j, u in enumerate(ustler):
                if u in adaylar:
                    bulunan[alan] = j
                    break
        if 'kalem' in bulunan and 'beyanname' in bulunan and 'satir' in bulunan:
            sutun = bulunan
            baslangic = i + 1
            break

    if not sutun:
        raise ValueError('Talep formunda KALEM / BEYANNAME NO / SATIR başlıkları bulunamadı')

    def al(satir, alan):
        j = sutun.get(alan)
        return satir[j] if j is not None and j < len(satir) else None

    kayitlar = {}
    for satir in satirlar[baslangic:]:
        decl = _norm_code(al(satir, 'beyanname'))
        if not decl or not re.match(r'^\d{2}\w+', decl):
            continue
        kalem = _num(al(satir, 'kalem'))
        satir_no = _num(al(satir, 'satir'))
        if not kalem or not satir_no:
            continue
        kayitlar[(decl, int(satir_no))] = {
            'kalem': int(kalem),
            'gtip':  _gtip6(al(satir, 'gtip')),
            'kap':   _num(al(satir, 'kap')),
            'brut':  _num(al(satir, 'brut')),
            'net':   _num(al(satir, 'net')),
            'adet':  _num(al(satir, 'adet')),
        }
    if not kayitlar:
        raise ValueError('Talep formunda beyanname satırı okunamadı')
    return kayitlar


# ── Masterfile ───────────────────────────────────────────────────────────────

def parse_masterfile(xlsx_bytes):
    """Masterfile'ı koli gruplarına ayırır ve ürün bazında ağırlık dağıtır.

    Dönen: (urunler, beyanname_gruplari)
      urunler[model_code] = {qty, brut, net, decl, item_no}
      beyanname_gruplari[(decl, item_no)] = {brut, net, kap, gtip}
    """
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    ws = wb.worksheets[0]

    satirlar = []
    for r in range(2, ws.max_row + 1):
        kod = _norm_code(ws.cell(r, MF_MODEL_CODE).value)
        if not kod:
            continue
        satirlar.append({
            'kod':   kod,
            'kap':   _num(ws.cell(r, MF_BOX_QTY).value),
            'adet':  _num(ws.cell(r, MF_TOTAL_QTY).value),
            'decl':  _norm_code(ws.cell(r, MF_DOC_NO).value),
            'item':  ws.cell(r, MF_ITEM_NO).value,
            'gtip':  _gtip6(ws.cell(r, MF_GTIP).value),
            'brut':  _num(ws.cell(r, MF_GROSS).value),
            'net':   _num(ws.cell(r, MF_NET).value),
        })

    # Koli grupları: brüt > 0 olan satır yeni grup başlatır, sonraki 0'lar ona dahildir
    gruplar = []
    aktif = None
    for s in satirlar:
        if s['brut'] > 0 or aktif is None:
            aktif = {'satirlar': [], 'brut': s['brut'], 'net': s['net'],
                     'kap': s['kap'], 'decl': s['decl'], 'item': s['item'],
                     'gtip': s['gtip']}
            gruplar.append(aktif)
        aktif['satirlar'].append(s)

    # Beyanname (decl, item_no) bazında toplamlar — T1 eşleştirmesi için
    beyannameler = defaultdict(lambda: {'brut': 0.0, 'net': 0.0, 'kap': 0.0, 'gtip': set()})
    for g in gruplar:
        b = beyannameler[(g['decl'], g['item'])]
        b['brut'] += g['brut']
        b['net']  += g['net']
        b['kap']  += g['kap']
        b['gtip'].add(g['gtip'])

    # Ürün bazında ağırlık dağıtımı — koli ağırlığı adet oranında paylaştırılır
    urunler = defaultdict(lambda: {'qty': 0.0, 'brut': 0.0, 'net': 0.0,
                                   'decl': None, 'item': None})
    for g in gruplar:
        toplam_adet = sum(s['adet'] for s in g['satirlar'])
        for s in g['satirlar']:
            u = urunler[s['kod']]
            u['qty'] += s['adet']
            if toplam_adet:
                pay = s['adet'] / toplam_adet
                u['brut'] += g['brut'] * pay
                u['net']  += g['net'] * pay
            u['decl'] = g['decl']
            u['item'] = g['item']

    return urunler, beyannameler


# ── T1 ↔ Masterfile eşleştirme ───────────────────────────────────────────────

def eslestir_kalemler(beyannameler, t1_items, talep=None):
    """(decl, item_no) → T1 kalem numarası. Eşleşmeyenler uyarı listesine düşer.

    Talep formu verilmişse kalem numarası doğrudan oradan alınır (beyanname no +
    satır no zaten yazılıdır); kalanlar T1 PDF'i üzerinden ağırlık/GTİP ile eşlenir.
    """
    kalem_map = {}
    uyarilar = []
    kullanilan = set()

    if talep:
        for anahtar in list(beyannameler):
            kayit = talep.get((anahtar[0], int(_num(anahtar[1]))))
            if kayit:
                kalem_map[anahtar] = kayit['kalem']
                kullanilan.add(kayit['kalem'])
        beyannameler = {k: v for k, v in beyannameler.items() if k not in kalem_map}

    # Sıkıdan gevşeğe: önce GTİP+brüt+net+kap, sonra kap serbest, sonra brüt+GTİP
    kriterler = [
        lambda t, b: (t['gtip6'] in b['gtip'] and abs(t['brut'] - b['brut']) < _TOL
                      and abs(t['net'] - b['net']) < _TOL and t['kap'] == round(b['kap'])),
        lambda t, b: (t['gtip6'] in b['gtip'] and abs(t['brut'] - b['brut']) < _TOL
                      and abs(t['net'] - b['net']) < _TOL),
        lambda t, b: (t['gtip6'] in b['gtip'] and abs(t['brut'] - b['brut']) < _TOL),
    ]

    kalanlar = list(beyannameler.items())
    for kriter in kriterler:
        yeni_kalanlar = []
        for anahtar, b in kalanlar:
            adaylar = [t for t in t1_items
                       if t['kalem'] not in kullanilan and kriter(t, b)]
            if len(adaylar) == 1:
                kalem_map[anahtar] = adaylar[0]['kalem']
                kullanilan.add(adaylar[0]['kalem'])
            else:
                yeni_kalanlar.append((anahtar, b))
        kalanlar = yeni_kalanlar

    for anahtar, b in kalanlar:
        uyarilar.append(
            f"Beyanname {anahtar[0]} / kalem {anahtar[1]} "
            f"(brüt {b['brut']:.2f} kg) T1'de eşleştirilemedi"
        )

    bos_kalanlar = [t['kalem'] for t in t1_items if t['kalem'] not in kullanilan]
    if bos_kalanlar:
        uyarilar.append(
            "T1'de karşılığı bulunamayan kalem(ler): " +
            ', '.join(str(k) for k in bos_kalanlar)
        )

    return kalem_map, uyarilar


def capraz_kontrol(beyannameler, t1_items, talep):
    """Masterfile / T1 PDF / talep formu arasındaki sapmaları listeler."""
    uyarilar = []
    if not talep:
        return uyarilar

    t1_kalem = {t['kalem']: t for t in t1_items}

    for anahtar, b in beyannameler.items():
        kayit = talep.get((anahtar[0], int(_num(anahtar[1]))))
        if not kayit:
            uyarilar.append(
                f"Beyanname {anahtar[0]} / satır {anahtar[1]} talep formunda yok"
            )
            continue
        etiket = f"Beyanname {anahtar[0]}-{anahtar[1]}"
        if abs(b['brut'] - kayit['brut']) > _TOL:
            uyarilar.append(f"{etiket}: brüt masterfile {b['brut']:.2f} / talep formu {kayit['brut']:.2f} kg")
        if abs(b['net'] - kayit['net']) > _TOL:
            uyarilar.append(f"{etiket}: net masterfile {b['net']:.2f} / talep formu {kayit['net']:.2f} kg")
        if round(b['kap']) != round(kayit['kap']):
            uyarilar.append(f"{etiket}: kap masterfile {b['kap']:g} / talep formu {kayit['kap']:g}")
        if kayit['gtip'] and kayit['gtip'] not in b['gtip']:
            uyarilar.append(f"{etiket}: GTİP masterfile {'/'.join(sorted(b['gtip']))} / talep formu {kayit['gtip']}")

        if not t1_items:
            continue  # T1 PDF verilmemiş — çapraz kontrol yapılamaz
        t1 = t1_kalem.get(kayit['kalem'])
        if t1 is None:
            uyarilar.append(f"{etiket}: talep formundaki {kayit['kalem']}. kalem T1 PDF'inde yok")
        else:
            if abs(t1['brut'] - kayit['brut']) > _TOL:
                uyarilar.append(f"T1 kalem {kayit['kalem']}: brüt T1 {t1['brut']:.2f} / talep formu {kayit['brut']:.2f} kg")
            if abs(t1['net'] - kayit['net']) > _TOL:
                uyarilar.append(f"T1 kalem {kayit['kalem']}: net T1 {t1['net']:.2f} / talep formu {kayit['net']:.2f} kg")
            if t1['gtip6'] and kayit['gtip'] and t1['gtip6'] != kayit['gtip']:
                uyarilar.append(f"T1 kalem {kayit['kalem']}: GTİP T1 {t1['gtip6']} / talep formu {kayit['gtip']}")

    fazla = [k for k in talep if k not in beyannameler]
    if fazla:
        uyarilar.append(
            "Talep formunda olup masterfile'da olmayan beyanname satırı: " +
            ', '.join(f"{d}-{s}" for d, s in fazla[:10])
        )
    return uyarilar


# ── Dosya tanıma ─────────────────────────────────────────────────────────────

ROL_ADLARI = {
    'invpl':      'INV + PL',
    'masterfile': 'Masterfile',
    't1':         'T1 transit beyannamesi',
    'talep':      'T1 talep formu',
}

# Doldurma kuralları koda gömülüdür; referans örnek faturanın yüklenmesine gerek
# yoktur. Yanlışlıkla bırakılırsa girdi sanılmasın diye ayrıca tanınır.
ROL_ORNEK = 'ornek'


def _excel_ilk_hucreler(dosya_bytes, satir_limiti=60):
    """Excel'in ilk sayfasındaki hücreleri büyük harfli metin listesi olarak verir."""
    if dosya_bytes[:2] == b'PK':
        wb = openpyxl.load_workbook(io.BytesIO(dosya_bytes), data_only=True, read_only=True)
        ws = wb.worksheets[0]
        sayfalar = wb.sheetnames
        hucreler = []
        for i, satir in enumerate(ws.iter_rows(values_only=True)):
            if i >= satir_limiti:
                break
            hucreler += [str(v).strip().upper() for v in satir if v is not None]
        wb.close()
        return hucreler, [s.strip().upper() for s in sayfalar]

    if xlrd is None:
        raise ValueError('Eski .xls dosyası için xlrd kurulu değil')
    wb = xlrd.open_workbook(file_contents=dosya_bytes)
    sh = wb.sheet_by_index(0)
    hucreler = []
    for r in range(min(satir_limiti, sh.nrows)):
        hucreler += [str(sh.cell_value(r, c)).strip().upper()
                     for c in range(sh.ncols) if sh.cell_value(r, c) != '']
    return hucreler, [s.upper() for s in wb.sheet_names()]


def dosya_rolu(ad, dosya_bytes):
    """Dosyanın hangi girdi olduğunu içeriğinden tespit eder."""
    if dosya_bytes[:4] == b'%PDF':
        return 't1'

    try:
        hucreler, sayfalar = _excel_ilk_hucreler(dosya_bytes)
    except Exception:
        return None
    kume = set(hucreler)

    # Talep formu — KALEM + BEYANNAME NO + SATIR başlıkları
    if 'BEYANNAME NO' in kume and 'KALEM' in kume and {'SATIR', 'SIRA'} & kume:
        return 'talep'

    # Masterfile — Model Code + Document Number
    if 'MODEL CODE' in kume and ('DOCUMENT NUMBER' in kume or 'TOTAL GROSS KG' in kume):
        return 'masterfile'

    # Fatura benzeri dosya — hedef INV+PL mi, referans örnek fatura mı?
    fatura_gibi = ('INV' in sayfalar or 'IV' in sayfalar or 'DECLARION NO' in kume
                   or ('ITEM CODE' in kume and 'HS CODE' in kume))
    if fatura_gibi:
        # Hedef dosyanın sayfası 'INV'dir. Örnek/referans fatura ('IV' sayfası,
        # 'IV NO' sütunu) doldurulacak dosya değildir.
        if 'INV' in sayfalar and 'IV NO' not in kume:
            return 'invpl'
        return ROL_ORNEK

    return None


def siniflandir_dosyalar(dosyalar):
    """[(ad, bytes)] → ({rol: bytes}, {rol: ad}, [uyari])"""
    secilen = {}
    adlar = {}
    uyarilar = []

    for ad, veri in dosyalar:
        if not veri:
            continue
        rol = dosya_rolu(ad, veri)
        if rol is None:
            uyarilar.append(f"'{ad}' tanınamadı, dikkate alınmadı")
            continue
        if rol == ROL_ORNEK:
            uyarilar.append(
                f"'{ad}' referans örnek fatura olarak görünüyor, atlandı — "
                f"doldurma kuralları uygulamada gömülü, bu dosyayı yüklemeye gerek yok"
            )
            continue
        if rol in secilen:
            uyarilar.append(
                f"'{ad}' ile '{adlar[rol]}' aynı türde ({ROL_ADLARI[rol]}); "
                f"'{adlar[rol]}' kullanıldı"
            )
            continue
        secilen[rol] = veri
        adlar[rol] = ad

    return secilen, adlar, uyarilar


# ── Ana akış ─────────────────────────────────────────────────────────────────

def t1_ayrimi_uret(invpl_bytes, masterfile_bytes, t1_pdf_bytes=None, talep_bytes=None):
    """INV+PL dosyasının M–R sütunlarını doldurur.

    T1 PDF ve talep formundan en az biri gerekir. İkisi de verilirse kalem
    numaraları talep formundan alınır, T1 PDF'i çapraz kontrol için kullanılır.

    Dönen: (xlsx_bytes, ozet_dict)
    """
    if not t1_pdf_bytes and not talep_bytes:
        raise ValueError('T1 PDF veya talep formundan en az biri gerekli')

    t1_items = parse_t1_pdf(t1_pdf_bytes) if t1_pdf_bytes else []
    if t1_pdf_bytes and not t1_items:
        raise ValueError('T1 PDF\'inde kalem listesi okunamadı')

    talep = parse_talep_formu(talep_bytes) if talep_bytes else None

    urunler, beyannameler = parse_masterfile(masterfile_bytes)
    if not urunler:
        raise ValueError('Masterfile\'da ürün satırı bulunamadı')

    kalem_map, uyarilar = eslestir_kalemler(beyannameler, t1_items, talep)
    uyarilar += capraz_kontrol(beyannameler, t1_items, talep)

    wb = openpyxl.load_workbook(io.BytesIO(invpl_bytes))
    ws = wb['INV'] if 'INV' in wb.sheetnames else wb.worksheets[0]

    dolu = 0
    top_brut = top_net = 0.0
    for r in range(1, ws.max_row + 1):
        kod = _norm_code(ws.cell(r, COL_ITEM_CODE).value)
        qty = _num(ws.cell(r, COL_QTY).value)
        if not kod or qty <= 0 or kod not in urunler:
            if kod and qty > 0 and kod not in urunler:
                uyarilar.append(f"Satır {r}: '{kod}' masterfile'da bulunamadı")
            continue

        u = urunler[kod]
        if abs(u['qty'] - qty) > 0.001:
            uyarilar.append(
                f"Satır {r}: '{kod}' adedi faturada {qty:g}, "
                f"masterfile'da {u['qty']:g}"
            )

        kalem = kalem_map.get((u['decl'], u['item']))
        ws.cell(r, COL_DECL).value   = u['decl']
        ws.cell(r, COL_INDEX).value  = kalem
        ws.cell(r, COL_TGROSS).value = u['brut']
        ws.cell(r, COL_TNET).value   = u['net']
        ws.cell(r, COL_UGROSS).value = u['brut'] / qty
        ws.cell(r, COL_UNET).value   = u['net'] / qty

        dolu += 1
        top_brut += u['brut']
        top_net  += u['net']

    if not dolu:
        raise ValueError(
            'INV+PL dosyasında doldurulacak satır bulunamadı — ürün kodları '
            'masterfile ile eşleşmiyor. Doğru INV+PL dosyasını yüklediğinizden '
            'emin olun.'
        )

    cikti = io.BytesIO()
    wb.save(cikti)

    ozet = {
        'satir': dolu,
        'dosyalar': {},
        't1_kalem': len(t1_items),
        'beyanname_grubu': len(beyannameler),
        'eslesen_kalem': len(kalem_map),
        'toplam_brut': round(top_brut, 2),
        'toplam_net': round(top_net, 2),
        't1_toplam_brut': round(sum(t['brut'] for t in t1_items), 2),
        't1_toplam_net': round(sum(t['net'] for t in t1_items), 2),
        't1_toplam_kap': sum(t['kap'] for t in t1_items),
        'talep_kalem': len(talep) if talep else 0,
        'kaynak': 'talep formu' if talep else 'T1 PDF',
        'uyarilar': uyarilar,
    }
    if talep and not t1_items:
        ozet['t1_kalem'] = len(talep)
        ozet['t1_toplam_brut'] = round(sum(k['brut'] for k in talep.values()), 2)
        ozet['t1_toplam_net'] = round(sum(k['net'] for k in talep.values()), 2)
        ozet['t1_toplam_kap'] = int(sum(k['kap'] for k in talep.values()))
    return cikti.getvalue(), ozet
